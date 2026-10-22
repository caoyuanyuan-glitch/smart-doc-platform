import sys
import unittest
from pathlib import Path

BACKEND_ROOT = Path(__file__).resolve().parents[1]
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

from app.api.convert import (  # noqa: E402
    _append_bookmap_topics,
    _append_root_container_xml,
    _build_topic_hierarchy,
    _clean_title,
    _content_to_dita_xml,
    _infer_docx_heading_level,
    _extract_template_frontmatter_subset,
    _docx_to_markdown,
    _docx_numbered_run_continues,
    _docx_table_to_markdown,
    _looks_like_docx_heading_text,
    _parse_md_sections,
    _parse_content,
    _normalize_native_markdown,
    _postprocess_section_tree,
    _reshape_docx_sections,
    _should_emit_docx_list_marker,
    _section_dita_type,
    _generated_topic_code,
    _copy_docx_media_to_output,
    _download_images_for_output,
    _topic_placeholder,
    _rewrite_template_frontmatter,
    _topic_kind_for_section,
    _idml_to_markdown,
    _idml_table_to_markdown,
    _idml_walk,
    _resolve_idml_source,
    _flatten_sections,
    _markdown_to_lossless_csv,
    _csv_to_markdown,
)

TOOLS_ROOT = BACKEND_ROOT / "tools"
if str(TOOLS_ROOT) not in sys.path:
    sys.path.insert(0, str(TOOLS_ROOT))

from verify_word_to_dita_package import _parse_dita_package  # noqa: E402


class GeneratedTopicCodeTest(unittest.TestCase):
    def test_cover_uses_ctt_prefix(self):
        self.assertTrue(_generated_topic_code(1, "cover").startswith("CTT"))

    def test_root_chapter_uses_cto_prefix(self):
        self.assertTrue(_generated_topic_code(2, "chapter").startswith("CTO"))

    def test_concept_uses_dto_prefix(self):
        self.assertTrue(_generated_topic_code(3, "concept").startswith("DTO"))

    def test_task_uses_dto_prefix(self):
        self.assertTrue(_generated_topic_code(4, "task").startswith("DTO"))

    def test_subsection_without_steps_uses_task_kind(self):
        kind = _topic_kind_for_section({"level": 2, "dita_type": "concept"}, "Parameter description")
        self.assertEqual(kind, "task")
        self.assertTrue(_generated_topic_code(10, kind).startswith("DTO"))


class TopicPlaceholderTest(unittest.TestCase):
    def test_placeholder_uses_filename_stem(self):
        self.assertEqual(_topic_placeholder({"filename": "DTC041002.dita"}), "DTC041002")


class ContentGenerationTest(unittest.TestCase):
    def test_chapter_topics_use_concept_shell(self):
        xml = _content_to_dita_xml(
            "测序",
            "第一段\n- 条目",
            "concept",
            "CTO041001",
            "zh-CN",
            topic_kind="chapter",
        )
        self.assertIn("<concept ", xml)
        self.assertIn("<conbody ", xml)
        self.assertIn('cms:imesofttype="chapterTopic"', xml)
        self.assertIn('id="concept-CTO041001"', xml)

    def test_task_topics_use_topic_shell(self):
        xml = _content_to_dita_xml(
            "进行磁珠分选",
            "1. 步骤一",
            "task",
            "DTO041002",
            "zh-CN",
            topic_kind="task",
        )
        self.assertIn("<topic ", xml)
        self.assertIn("<body ", xml)
        self.assertIn('cms:imesofttype="sDitaTopic"', xml)
        self.assertIn('id="topic-DTO041002"', xml)

    def test_concept_topics_use_sdita_topic_shell(self):
        xml = _content_to_dita_xml(
            "产品概述",
            "ZARC 用于基因数据压缩。",
            "concept",
            "DTC041002",
            "zh-CN",
            topic_kind="concept",
        )
        self.assertIn("<topic ", xml)
        self.assertIn("<body ", xml)
        self.assertIn('cms:imesofttype="sDitaTopic"', xml)
        self.assertIn('id="topic-DTC041002"', xml)
        self.assertNotIn("<concept ", xml)
        self.assertNotIn("sDitaConcept", xml)

    def test_cover_topics_use_cover_metadata(self):
        xml = _content_to_dita_xml(
            "封面CN",
            "编号",
            "concept",
            "CTT041003",
            "zh-CN",
            topic_kind="cover",
        )
        self.assertIn('outputclass="cover A4"', xml)
        self.assertIn('cms:imesofttype="sCoverTopic"', xml)

    def test_title_quotes_are_escaped_in_attributes(self):
        xml = _content_to_dita_xml(
            '末端修复产物纯化和加 "A" 尾',
            '正文',
            'concept',
            'DTC041206',
            'zh-CN',
            topic_kind='concept',
        )
        self.assertIn('cms:title="末端修复产物纯化和加 &quot;A&quot; 尾"', xml)

    def test_warning_paragraphs_become_note(self):
        xml = _content_to_dita_xml(
            "连接产物纯化",
            "警告：吸取上清并丢弃。\n请勿触碰磁珠。",
            "concept",
            "DTC041200",
            "zh-CN",
            topic_kind="concept",
        )
        self.assertIn('<note type="warning"><p>吸取上清并丢弃。</p></note>', xml)
        self.assertIn('<note type="warning"><p>触碰磁珠。</p></note>', xml)

    def test_bold_attention_label_becomes_caution_note(self):
        xml = _content_to_dita_xml(
            "细胞计数",
            "[[B]]注意：[[/B]]建议细胞浓度计数时计数活细胞的浓度。",
            "concept",
            "DTC041200",
            "zh-CN",
            topic_kind="concept",
        )
        self.assertIn(
            '<note type="caution"><p>建议细胞浓度计数时计数活细胞的浓度。</p></note>', xml
        )
        self.assertNotIn("<b>注意", xml)

    def test_fully_bold_note_line_keeps_emphasis(self):
        xml = _content_to_dita_xml(
            "样本加载",
            "[[B]]注意：请先进行样本加载，再进行试剂加载[[/B]]",
            "concept",
            "DTC041200",
            "zh-CN",
            topic_kind="concept",
        )
        self.assertIn(
            '<note type="caution"><p><b>请先进行样本加载，再进行试剂加载</b></p></note>', xml
        )

    def test_label_only_note_groups_following_enumerated_items(self):
        xml = _content_to_dita_xml(
            "细胞计数",
            "[[B]]注意：[[/B]]\n\n【1】推荐使用阔口吸头。\n\n【2】计数务必准确。",
            "concept",
            "DTC041200",
            "zh-CN",
            topic_kind="concept",
        )
        self.assertIn('<note type="caution">', xml)
        self.assertIn("<li><p>推荐使用阔口吸头。</p></li>", xml)
        self.assertIn("<li><p>计数务必准确。</p></li>", xml)
        note_start = xml.index('<note type="caution">')
        self.assertLess(xml.index("推荐使用阔口吸头。"), xml.index("</note>", note_start))
        self.assertLess(xml.index("推荐使用阔口吸头。"), xml.index("计数务必准确。"))

    def test_inline_enumerated_note_items_group_into_bullet_list(self):
        xml = _content_to_dita_xml(
            "试剂装载",
            "[[B]]注意：[[/B]]【1】运行8个样本，可直接取1管新试剂。\n\n【2】黑色单圆圈表示用2.0冻存管。\n\n【3】装载完成后盖上盖板。",
            "concept",
            "DTC041201",
            "zh-CN",
            topic_kind="concept",
        )
        self.assertIn('<note type="caution">', xml)
        self.assertIn("<ul>", xml)
        self.assertIn("<li><p>运行8个样本，可直接取1管新试剂。</p></li>", xml)
        self.assertIn("<li><p>黑色单圆圈表示用2.0冻存管。</p></li>", xml)
        self.assertIn("<li><p>装载完成后盖上盖板。</p></li>", xml)
        self.assertEqual(xml.count("<note "), 1)

    def test_english_label_only_note_groups_bracket_items(self):
        xml = _content_to_dita_xml(
            "Cell counting",
            "Note:\n\n[1] Use wide-bore pipette tips.\n\n[2] Count at least three times.",
            "concept",
            "DTC041230",
            "en-US",
            topic_kind="concept",
        )
        self.assertIn('<note type="tip">', xml)
        self.assertIn("<ul>", xml)
        self.assertIn("<li><p>Use wide-bore pipette tips.</p></li>", xml)
        self.assertIn("<li><p>Count at least three times.</p></li>", xml)
        self.assertEqual(xml.count("<note "), 1)
        self.assertNotIn("<p>[2]", xml)

    def test_english_inline_note_bracket_items_group_into_bullet_list(self):
        xml = _content_to_dita_xml(
            "Reagent loading",
            "Note: [1] Retrieve DNA Clean Beads in advance.\n\n[2] A single black circle indicates a 2.0 mL cryovial.\n\n[3] Place the low temperature cover plate.",
            "concept",
            "DTC041231",
            "en-US",
            topic_kind="concept",
        )
        self.assertIn('<note type="tip">', xml)
        self.assertIn("<ul>", xml)
        self.assertIn("<li><p>Retrieve DNA Clean Beads in advance.</p></li>", xml)
        self.assertIn("<li><p>A single black circle indicates a 2.0 mL cryovial.</p></li>", xml)
        self.assertIn("<li><p>Place the low temperature cover plate.</p></li>", xml)
        self.assertEqual(xml.count("<note "), 1)
        self.assertNotIn("<p>[2]", xml)

    def test_table_before_note_stays_outside_note(self):
        xml = _content_to_dita_xml(
            "试剂准备",
            "| 组分 | 体积 |\n| :--- | :--- |\n| Fe (II) Solution | 1 μL |\n\n[[B]]注意：[[/B]]【1】将DNA Clean Beads提前取出。\n\n【2】提前将Lysis Buffer-V3取出。",
            "concept",
            "DTC041226",
            "zh-CN",
            topic_kind="concept",
        )
        self.assertIn('<note type="caution">', xml)
        # The table must be emitted before the note, not nested inside it.
        self.assertLess(xml.index('<table>'), xml.index('<note type="caution">'))
        self.assertNotIn('<table>\n', xml[xml.index('<note'):])
        self.assertIn('<li><p>将DNA Clean Beads提前取出。</p></li>', xml)

    def test_table_heading_becomes_table_title(self):
        xml = _content_to_dita_xml(
            "表 30 Fe (II) 稀释液配置",
            "表 30 Fe (II) 稀释液配置\n| 组分 | 体积 |\n| :--- | :--- |\n| Fe (II) Solution | 1 μL |",
            "task",
            "DTO041201",
            "zh-CN",
            topic_kind="task",
        )
        self.assertIn('<table>\n          <title>Fe (II) 稀释液配置</title>', xml)

    def test_table_has_one_colspec_per_column(self):
        """IME resizes columns via <colspec>; a table without them has fixed widths."""
        xml = _content_to_dita_xml(
            "组分清单",
            "| 组分 | 体积 | 备注 |\n| :--- | :--- | :--- |\n| A | 1 μL | 无 |",
            "concept",
            "DTC041202",
            "zh-CN",
            topic_kind="concept",
        )
        self.assertIn('<tgroup cols="3">', xml)
        self.assertEqual(xml.count("<colspec "), 3)
        self.assertIn('<colspec colname="COLSPEC0" colwidth="1.00*"/>', xml)
        self.assertIn('<colspec colname="COLSPEC2" colwidth="1.00*"/>', xml)
        # colspec must precede the row groups.
        self.assertLess(xml.index("<colspec "), xml.index("<tbody>"))

    def test_colw_marker_sets_relative_colspec_widths(self):
        xml = _content_to_dita_xml(
            "组分清单",
            "[[COLW:0.30*,0.70*]]\n| 组分 | 体积 |\n| :--- | :--- |\n| A | 1 μL |",
            "concept",
            "DTC041202",
            "zh-CN",
            topic_kind="concept",
        )
        self.assertIn('<colspec colname="COLSPEC0" colwidth="0.30*"/>', xml)
        self.assertIn('<colspec colname="COLSPEC1" colwidth="0.70*"/>', xml)
        self.assertNotIn("[[COLW:", xml)

    def test_empty_header_single_cell_callout_defaults_to_tip(self):
        """IDML tip boxes surface as a 1-column table with an empty header; they
        must render as <note type="tip"> instead of a title-less table."""
        xml = _content_to_dita_xml(
            "移除载片",
            "1. 移除载片和试剂槽。\n\n|  |\n| :--- |\n| 请勿将手置于仓门和仪器外壳之间，以防夹手。 |",
            "task",
            "DTO041203",
            "zh-CN",
            topic_kind="task",
        )
        self.assertIn('<note type="tip">', xml)
        self.assertNotIn("<table>", xml)
        # The note stays inside the step so numbering keeps flowing.
        self.assertLess(xml.index("<li>"), xml.index('<note type="tip">'))
        self.assertLess(xml.index('<note type="tip">'), xml.index("</ol>"))

    def test_callout_marker_selects_warning_type(self):
        xml = _content_to_dita_xml(
            "生物危害",
            "| [[NOTE:warning]] |\n| :--- |\n| 请按医疗废弃物标准处理。 |",
            "concept",
            "DTC041205",
            "zh-CN",
            topic_kind="concept",
        )
        self.assertIn('<note type="warning">', xml)
        self.assertNotIn("<table>", xml)

    def test_adjacent_empty_header_callouts_stay_separate_notes(self):
        xml = _content_to_dita_xml(
            "自检",
            "1. 点击自检。\n\n|  |\n| :--- |\n| 请勿将手置于仓门之间。 |\n\n|  |\n| :--- |\n| 测序开始前，切勿手动关闭仓门。 |",
            "task",
            "DTO041206",
            "zh-CN",
            topic_kind="task",
        )
        self.assertEqual(xml.count('<note type="tip">'), 2)
        self.assertNotIn("<table>", xml)

    def test_empty_header_callout_after_steps_is_tip(self):
        xml = _content_to_dita_xml(
            "加样",
            "|  |\n| :--- |\n| 为避免产生气泡，缓慢注入样本。 |",
            "concept",
            "DTC041204",
            "zh-CN",
            topic_kind="concept",
        )
        self.assertIn('<note type="tip">', xml)
        self.assertNotIn("<table>", xml)

    def test_adjacent_tables_keep_their_own_titles(self):
        xml = _content_to_dita_xml(
            "连续表格",
            "表 11 试剂准备\n| 试剂名称 | 要求 |\n| :--- | :--- |\n| TE Buffer | 室温暂存 |\n\n表 12 样本准备\n| 样本名称 | 用途 |\n| :--- | :--- |\n| DNA | 待检测 |",
            "concept",
            "DTC041205",
            "zh-CN",
            topic_kind="concept",
        )
        self.assertIn('<title>试剂准备</title>', xml)
        self.assertIn('<title>样本准备</title>', xml)

    def test_table_number_and_following_line_form_title(self):
        xml = _content_to_dita_xml(
            "Components",
            "Table 2\nMGIEasy Whole Genome Methylation Sequencing Library Prep Kit (16 RXN) (Cat. No.: 940-001530-00)\n| Item | Component |\n| :--- | :--- |\n| A | B |",
            "concept",
            "DTC041212",
            "en-US",
            topic_kind="concept",
        )
        self.assertIn('<title>MGIEasy Whole Genome Methylation Sequencing Library Prep Kit (16 RXN) (Cat. No.: 940-001530-00)</title>', xml)

    def test_table_caption_without_table_body_is_preserved_as_paragraph(self):
        xml = _content_to_dita_xml(
            "Annealing",
            "Table 58 Annealing reaction mix\n1. Place the conditions.",
            "task",
            "DTO041213",
            "en-US",
            topic_kind="task",
        )
        self.assertIn('<p>Table 58 Annealing reaction mix</p>', xml)
        self.assertNotIn('<table>', xml)

    def test_embedded_table_caption_row_splits_into_next_table(self):
        xml = _content_to_dita_xml(
            "Barcode using guide",
            "Table 73 Perfect balanced 8 barcode Pooling strategy (8 barcode from one entire column)\n| Sample 1 | A | G |\n| :--- | :--- | :--- |\n| Signal % | 25.0 | 25.0 |\n| Table 74 Unbalanced 9 barcode Poolin |  | g strategy (barcode from different columns) |\n| Sample 1 | A | T |\n| A signal % | 33.3 | 0 |",
            "concept",
            "DTC041222",
            "en-US",
            topic_kind="concept",
        )
        self.assertEqual(xml.count('<table>'), 2)
        self.assertIn('<title>Perfect balanced 8 barcode Pooling strategy (8 barcode from one entire column)</title>', xml)
        self.assertIn('<title>Unbalanced 9 barcode Pooling strategy (barcode from different columns)</title>', xml)

    def test_markdown_lists_map_to_ol_and_ul(self):
        xml = _content_to_dita_xml(
            "准备",
            "1. 第一步\n2. 第二步\n\n- 提示一\n- 提示二",
            "task",
            "DTO041202",
            "zh-CN",
            topic_kind="task",
        )
        self.assertIn('<ol>', xml)
        self.assertIn('<ul>', xml)

    def test_note_with_unordered_list_stays_inside_note(self):
        xml = _content_to_dita_xml(
            "准备",
            "注意事项：\n- 提示一\n- 提示二",
            "concept",
            "DTC041225",
            "zh-CN",
            topic_kind="concept",
        )
        self.assertIn('<note type="tip">', xml)
        self.assertIn('<ul>', xml)
        self.assertIn('提示一', xml)
        self.assertIn('提示二', xml)
        self.assertIn('<note type="tip">\n        <ul>', xml)

    def test_figure_caption_wraps_pending_image(self):
        xml = _content_to_dita_xml(
            "图片示例",
            "![docx_image_001.png](image/docx_image_001.png)\n图 1 示例图片",
            "concept",
            "DTC041203",
            "zh-CN",
            topic_kind="concept",
        )
        self.assertIn('<fig><title>示例图片</title><image href="image/docx_image_001.png" placement="break"></image></fig>', xml)
        self.assertNotIn('<alt>', xml)

    def test_english_figure_caption_drops_number_prefix(self):
        xml = _content_to_dita_xml(
            "Figure demo",
            "![docx_image_001.png](image/docx_image_001.png)\nFigure 7 PCR BC Primer-96 layout",
            "concept",
            "DTC041217",
            "en-US",
            topic_kind="concept",
        )
        self.assertIn('<fig><title>PCR BC Primer-96 layout</title><image href="image/docx_image_001.png" placement="break"></image></fig>', xml)
        self.assertNotIn('Figure 7', xml)

    def test_figure_caption_written_before_image_binds_to_following_image(self):
        xml = _content_to_dita_xml(
            "封面示例",
            "图1 前置图题\n![docx_image_001.png](image/docx_image_001.png)",
            "concept",
            "DTC041204",
            "zh-CN",
            topic_kind="concept",
        )
        self.assertIn('<fig><title>前置图题</title><image href="image/docx_image_001.png" placement="break"></image></fig>', xml)
        self.assertNotIn('<p>图1 前置图题</p>', xml)

    def test_caption_with_two_following_images_stays_in_one_fig(self):
        xml = _content_to_dita_xml(
            "Consumables Loading",
            "Figure 3-5 Schematic Diagram of Slide Loading\n"
            "![image_008.png](image/image_008.png)\n"
            "![image_009.jpeg](image/image_009.jpeg)",
            "task",
            "DTO041029",
            "en-US",
            topic_kind="task",
        )
        self.assertIn(
            '<fig><title>Schematic Diagram of Slide Loading</title>'
            '<image href="image/image_008.png" placement="break"></image>'
            '<image href="image/image_009.jpeg" placement="break"></image></fig>',
            xml,
        )
        self.assertNotIn('<image href="image/image_009.jpeg" placement="break"></image>\n', xml.split("</fig>")[-1])

    def test_consecutive_figures_keep_separate_images(self):
        xml = _content_to_dita_xml(
            "Sequencing",
            "![image_021.png](image/image_021.png)\n"
            "Figure 4-1 cDNA Library Structure\n"
            "![image_022.png](image/image_022.png)\n"
            "Figure 4-2 Oligo Library Structure\n",
            "concept",
            "CTO041037",
            "en-US",
            topic_kind="concept",
        )
        self.assertIn(
            '<fig><title>cDNA Library Structure</title>'
            '<image href="image/image_021.png" placement="break"></image></fig>',
            xml,
        )
        self.assertIn(
            '<fig><title>Oligo Library Structure</title>'
            '<image href="image/image_022.png" placement="break"></image></fig>',
            xml,
        )

    def test_image_followed_by_unnumbered_caption_becomes_fig(self):
        xml = _content_to_dita_xml(
            "Silicone Pad for Replacing the PCR Cover Plate",
            "1. Press below the latch of the PCR module top cover plate.\n"
            "![image_026.png](image/image_026.png)\n"
            "PCR module top cover plate latch position\n",
            "task",
            "DTO041052",
            "en-US",
            topic_kind="task",
        )
        self.assertIn(
            '<fig><title>PCR module top cover plate latch position</title>'
            '<image href="image/image_026.png" placement="break"></image></fig>',
            xml,
        )
        self.assertNotIn("<p>PCR module top cover plate latch position</p>", xml)

    def test_unmatched_figure_caption_stays_paragraph(self):
        xml = _content_to_dita_xml(
            "无图示例",
            "图1 孤立图题\n\n正文段落",
            "concept",
            "DTC041205",
            "zh-CN",
            topic_kind="concept",
        )
        self.assertIn('<p>图1 孤立图题</p>', xml)

    def test_inline_sup_markup_becomes_sup_tag(self):
        xml = _content_to_dita_xml(
            "About",
            "Agilent[[SUP]]®[[/SUP]] and Thermo Fisher[[SUP]]™[[/SUP]] are trademarks.",
            "concept",
            "DTC041223",
            "en-US",
            topic_kind="concept",
        )
        self.assertIn('<p>Agilent<sup>®</sup> and Thermo Fisher<sup>™</sup> are trademarks.</p>', xml)

    def test_inline_bold_markup_becomes_b_tag(self):
        xml = _content_to_dita_xml(
            "Purity",
            "It is strongly recommended to use [[B]]high quality DNA[[/B]] for library preparation.",
            "concept",
            "DTC041224",
            "en-US",
            topic_kind="concept",
        )
        self.assertIn('<p>It is strongly recommended to use <b>high quality DNA</b> for library preparation.</p>', xml)

    def test_ordered_list_keeps_embedded_blocks_in_same_ol(self):
        xml = _content_to_dita_xml(
            "步骤示例",
            "1. 第一步\n![docx_image_001.png](image/docx_image_001.png)\n图 1 示例图片\n警告：保持低温。\n表 1 参数\n| 名称 | 值 |\n| :--- | :--- |\n| 温度 | 4℃ |\n2. 第二步",
            "task",
            "DTO041210",
            "zh-CN",
            topic_kind="task",
        )
        self.assertEqual(xml.count('<ol>'), 1)
        self.assertIn('<li><p>第一步</p>', xml)
        self.assertIn('<fig><title>示例图片</title><image href="image/docx_image_001.png" placement="break"></image></fig>', xml)
        self.assertIn('<note type="warning"><p>保持低温。</p></note>', xml)
        self.assertIn('<table>\n          <title>参数</title>', xml)
        self.assertIn('<li><p>第二步</p></li>', xml)

    def test_plain_list_items_wrap_text_in_p(self):
        xml = _content_to_dita_xml(
            "适用场景",
            "- 大规模归档\n- 降低存储成本",
            "concept",
            "DTC041240",
            "zh-CN",
            topic_kind="concept",
        )
        self.assertIn("<li><p>大规模归档</p></li>", xml)
        self.assertIn("<li><p>降低存储成本</p></li>", xml)

    def test_table_cell_text_wraps_in_p(self):
        xml = _content_to_dita_xml(
            "参数",
            "| 参数 | 说明 |\n| :--- | :--- |\n| -c | 压缩 |",
            "concept",
            "DTC041241",
            "zh-CN",
            topic_kind="concept",
        )
        self.assertIn("<entry><p>参数</p></entry>", xml)
        self.assertIn("<entry><p>-c</p></entry>", xml)

    def test_standalone_image_has_no_alt_tag(self):
        xml = _content_to_dita_xml(
            "图片示例",
            "![docx_image_003.png](image/docx_image_003.png)",
            "concept",
            "DTC041211",
            "zh-CN",
            topic_kind="concept",
        )
        self.assertIn('<image href="image/docx_image_003.png" placement="break"></image>', xml)
        self.assertNotIn('<alt>', xml)

    def test_note_icon_image_is_skipped_when_adjacent_to_warning_text(self):
        xml = _content_to_dita_xml(
            "连接产物纯化",
            "- 添加试剂、转移上清时请勿触碰、吸取磁珠。\n![docx_image_002.png](image/docx_image_002.png)",
            "concept",
            "DTC041204",
            "zh-CN",
            topic_kind="concept",
        )
        self.assertIn('<note type="warning">', xml)
        self.assertNotIn('image/docx_image_002.png', xml)

    def test_english_warning_paragraph_becomes_note(self):
        xml = _content_to_dita_xml(
            "Cleanup",
            "WARNING: Keep the tube on ice.",
            "concept",
            "DTC041214",
            "en-US",
            topic_kind="concept",
        )
        self.assertIn('<note type="warning"><p>Keep the tube on ice.</p></note>', xml)

    def test_english_note_label_becomes_tip_note(self):
        xml = _content_to_dita_xml(
            "Login",
            "Note: Do not click [Login] repeatedly.",
            "concept",
            "DTC041230",
            "en-US",
            topic_kind="concept",
        )
        self.assertIn('<note type="tip">', xml)
        self.assertIn("<uicontrol>Login</uicontrol>", xml)
        self.assertNotIn("<p>Note:", xml)

    def test_unmatched_bold_close_in_english_note_is_not_uicontrol(self):
        xml = _content_to_dita_xml(
            "Login",
            "Note: [[/B]]Please ensure that the sample carrier is pushed in.",
            "concept",
            "DTC041232",
            "en-US",
            topic_kind="concept",
        )
        self.assertIn('<note type="tip">', xml)
        self.assertNotIn("<uicontrol>/B</uicontrol>", xml)
        self.assertIn("Please ensure that the sample carrier is pushed in.", xml)

    def test_note_label_prefixed_by_bold_space_still_becomes_note(self):
        xml = _content_to_dita_xml(
            "试剂准备",
            "[[B]] [[/B]][[B]]注意：[[/B]]【1】将DNA Clean Beads提前取出。\n\n【2】提前将Lysis Buffer-V3取出。",
            "concept",
            "DTC041225",
            "zh-CN",
            topic_kind="concept",
        )
        self.assertIn('<note type="caution">', xml)
        self.assertIn('<ul>', xml)
        self.assertIn('将DNA Clean Beads提前取出', xml)
        self.assertNotIn('<p><b> </b><b>注意：</b>', xml)

    def test_stoppoint_paragraph_becomes_note(self):
        xml = _content_to_dita_xml(
            "Cleanup",
            "StopPoint: Store at -20C for up to 24 h.",
            "concept",
            "DTC041216",
            "en-US",
            topic_kind="concept",
        )
        self.assertIn('<note type="tip"><p>Store at -20C for up to 24 h.</p></note>', xml)

    def test_multiline_english_note_uses_icon_line_and_following_text(self):
        xml = _content_to_dita_xml(
            "Cleanup",
            "Stop point.\n![docx_image_016.png](/workspace/H-940-001530-00-01 MGIEasy Whole Genome Methylation Sequencing Library Prep Kit User Manual 3.0-2503_media/docx_image_016.png)\nProduct(s) can be stored at -20 C.",
            "concept",
            "DTC041218",
            "en-US",
            topic_kind="concept",
        )
        self.assertIn('<note type="tip">', xml)
        self.assertIn('<p>Product(s) can be stored at -20 C.</p>', xml)
        self.assertNotIn('docx_image_016.png', xml)

    def test_warning_text_with_icon_becomes_warning_note(self):
        xml = _content_to_dita_xml(
            "QC",
            "Do not perform multi-sample pooling with PCR product.\n![docx_image_109.png](/workspace/H-940-001530-00-01 MGIEasy Whole Genome Methylation Sequencing Library Prep Kit User Manual 3.0-2503_media/docx_image_109.png)",
            "concept",
            "DTC041219",
            "en-US",
            topic_kind="concept",
        )
        self.assertIn('<note type="warning"><p>Do not perform multi-sample pooling with PCR product.</p></note>', xml)
        self.assertNotIn('docx_image_109.png', xml)

    def test_stop_point_with_only_page_number_tail_is_dropped(self):
        xml = _content_to_dita_xml(
            "Workflow",
            "- Stop point.\n![docx_image_016.png](/workspace/H-940-001530-00-01 MGIEasy Whole Genome Methylation Sequencing Library Prep Kit User Manual 3.0-2503_media/docx_image_016.png)\n9",
            "concept",
            "DTC041220",
            "en-US",
            topic_kind="concept",
        )
        self.assertNotIn('<note type="tip"><p>Stop point</p></note>', xml)
        self.assertNotIn('<p>9</p>', xml)

    def test_plain_page_number_paragraph_is_dropped(self):
        xml = _content_to_dita_xml(
            "Appendix",
            "正文段落\n\n12",
            "concept",
            "DTC041221",
            "zh-CN",
            topic_kind="concept",
        )
        self.assertIn('<p>正文段落</p>', xml)
        self.assertNotIn('<p>12</p>', xml)

    def test_page_reference_in_sentence_is_removed(self):
        xml = _content_to_dita_xml(
            "Barcode using guide (96 RXN)",
            "1. For sample numbers < 8 with the same data volume, refer to Appendix on page 51 to select barcodes.",
            "concept",
            "DTC041226",
            "en-US",
            topic_kind="concept",
        )
        self.assertIn('refer to Appendix to select barcodes.', xml)
        self.assertNotIn('page 51', xml)

    def test_duplicate_table_header_row_is_removed(self):
        xml = _content_to_dita_xml(
            "Workflow",
            "Table 1 Workflow\n| Section | Workflow | Hands-on time | Total time |\n| :--- | :--- | :--- | :--- |\n| 2.1 | Sample shearing | 2 min | 10 min |\n| Section | Workflow | Hands-on time | Total time |\n| 2.2 | Cleanup | 5 min | 20 min |",
            "concept",
            "DTC041227",
            "en-US",
            topic_kind="concept",
        )
        self.assertEqual(xml.count('<entry><p>Section</p></entry>'), 1)
        self.assertIn('<entry><p>2.2</p></entry>', xml)

    def test_repeated_grouped_first_column_is_collapsed_in_table_body(self):
        xml = _content_to_dita_xml(
            "Components",
            "Table 2 Components\n| Item & Cat. No. | Component | Spec |\n| :--- | :--- | :--- |\n| Module A | Buffer | 10 uL |\n| Module A | Enzyme | 20 uL |",
            "concept",
            "DTC041228",
            "en-US",
            topic_kind="concept",
        )
        self.assertIn('<entry><p>Module A</p></entry>', xml)
        self.assertIn('<entry></entry>\n                <entry><p>Enzyme</p></entry>', xml)

    def test_first_body_row_that_only_repeats_header_prefix_is_removed(self):
        xml = _content_to_dita_xml(
            "Shearing condition",
            """Table 68 Conditions
| S220 | Vessel | microTUBE |
| :--- | :--- | :--- |
| S220 | Vessel | |
|  | Sample Volume | 55 μL |""",
            "concept",
            "DTC041229",
            "en-US",
            topic_kind="concept",
        )
        self.assertEqual(xml.count('<entry><p>S220</p></entry>'), 1)
        self.assertIn('<entry><p>Sample Volume</p></entry>', xml)

    def test_do_not_sentence_stays_paragraph(self):
        xml = _content_to_dita_xml(
            "Cleanup",
            "Do not touch the magnetic beads.",
            "concept",
            "DTC041215",
            "en-US",
            topic_kind="concept",
        )
        self.assertIn('<p>Do not touch the magnetic beads.</p>', xml)
        self.assertNotIn('<note type="warning">', xml)


class DocxHeadingDetectionTest(unittest.TestCase):
    def test_numbered_short_title_is_recognized_as_heading(self):
        self.assertTrue(_looks_like_docx_heading_text('6.2.3 Sample Barcode 使⽤规则 (96 RXN)'))

    def test_parenthesized_short_title_is_recognized_as_heading(self):
        self.assertTrue(_looks_like_docx_heading_text('DNB 加载 (StandardMPS 2.0)'))

    def test_formula_line_is_not_recognized_as_heading(self):
        self.assertFalse(_looks_like_docx_heading_text('Formula 1 Conversion between 1 pmol of PCR product and mass in ng'))
        self.assertFalse(_looks_like_docx_heading_text('Sample mass (ng) = Sample concentration (ng/μL) × Sample volume (μL)'))

    def test_bold_table_caption_is_not_recognized_as_heading(self):
        self.assertFalse(_looks_like_docx_heading_text('[[B]]Table[[/B]][[B]] [[/B]][[B]]10[[/B]][[B]] [[/B]][[B]]Workflow[[/B]]'))

    def test_letter_numbered_table_caption_is_not_recognized_as_heading(self):
        self.assertFalse(_looks_like_docx_heading_text('表B-1  Barcode Primer使用规则'))

    def test_product_line_with_two_parentheticals_is_not_recognized_as_heading(self):
        self.assertFalse(
            _looks_like_docx_heading_text('MGISEQ-2000RS 高通量测序试剂套装(FCL PE100) (Cat. No. 1000012554)')
        )

    def test_bold_label_above_table_is_not_inferred_as_heading(self):
        class StubRun:
            def __init__(self, text, bold):
                self.text = text
                self.bold = bold

        class StubParagraph:
            def __init__(self, runs):
                self.runs = runs

        class StubTable:
            rows = [["试剂名称", "1RXN"]]

        bold_label = StubParagraph([StubRun("常温试剂", True)])
        self.assertEqual(_infer_docx_heading_level(bold_label, "常温试剂", StubTable(), "Normal"), 0)

        phase_label = StubParagraph([StubRun("Cell phase (1.13×)", True)])
        self.assertEqual(
            _infer_docx_heading_level(phase_label, "Cell phase (1.13×)", StubTable(), "Normal"),
            0,
        )

        numbered_headline = StubParagraph([StubRun("3.8.1将试剂转移到制备卡的特定孔位。", None)])
        self.assertEqual(
            _infer_docx_heading_level(numbered_headline, "3.8.1将试剂转移到制备卡的特定孔位。", StubTable(), "3级正文"),
            2,
        )

    def test_numbered_short_list_item_is_not_inferred_as_heading(self):
        from unittest import mock

        paragraph = object()
        with mock.patch("app.api.convert._get_docx_numbering_format", return_value=("decimal", 0)):
            self.assertEqual(
                _infer_docx_heading_level(paragraph, "cDNA产物扩增循环数参考", None, "List Paragraph"),
                0,
            )
            self.assertEqual(
                _infer_docx_heading_level(
                    paragraph, "6.2.3 Sample Barcode 使⽤规则 (96 RXN)", None, "List Paragraph"
                ),
                1,
            )

    def test_short_numbered_list_item_keeps_its_marker(self):
        from unittest import mock

        paragraph = object()
        neighbour = object()
        with mock.patch(
            "app.api.convert._docx_numbering_id",
            side_effect=lambda item: 9 if item in (paragraph, neighbour) else None,
        ), mock.patch("app.api.convert._get_docx_numbering_format", return_value=("decimal", 0)):
            self.assertTrue(_docx_numbered_run_continues(paragraph, neighbour, None))
            self.assertTrue(
                _should_emit_docx_list_marker(paragraph, "cDNA产物扩增循环数参考", neighbour, None)
            )
        with mock.patch("app.api.convert._docx_numbering_id", return_value=None):
            self.assertFalse(_docx_numbered_run_continues(paragraph, neighbour, None))

    def test_clean_title_drops_manual_chapter_section_numbering(self):
        self.assertEqual(_clean_title("第一章 产品信息"), "产品信息")
        self.assertEqual(_clean_title("2.1 注意事项及实验前准备工作"), "注意事项及实验前准备工作")
        self.assertEqual(_clean_title("3.1.1 机器准备"), "机器准备")
        self.assertEqual(_clean_title("4.1 MGISEQ-2000RS测序平台实验流程"), "MGISEQ-2000RS测序平台实验流程")
        self.assertEqual(_clean_title("1. 产品概述"), "产品概述")
        self.assertEqual(_clean_title("1. Product Overview"), "Product Overview")
        self.assertEqual(_clean_title("附录A DNA Clean Beads使用前注意事项"), "DNA Clean Beads使用前注意事项")
        self.assertEqual(_clean_title("附录"), "附录")
        self.assertEqual(_clean_title("4.1.1 仪器、试剂和耗材准备："), "仪器、试剂和耗材准备")

    def test_clean_title_keeps_numbers_that_are_content(self):
        self.assertEqual(_clean_title("3'RNA自动化标准流程"), "3'RNA自动化标准流程")
        self.assertEqual(_clean_title("细胞相（1.13×）"), "细胞相（1.13×）")
        self.assertEqual(_clean_title("2.0 mL"), "2.0 mL")

    def test_reshape_keeps_numbered_docx_heading_as_separate_topic(self):
        markdown = "# Barcode 引物使⽤注意\n\n1. 规则 A\n\n### 6.2.3 Sample Barcode 使⽤规则 (96 RXN)\n\n1. 规则 B\n"
        sections = _postprocess_section_tree(_reshape_docx_sections(_parse_md_sections(markdown)))
        root = sections[0]
        child_titles = [child.get('title') for child in root.get('sections', [])]
        self.assertIn('6.2.3 Sample Barcode 使⽤规则 (96 RXN)', child_titles)

    def test_parenthesized_docx_heading_becomes_separate_section(self):
        markdown = "# MGISEQ-2000RS DNB 加载\n\n### DNB 加载 (StandardMPS)\n\n#### 准备\n\n试剂说明\n"
        sections = _postprocess_section_tree(_reshape_docx_sections(_parse_md_sections(markdown)))
        root = sections[0]
        child_titles = [child.get('title') for child in root.get('sections', [])]
        self.assertIn('DNB 加载 (StandardMPS)', child_titles)

    def test_cover_keeps_version_history_as_child_topic(self):
        sections = [{
            "title": "Cover",
            "content": "DNBelab C-Yellow 16",
            "sections": [{
                "title": "Version History",
                "content": "| Version | Date |\n| --- | --- |\n| V1.0 | March 10, 2026 |\n\nTips: Please download the latest IFU.",
                "sections": [],
            }],
        }]
        processed = _postprocess_section_tree(sections)
        cover = processed[0]
        self.assertEqual([child["title"] for child in cover["sections"]], ["Version History"])
        self.assertIn("DNBelab C-Yellow 16", cover["content"])
        self.assertNotIn("Version History", cover["content"])
        self.assertIn("March 10, 2026", cover["sections"][0]["content"])

    def test_version_history_table_stays_with_version_history_heading(self):
        import tempfile
        from docx import Document

        with tempfile.TemporaryDirectory() as tmp:
            path = str(Path(tmp) / "version_history.docx")
            doc = Document()
            doc.add_paragraph("Version History")
            table = doc.add_table(rows=2, cols=3)
            table.cell(0, 0).text = "Automated Version"
            table.cell(0, 1).text = "Revision Date"
            table.cell(0, 2).text = "Revised Summary"
            table.cell(1, 0).text = "V1.0"
            table.cell(1, 1).text = "March 10, 2026"
            table.cell(1, 2).text = "First release"
            doc.add_paragraph("Tips: Please download the latest version of the Manual.")
            doc.save(path)
            markdown = _docx_to_markdown(path)

        self.assertIn("Version History", markdown)
        self.assertIn("Automated Version", markdown)
        self.assertLess(markdown.index("Version History"), markdown.index("Automated Version"))
        self.assertNotIn("# Cover", markdown)
        sections = _postprocess_section_tree(_reshape_docx_sections(_parse_md_sections(markdown)))
        version_history = next(
            section for section in sections
            if "version history" in (section.get("title") or "").lower()
        )
        self.assertIn("Automated Version", version_history["content"])
        self.assertIn("First release", version_history["content"])

    def test_real_docx_markdown_contains_expected_split_headings(self):
        markdown = _docx_to_markdown('/workspace/.monkeycode-tmp-files/8a916735-H-940-001530-00 MGIEasy 全基因组甲基化建库试剂盒使用说明书 3.0 -250317-1.docx')
        self.assertIn('## Barcode 引物使⽤注意', markdown)

    def test_real_docx_markdown_merges_wrapped_sentence_fragments(self):
        markdown = _docx_to_markdown('/workspace/.monkeycode-tmp-files/8a916735-H-940-001530-00 MGIEasy 全基因组甲基化建库试剂盒使用说明书 3.0 -250317-1.docx')
        self.assertIn('1. 将样本管瞬时离心，再置于磁力架上静置 2~5 min 至液体澄清，小心吸取上清并丢弃。', markdown)
        self.assertNotIn('将样本管瞬时离心，再置于磁力架上静置 2~5 min 至液体澄清，\n\n小心吸取上清并丢弃。', markdown)
        self.assertNotIn('将样本管瞬时离心，再置于磁力架上静置 2~5 min 至液体澄清，\n小心吸取上清并丢弃。', markdown)

    def test_real_docx_markdown_preserves_superscript_placeholders(self):
        markdown = _docx_to_markdown('/workspace/H-940-001530-00-01 MGIEasy Whole Genome Methylation Sequencing Library Prep Kit User Manual 3.0-2503.docx')
        self.assertIn('Agilent[[SUP]]®[[/SUP]]', markdown)

    def test_real_docx_markdown_preserves_bold_placeholders(self):
        markdown = _docx_to_markdown('/workspace/H-940-001530-00-01 MGIEasy Whole Genome Methylation Sequencing Library Prep Kit User Manual 3.0-2503.docx')
        self.assertIn('[[B]]Revision[[/B]][[B]] [[/B]][[B]]history[[/B]]', markdown)

    def test_topic_title_attribute_strips_inline_bold_placeholders(self):
        xml = _content_to_dita_xml(
            '[[B]]Table[[/B]][[B]] [[/B]][[B]]10[[/B]][[B]] [[/B]][[B]]Workflow[[/B]]',
            '',
            'concept',
            'DTC999999',
            'en-US',
            topic_kind='concept',
        )
        self.assertIn('cms:title="Table 10 Workflow"', xml)
        self.assertIn('<title id="title_DTC999999">Table 10 Workflow</title>', xml)

    def test_wrapped_sentence_keeps_independent_explanation_line(self):
        from app.api.convert import _merge_docx_wrapped_paragraphs

        lines = [
            '1. 将样本管从磁力架上取下，加入 32 μL TE Buffer 进行 DNA 洗脱',
            '',
            '也可根据实际需求，适当减少洗脱体积。',
        ]
        self.assertEqual(_merge_docx_wrapped_paragraphs(lines), lines)


class SectionPostprocessTest(unittest.TestCase):
    def test_same_named_bold_table_section_copies_content_to_parent(self):
        sections = _postprocess_section_tree([
            {
                "title": "Workflow",
                "content": "",
                "sections": [
                    {
                        "title": "[[B]]Table[[/B]][[B]] [[/B]][[B]]10[[/B]][[B]] [[/B]][[B]]Workflow[[/B]]",
                        "content": "| Section | Workflow |",
                        "sections": [],
                    }
                ],
            }
        ])
        self.assertEqual(len(sections), 1)
        self.assertEqual(sections[0]["title"], "Workflow")
        self.assertIn('| Section | Workflow |', sections[0]["content"])
        self.assertEqual(len(sections[0]["sections"]), 1)
        self.assertEqual(sections[0]["sections"][0]["title"], '[[B]]Table[[/B]][[B]] [[/B]][[B]]10[[/B]][[B]] [[/B]][[B]]Workflow[[/B]]')

    def test_empty_appendix_wrapper_promotes_children(self):
        sections = _postprocess_section_tree([
            {
                "title": "Appendix",
                "content": "",
                "sections": [
                    {
                        "title": "Precautions Before Using DNA Clean Beads",
                        "content": "Equilibrate beads to room temperature.",
                        "sections": [],
                    },
                    {
                        "title": "Rules for the Use of Barcode Primer",
                        "content": "",
                        "sections": [
                            {
                                "title": "Barcode Primer Use Rules",
                                "content": "Use unique barcodes.",
                                "sections": [],
                            }
                        ],
                    },
                ],
            }
        ])
        self.assertEqual(
            [sec["title"] for sec in sections],
            [
                "Precautions Before Using DNA Clean Beads",
                "Rules for the Use of Barcode Primer",
            ],
        )
        self.assertTrue(all(sec.get("is_appendix") for sec in sections))
        self.assertEqual(
            [child["title"] for child in sections[1]["sections"]],
            ["Barcode Primer Use Rules"],
        )

    def test_empty_chinese_appendix_wrapper_promotes_children(self):
        sections = _postprocess_section_tree([
            {
                "title": "附录",
                "content": "",
                "sections": [
                    {
                        "title": "DNA Clean Beads使用前注意事项",
                        "content": "磁珠需平衡至室温。",
                        "sections": [],
                    }
                ],
            }
        ])
        self.assertEqual([sec["title"] for sec in sections], ["DNA Clean Beads使用前注意事项"])
        self.assertTrue(sections[0].get("is_appendix"))

    def test_appendix_with_body_stays_as_parent(self):
        sections = _postprocess_section_tree([
            {
                "title": "Appendix",
                "content": "See the following items.",
                "sections": [
                    {"title": "Notes", "content": "Keep dry.", "sections": []},
                ],
            }
        ])
        self.assertEqual(len(sections), 1)
        self.assertEqual(sections[0]["title"], "Appendix")
        self.assertFalse(sections[0].get("is_appendix"))
        self.assertEqual([child["title"] for child in sections[0]["sections"]], ["Notes"])


class BookmapTreeTest(unittest.TestCase):
    def test_nested_topics_keep_filenames_and_templates(self):
        lines = []
        node = {
            "topic": {
                "title": "产品信息",
                "filename": "CTO041001.dita",
                "id": "CTO041001",
                "topic_kind": "chapter",
            },
            "children": [
                {
                    "topic": {
                        "title": "实验前准备",
                        "filename": "DTC041002.dita",
                        "id": "DTC041002",
                        "topic_kind": "concept",
                    },
                    "children": [
                        {
                            "topic": {
                                "title": "准备试剂",
                                "filename": "DTO041003.dita",
                                "id": "DTO041003",
                                "topic_kind": "task",
                            },
                            "children": [],
                        }
                    ],
                }
            ],
        }

        _append_root_container_xml(
            lines,
            node,
            iter(range(1, 2)),
            iter(range(1, 10)),
            "  ",
            1,
            "20260811_090000",
            "zh-CN",
        )

        output = "\n".join(lines)
        self.assertIn('href="CTO041001.dita"', output)
        self.assertIn('cms:template="chapterTopic"', output)
        self.assertIn('href="DTC041002.dita"', output)
        self.assertNotIn('sDitaConcept', output)
        self.assertIn('href="DTO041003.dita"', output)
        self.assertEqual(output.count('cms:template="sDitaTopic"'), 2)

    def test_topic_hierarchy_uses_levels_for_nesting(self):
        topics = [
            {"title": "Cover", "filename": "CTT041001.dita", "id": "CTT041001", "topic_kind": "cover", "level": 1},
            {"title": "About the user manual", "filename": "CTO041002.dita", "id": "CTO041002", "topic_kind": "chapter", "level": 2},
            {"title": "Manufacturer information", "filename": "DTC041003.dita", "id": "DTC041003", "topic_kind": "concept", "level": 2},
            {"title": "Contact information", "filename": "DTC041004.dita", "id": "DTC041004", "topic_kind": "concept", "level": 3},
        ]

        roots = _build_topic_hierarchy(topics, set())
        self.assertEqual([node["topic"]["title"] for node in roots], ["Cover"])
        self.assertEqual([node["topic"]["title"] for node in roots[0]["children"]], ["About the user manual", "Manufacturer information"])
        self.assertEqual([node["topic"]["title"] for node in roots[0]["children"][1]["children"]], ["Contact information"])

    def test_frontmatter_topics_after_cover_attach_under_cover(self):
        import re

        lines = []
        topic_roots = [
            {"topic": {"title": "Cover", "filename": "CTT041001.dita", "id": "CTT041001", "topic_kind": "cover"}, "children": []},
            {"topic": {"title": "About the user manual", "filename": "DTC041002.dita", "id": "DTC041002", "topic_kind": "concept"}, "children": []},
            {"topic": {"title": "Manufacturer information", "filename": "DTC041003.dita", "id": "DTC041003", "topic_kind": "concept"}, "children": []},
            {"topic": {"title": "Revision history", "filename": "DTC041004.dita", "id": "DTC041004", "topic_kind": "concept"}, "children": []},
            {"topic": {"title": "Product overview", "filename": "CTO041005.dita", "id": "CTO041005", "topic_kind": "chapter"}, "children": []},
        ]

        _append_bookmap_topics(lines, topic_roots, "20260813_120000", "en-US")
        output = "\n".join(lines)

        self.assertEqual(output.count('navtitle="About the user manual"'), 1)
        self.assertEqual(output.count('navtitle="Manufacturer information"'), 1)
        self.assertEqual(output.count('navtitle="Revision history"'), 1)
        cover_block = re.search(r'<frontmatter[\s\S]*?</frontmatter>', output)
        self.assertIsNotNone(cover_block)
        frontmatter_xml = cover_block.group(0)
        self.assertIn('navtitle="About the user manual"', frontmatter_xml)
        self.assertIn('navtitle="Manufacturer information"', frontmatter_xml)
        self.assertIn('navtitle="Revision history"', frontmatter_xml)

    def test_each_appendix_emits_independent_container_and_chapter_topic(self):
        processed = _postprocess_section_tree([
            {
                "title": "Appendix",
                "content": "",
                "sections": [
                    {
                        "title": "Precautions Before Using DNA Clean Beads",
                        "content": "Equilibrate beads.",
                        "sections": [],
                    },
                    {
                        "title": "Rules for the Use of Barcode Primer",
                        "content": "",
                        "sections": [
                            {
                                "title": "Barcode Primer Use Rules",
                                "content": "Use unique barcodes.",
                                "sections": [],
                            }
                        ],
                    },
                ],
            }
        ])
        flat = _flatten_sections(processed)
        topics = []
        for index, section in enumerate(flat, start=1):
            kind = "chapter" if int(section.get("level", 1) or 1) <= 1 else "concept"
            topics.append({
                "title": section["title"],
                "filename": f"T{index:03d}.dita",
                "id": f"T{index:03d}",
                "topic_kind": kind,
                "level": section["level"],
                "is_appendix": section.get("is_appendix", False),
            })
        roots = _build_topic_hierarchy(topics, set())
        lines = []
        _append_bookmap_topics(lines, roots, "20261020_120000", "en-US")
        output = "\n".join(lines)

        self.assertEqual(output.count("<appendix "), 2)
        self.assertNotIn('navtitle="Appendix"', output)
        self.assertIn('<appendix navtitle="Precautions Before Using DNA Clean Beads"', output)
        self.assertIn('<appendix navtitle="Rules for the Use of Barcode Primer"', output)
        self.assertIn('cms:template="chapterTopic"', output)
        self.assertIn('navtitle="Barcode Primer Use Rules"', output)
        self.assertLess(
            output.index('<appendix navtitle="Rules for the Use of Barcode Primer"'),
            output.index('navtitle="Barcode Primer Use Rules"'),
        )


class TemplateSubsetTest(unittest.TestCase):
    def test_frontmatter_subset_keeps_cover_and_about_but_drops_version_history(self):
        xml = '''<frontmatter navtitle="Preface">
  <topicref navtitle="封面CN" href="cover.dita"/>
  <topicref navtitle="关于说明书" href="about.dita"/>
  <topicref navtitle="版本记录" href="revision.dita"/>
  <booklists><toc/></booklists>
</frontmatter>'''
        subset = _extract_template_frontmatter_subset(xml)
        self.assertIn('href="cover.dita"', subset)
        self.assertIn('href="about.dita"', subset)
        self.assertIn('<booklists><toc/></booklists>', subset)
        # The source document supplies its own version history.
        self.assertNotIn('href="revision.dita"', subset)

    def test_rewrite_frontmatter_injects_version_history_before_booklists(self):
        xml = '<frontmatter><topicref navtitle="封面CN" href="cover.dita" keys="OLD" cms:title="封面CN" cms:placeHolder="cover"/><booklists><toc/></booklists></frontmatter>'
        topics = [
            {"title": "版本历史", "filename": "DTC041004.dita", "id": "DTC041004", "topic_kind": "concept", "level": 2},
        ]
        rewritten, used_files = _rewrite_template_frontmatter(xml, topics)
        self.assertIn('href="cover.dita"', rewritten)
        self.assertIn('href="DTC041004.dita"', rewritten)
        self.assertIn('navtitle="版本历史"', rewritten)
        self.assertLess(rewritten.index('href="DTC041004.dita"'), rewritten.index("<booklists"))
        self.assertEqual(used_files, {"cover.dita", "DTC041004.dita"})

    def test_rewrite_frontmatter_injects_english_version_history(self):
        xml = (
            '<frontmatter>'
            '<topicref navtitle="Cover" href="cover.dita"/>'
            '<topicref navtitle="About the instructions for use" href="about.dita"/>'
            '<booklists><toc/></booklists>'
            '</frontmatter>'
        )
        topics = [
            {"title": "Version History", "filename": "DTC041004.dita", "id": "DTC041004", "topic_kind": "concept", "level": 2},
        ]
        rewritten, used_files = _rewrite_template_frontmatter(xml, topics)
        self.assertIn('navtitle="Version History"', rewritten)
        self.assertLess(rewritten.index('href="DTC041004.dita"'), rewritten.index("<booklists"))
        self.assertIn("cover.dita", used_files)
        self.assertIn("DTC041004.dita", used_files)


class DocxTableMarkdownTest(unittest.TestCase):
    def test_repeated_merged_cell_values_are_collapsed(self):
        tc_objects = {}

        class Cell:
            def __init__(self, text, tc_id):
                self.text = text
                self._tc = tc_objects.setdefault(tc_id, object())

        class Row:
            def __init__(self, cells):
                self.cells = [Cell(text, tc_id) for text, tc_id in cells]

        class Table:
            def __init__(self, rows):
                self.rows = [Row(row) for row in rows]

        table = Table([
            [["Component", 1], ["Volume", 2]],
            [["Buffer", 3], ["10 uL", 4]],
            [["Buffer", 3], ["20 uL", 5]],
        ])

        lines = _docx_table_to_markdown(table)
        self.assertEqual(lines[2], '| Buffer | 10 uL |')
        self.assertEqual(lines[3], '| Buffer | 20 uL |')

    def test_merged_cell_does_not_clear_nonmerged_neighbor_content(self):
        tc_objects = {}

        class Cell:
            def __init__(self, text, tc_id):
                self.text = text
                self._tc = tc_objects.setdefault(tc_id, object())

        class Row:
            def __init__(self, cells):
                self.cells = [Cell(text, tc_id) for text, tc_id in cells]

        class Table:
            def __init__(self, rows):
                self.rows = [Row(row) for row in rows]

        table = Table([
            [["A", 1], ["B", 2], ["C", 3]],
            [["A", 4], ["Merged", 5], ["Tail 1", 6]],
            [["D", 7], ["Merged", 5], ["Tail 2", 8]],
        ])

        lines = _docx_table_to_markdown(table)
        self.assertEqual(lines[2], '| A | Merged | Tail 1 |')
        self.assertEqual(lines[3], '| D | Merged | Tail 2 |')

    def test_identical_text_without_shared_cell_object_is_preserved(self):
        class Cell:
            def __init__(self, text):
                self.text = text
                self._tc = object()

        class Row:
            def __init__(self, cells):
                self.cells = [Cell(text) for text in cells]

        class Table:
            def __init__(self, rows):
                self.rows = [Row(row) for row in rows]

        table = Table([
            ["Component", "Volume"],
            ["Buffer", "10 uL"],
            ["Buffer", "20 uL"],
        ])

        lines = _docx_table_to_markdown(table)
        self.assertEqual(lines[2], '| Buffer | 10 uL |')
        self.assertEqual(lines[3], '| Buffer | 20 uL |')

    def _real_table(self, rows, cols, fill, merges=None, header_rows=0, heading=""):
        from docx import Document
        from docx.oxml import OxmlElement

        doc = Document()
        table = doc.add_table(rows=rows, cols=cols)
        for (r, c), text in fill.items():
            table.cell(r, c).text = text
        for (r1, c1), (r2, c2) in merges or []:
            table.cell(r1, c1).merge(table.cell(r2, c2))
        for row_index in range(header_rows):
            tr_pr = table.rows[row_index]._tr.get_or_add_trPr()
            tr_pr.append(OxmlElement("w:tblHeader"))
        table._convert_heading = heading
        return table

    def test_vmerge_without_tblheader_has_no_thead(self):
        table = self._real_table(
            3,
            3,
            {
                (0, 0): "Kit type",
                (0, 1): "Kit A",
                (0, 2): "Cat. 1",
                (1, 1): "Kit B",
                (1, 2): "Cat. 2",
                (2, 1): "Kit C",
                (2, 2): "Cat. 3",
            },
            merges=[((0, 0), (2, 0))],
        )
        lines = _docx_table_to_markdown(table)
        self.assertEqual(lines[0], "[[THEAD:0]]")
        self.assertIn("[[SPAN:1,3]]", lines[1])
        xml = _content_to_dita_xml(
            "Applicable Reagent Kits",
            "\n".join(lines),
            "concept",
            "DTC041006",
            "en-US",
            topic_kind="concept",
        )
        self.assertNotIn("<thead>", xml)
        self.assertIn('morerows="2"', xml)
        self.assertEqual(xml.count("Kit type"), 1)

    def test_label_first_row_without_merge_becomes_thead(self):
        table = self._real_table(
            2,
            4,
            {
                (0, 0): "Workbench Position",
                (0, 1): "Plate Position Description",
                (0, 2): "Official Name",
                (0, 3): "Displayed Name",
                (1, 0): "POS11",
                (1, 1): "PCR Instrument Module",
                (1, 2): "/",
                (1, 3): "/",
            },
        )
        lines = _docx_table_to_markdown(table)
        self.assertEqual(lines[0], "[[THEAD:1]]")

    def test_all_tblheader_rows_keep_single_thead(self):
        table = self._real_table(
            2,
            5,
            {
                (0, 0): "Target Input Cell Count (cells)",
                (0, 1): "5000",
                (0, 2): "10000",
                (0, 3): "20000",
                (0, 4): "30000",
                (1, 0): "Recommended Concentration (cells/μL)",
                (1, 1): "272 <N<1000",
                (1, 2): "544 <N<2000",
                (1, 3): "1087 <N<2000",
                (1, 4): "1631 <N<2000",
            },
            header_rows=2,
        )
        lines = _docx_table_to_markdown(table)
        self.assertEqual(lines[0], "[[THEAD:1]]")
        xml = _content_to_dita_xml(
            "Total Cell Count and Concentration Requirement",
            "Table 2-1 Recommended Cell Input\n" + "\n".join(lines),
            "concept",
            "DTC041014",
            "en-US",
            topic_kind="concept",
        )
        thead = xml.split("<thead>")[1].split("</thead>")[0]
        tbody = xml.split("<tbody>")[1].split("</tbody>")[0]
        self.assertIn("Target Input Cell Count", thead)
        self.assertNotIn("Recommended Concentration", thead)
        self.assertIn("Recommended Concentration", tbody)

    def test_horizontal_merge_last_row_emits_namest(self):
        table = self._real_table(
            2,
            7,
            {
                (0, 0): "Sample size/lane",
                (0, 1): "Method 1",
                (0, 2): "Method 2",
                (0, 3): "Method 3",
                (0, 4): "Method 4",
                (0, 5): "Method 5",
                (0, 6): "Method 6",
                (1, 0): "8",
                (1, 1): "Choose any 2 out of 4 sets of barcode primers.",
                (1, 5): "17~24",
                (1, 6): "25~32",
            },
            merges=[((1, 1), (1, 4))],
        )
        lines = _docx_table_to_markdown(table)
        self.assertIn("[[TWIDE:180mm]]", lines)
        self.assertIn("[[SPAN:4,1]]", "\n".join(lines))
        xml = _content_to_dita_xml(
            "Barcode Primer Use Rules",
            "\n".join(lines),
            "concept",
            "DTC041047",
            "en-US",
            topic_kind="concept",
        )
        self.assertIn('<table outputclass="180mm">', xml)
        self.assertIn('namest="COLSPEC1" nameend="COLSPEC4"', xml)

    def test_sequencing_parameters_table_is_extra_wide(self):
        xml = _content_to_dita_xml(
            "Sequencing Parameters",
            "| Library Type | cDNA | Oligo |\n| :--- | :--- | :--- |\n| Reads 1 | 47 | 32 |",
            "concept",
            "DTC041040",
            "en-US",
            topic_kind="concept",
        )
        self.assertIn('<table outputclass="180mm">', xml)

    def test_pcr_module_vmerge_keeps_single_description(self):
        table = self._real_table(
            3,
            4,
            {
                (0, 0): "Workbench Position",
                (0, 1): "Plate Position Description",
                (0, 2): "Official Name",
                (0, 3): "Displayed Name",
                (1, 0): "POS11",
                (1, 1): "PCR Instrument Module",
                (1, 2): "/",
                (1, 3): "/",
                (2, 0): "POS12",
                (2, 2): "0.2 mL PCR plate",
                (2, 3): "PCR Plate",
            },
            merges=[((1, 1), (2, 1))],
        )
        xml = _content_to_dita_xml(
            "Consumables Loading",
            "\n".join(_docx_table_to_markdown(table)),
            "task",
            "DTO041029",
            "en-US",
            topic_kind="task",
        )
        self.assertEqual(xml.count("PCR Instrument Module"), 1)
        self.assertIn('morerows="1"', xml)


class DocxMediaCopyTest(unittest.TestCase):
    def test_docx_media_is_copied_to_output_image_dir(self):
        import tempfile
        import zipfile

        with tempfile.TemporaryDirectory() as tmpdir:
            docx_path = Path(tmpdir) / "sample.docx"
            with zipfile.ZipFile(docx_path, "w") as zf:
                zf.writestr("word/media/image1.png", b"abc")
                zf.writestr("word/media/image2.jpg", b"def")

            output_base = Path(tmpdir) / "out"
            output_base.mkdir()

            copied = _copy_docx_media_to_output(str(docx_path), str(output_base))

            self.assertEqual(copied, {"image/docx_media_001.png", "image/docx_media_002.jpg"})
            self.assertTrue((output_base / "image" / "docx_media_001.png").exists())
            self.assertTrue((output_base / "image" / "docx_media_002.jpg").exists())


class ImageDownloadReuseTest(unittest.TestCase):
    def test_download_reuses_existing_output_image(self):
        import tempfile
        from pathlib import Path

        with tempfile.TemporaryDirectory() as tmpdir:
            output_base = Path(tmpdir) / "out"
            image_dir = output_base / "image"
            image_dir.mkdir(parents=True)
            existing = image_dir / "docx_media_001.png"
            existing.write_bytes(b"abc")

            mapping = _download_images_for_output([
                {"alt": "x", "path": str(existing)}
            ], str(output_base))

            self.assertEqual(mapping[str(existing)], "image/docx_media_001.png")

    def test_preserve_names_keeps_original_image_names(self):
        import tempfile
        from pathlib import Path

        with tempfile.TemporaryDirectory() as tmpdir:
            src = Path(tmpdir) / "flow sequencing flow.svg"
            src.write_bytes(b"<svg xmlns='http://www.w3.org/2000/svg'/>")
            output_base = Path(tmpdir) / "out"

            mapping = _download_images_for_output(
                [{"alt": "x", "path": str(src)}],
                str(output_base),
                preserve_names=True,
            )

            self.assertEqual(mapping[str(src)], "image/flow sequencing flow.svg")
            self.assertTrue((output_base / "image" / "flow sequencing flow.svg").exists())


class DitaPackageTitleFilterTest(unittest.TestCase):
    def test_frontmatter_titles_are_filtered(self):
        import tempfile
        import zipfile

        with tempfile.TemporaryDirectory() as tmpdir:
            zip_path = Path(tmpdir) / "sample.zip"
            with zipfile.ZipFile(zip_path, "w") as zf:
                zf.writestr("sample.ditamap", '<topicref navtitle="封面CN" href="a.dita"/><topicref navtitle="正文" href="b.dita"/>')
                zf.writestr("a.dita", '<topic><title>A</title></topic>')
                zf.writestr("b.dita", '<topic><title>B</title></topic>')

            metrics = _parse_dita_package(zip_path)
            self.assertNotIn("封面CN", metrics["title_counts"])


IDML_PKG_NS = "http://ns.adobe.com/AdobeInDesign/idml/1.0/packaging"


class LosslessCsvRoundTripTest(unittest.TestCase):
    def _round_trip(self, content):
        return _csv_to_markdown(_markdown_to_lossless_csv(content))

    def test_trailing_newline_is_preserved(self):
        content = "# Title\n\nbody\n"
        self.assertEqual(self._round_trip(content), content)

    def test_no_trailing_newline_is_preserved(self):
        content = "# Title\n\nbody"
        self.assertEqual(self._round_trip(content), content)

    def test_empty_content(self):
        self.assertEqual(self._round_trip(""), "")

    def test_blank_lines_and_multiple_trailing_newlines(self):
        content = "a\n\nb\n\n\n"
        self.assertEqual(self._round_trip(content), content)

    def test_crlf_line_endings(self):
        content = "a\r\nb\r\n"
        self.assertEqual(self._round_trip(content), content)

    def test_single_line_without_newline(self):
        content = "单行无换行"
        self.assertEqual(self._round_trip(content), content)


def _write_idml_package(path, *, story_files, spread_bodies, styles_body=None):
    """Build a minimal IDML package. story_files: {filename: xml_body}."""
    import zipfile

    designmap_entries = "\n".join(
        f'  <idPkg:Spread src="Spreads/Spread_u1.xml"/>' for _ in [0]
    )
    story_entries = "\n".join(
        f'  <idPkg:Story src="Stories/{name}"/>' for name in story_files
    )
    designmap = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        f'<Document DOMVersion="17.0" xmlns:idPkg="{IDML_PKG_NS}">\n'
        f"{designmap_entries}\n{story_entries}\n</Document>"
    )
    spread = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        f'<idPkg:Spread DOMVersion="17.0" xmlns:idPkg="{IDML_PKG_NS}">\n'
        "  <Spread Self=\"sp1\">\n"
        f"{spread_bodies}\n"
        "  </Spread>\n</idPkg:Spread>"
    )
    if styles_body is None:
        styles_body = (
            '    <ParagraphStyle Self="ParagraphStyle/Heading1" Name="Heading 1"/>\n'
            '    <ParagraphStyle Self="ParagraphStyle/Heading2" Name="Heading 2"/>\n'
            '    <ParagraphStyle Self="ParagraphStyle/Body" Name="Body Text"/>\n'
            '    <ParagraphStyle Self="ParagraphStyle/StepsList" Name="Steps List">\n'
            "      <Properties><BulletsAndNumberingListType>NumberedList</BulletsAndNumberingListType></Properties>\n"
            "    </ParagraphStyle>\n"
        )
    styles = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        f'<idPkg:Styles DOMVersion="17.0" xmlns:idPkg="{IDML_PKG_NS}">\n'
        '  <RootParagraphStyleGroup Self="root">\n'
        f"{styles_body}"
        "  </RootParagraphStyleGroup>\n</idPkg:Styles>"
    )
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("mimetype", "application/vnd.adobe.indesign-idml-package")
        zf.writestr("designmap.xml", designmap)
        zf.writestr("Spreads/Spread_u1.xml", spread)
        for name, body in story_files.items():
            zf.writestr(f"Stories/{name}", body)
        zf.writestr("Resources/Styles.xml", styles)


def _idml_story(body, self_id="s"):
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        f'<idPkg:Story DOMVersion="17.0" xmlns:idPkg="{IDML_PKG_NS}">\n'
        f'  <Story Self="{self_id}">{body}</Story>\n</idPkg:Story>'
    )


class IdmlConversionTest(unittest.TestCase):
    def test_paragraph_styles_map_to_heading_levels(self):
        import tempfile

        body = (
            '<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/Heading1">'
            '<CharacterStyleRange><Content>产品信息</Content></CharacterStyleRange>'
            "</ParagraphStyleRange>"
            '<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/Body">'
            '<CharacterStyleRange><Content>本节介绍产品。</Content></CharacterStyleRange>'
            "</ParagraphStyleRange>"
        )
        with tempfile.TemporaryDirectory() as tmpdir:
            idml_path = Path(tmpdir) / "manual.idml"
            _write_idml_package(
                idml_path,
                story_files={"Story_u1.xml": _idml_story(body)},
                spread_bodies='    <TextFrame Self="tf1" ParentStory="Story/u1" ItemTransform="1 0 0 1 20 60"/>',
            )
            md = _idml_to_markdown(str(idml_path))
            self.assertIn("# 产品信息", md)
            self.assertIn("本节介绍产品。", md)
            sections = _flatten_sections(
                _postprocess_section_tree(_reshape_docx_sections(_parse_md_sections(md)))
            )
            titles = [s["title"] for s in sections]
            self.assertIn("产品信息", titles)

    def test_reading_order_follows_story_ids_not_filenames(self):
        import tempfile

        with tempfile.TemporaryDirectory() as tmpdir:
            idml_path = Path(tmpdir) / "order.idml"
            # Self ids differ from filenames: filename sort (aaa before zzz)
            # would yield AAA first, but the spread references sid1 (ZZZ) first.
            files = {
                "Story_aaa.xml": _idml_story(
                    self_id="sid2",
                    body='<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/Heading1">'
                    '<CharacterStyleRange><Content>AAA</Content></CharacterStyleRange>'
                    "</ParagraphStyleRange>",
                ),
                "Story_zzz.xml": _idml_story(
                    self_id="sid1",
                    body='<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/Heading1">'
                    '<CharacterStyleRange><Content>ZZZ</Content></CharacterStyleRange>'
                    "</ParagraphStyleRange>",
                ),
            }
            _write_idml_package(
                idml_path,
                story_files=files,
                spread_bodies=(
                    '    <TextFrame Self="tf2" ParentStory="sid1"/>\n'
                    '    <TextFrame Self="tf1" ParentStory="sid2"/>'
                ),
            )
            md = _idml_to_markdown(str(idml_path))
            self.assertLess(md.index("ZZZ"), md.index("AAA"))

    def test_container_zip_prefers_idml_over_indd(self):
        import tempfile
        import zipfile

        with tempfile.TemporaryDirectory() as tmpdir:
            inner = Path(tmpdir) / "manual.idml"
            _write_idml_package(
                inner,
                story_files={
                    "Story_u1.xml": _idml_story(
                        '<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/Heading1">'
                        '<CharacterStyleRange><Content>内容</Content></CharacterStyleRange>'
                        "</ParagraphStyleRange>"
                    )
                },
                spread_bodies='    <TextFrame Self="tf1" ParentStory="Story/u1" ItemTransform="1 0 0 1 20 60"/>',
            )
            container = Path(tmpdir) / "package.zip"
            with zipfile.ZipFile(container, "w") as zf:
                zf.writestr("manual.idml", inner.read_bytes())
                zf.writestr("manual.indd", b"\x00binary-indd\x00")
            resolved = _resolve_idml_source(str(container))
            self.assertIsNotNone(resolved)
            self.assertTrue(resolved.lower().endswith(".idml"))

    def test_container_zip_without_idml_returns_none(self):
        import tempfile
        import zipfile

        with tempfile.TemporaryDirectory() as tmpdir:
            container = Path(tmpdir) / "package.zip"
            with zipfile.ZipFile(container, "w") as zf:
                zf.writestr("manual.indd", b"\x00binary-indd\x00")
            self.assertIsNone(_resolve_idml_source(str(container)))

    def _build_idml_with_external_link(self, tmpdir, prefix=""):
        """Container holds only an .idml plus a sibling Links/ image folder."""
        import zipfile

        body = (
            '<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/Heading1">'
            '<CharacterStyleRange><Content>封面标题</Content></CharacterStyleRange>'
            "</ParagraphStyleRange>"
        )
        inner = Path(tmpdir) / "manual.idml"
        _write_idml_package(
            inner,
            story_files={"Story_u1.xml": _idml_story(body)},
            spread_bodies=(
                '<Rectangle Self="r1"><Image Self="i1">'
                '<Link Self="l1" LinkResourceURI="file:Links/image1.png"/></Image></Rectangle>'
                '<TextFrame Self="tf1" ParentStory="s"/>'
            ),
        )
        container = Path(tmpdir) / "user_upload.zip"
        with zipfile.ZipFile(container, "w") as zf:
            zf.writestr(f"{prefix}manual.idml", inner.read_bytes())
            zf.writestr(f"{prefix}Links/image1.png", b"\x89PNG\r\n\x1a\nREALIMAGE")
        return str(container)

    def test_container_merges_sibling_links_folder(self):
        import tempfile

        with tempfile.TemporaryDirectory() as tmpdir:
            container = self._build_idml_with_external_link(tmpdir)
            resolved = _resolve_idml_source(container)
            self.assertIsNotNone(resolved)
            md = _idml_to_markdown(resolved)
            self.assertIn("image1.png", md)
            media_dir = Path(resolved).with_name(Path(resolved).stem + "_idml_media")
            self.assertTrue((media_dir / "image1.png").exists())

    def test_container_merges_links_folder_under_prefix_dir(self):
        import tempfile

        with tempfile.TemporaryDirectory() as tmpdir:
            container = self._build_idml_with_external_link(tmpdir, prefix="package/")
            resolved = _resolve_idml_source(container)
            md = _idml_to_markdown(resolved)
            self.assertIn("image1.png", md)

    def test_container_skips_unreferenced_resources(self):
        import tempfile
        import zipfile

        with tempfile.TemporaryDirectory() as tmpdir:
            body = (
                '<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/Heading1">'
                '<CharacterStyleRange><Content>正文</Content></CharacterStyleRange>'
                "</ParagraphStyleRange>"
            )
            inner = Path(tmpdir) / "manual.idml"
            _write_idml_package(
                inner,
                story_files={"Story_u1.xml": _idml_story(body)},
                spread_bodies=(
                    '<Rectangle Self="r1"><Image Self="i1">'
                    '<Link Self="l1" LinkResourceURI="file:Links/used.png"/></Image></Rectangle>'
                    '<TextFrame Self="tf1" ParentStory="s"/>'
                ),
            )
            container = Path(tmpdir) / "user_upload.zip"
            with zipfile.ZipFile(container, "w") as zf:
                zf.writestr("manual.idml", inner.read_bytes())
                zf.writestr("Links/used.png", b"\x89PNG\r\n\x1a\nUSED")
                zf.writestr("Links/unused.psd", b"\x00" + b"x" * (2 * 1024 * 1024))
            resolved = _resolve_idml_source(str(container))
            names = zipfile.ZipFile(resolved).namelist()
            self.assertTrue(any(n.endswith("used.png") for n in names))
            self.assertFalse(any(n.endswith("unused.psd") for n in names))

    def test_deferred_image_formats_are_skipped(self):
        """Formats needing AI->SVG first (e.g. .ai) are not carried or emitted."""
        import tempfile
        import zipfile

        with tempfile.TemporaryDirectory() as tmpdir:
            body = (
                '<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/Heading1">'
                '<CharacterStyleRange><Content>正文</Content></CharacterStyleRange>'
                "</ParagraphStyleRange>"
            )
            inner = Path(tmpdir) / "manual.idml"
            _write_idml_package(
                inner,
                story_files={"Story_u1.xml": _idml_story(body)},
                spread_bodies=(
                    '<Rectangle Self="r1"><Image Self="i1">'
                    '<Link Self="l1" LinkResourceURI="file:Links/diagram.ai"/></Image></Rectangle>'
                    '<Rectangle Self="r2"><Image Self="i2">'
                    '<Link Self="l2" LinkResourceURI="file:Links/shot.png"/></Image></Rectangle>'
                    '<TextFrame Self="tf1" ParentStory="s"/>'
                ),
            )
            container = Path(tmpdir) / "user_upload.zip"
            with zipfile.ZipFile(container, "w") as zf:
                zf.writestr("manual.idml", inner.read_bytes())
                zf.writestr("Links/diagram.ai", b"%PDF fake-ai")
                zf.writestr("Links/shot.png", b"\x89PNG\r\n\x1a\nPNG")
            resolved = _resolve_idml_source(str(container))
            names = zipfile.ZipFile(resolved).namelist()
            self.assertFalse(any(n.endswith(".ai") for n in names))
            md = _idml_to_markdown(resolved)
            self.assertNotIn("diagram.ai", md)
            self.assertIn("shot.png", md)

    def _idml_referencing_ai(self, tmpdir, extra_links):
        import zipfile

        body = (
            '<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/Heading1">'
            '<CharacterStyleRange><Content>正文</Content></CharacterStyleRange>'
            "</ParagraphStyleRange>"
        )
        inner = Path(tmpdir) / "manual.idml"
        _write_idml_package(
            inner,
            story_files={"Story_u1.xml": _idml_story(body)},
            spread_bodies=(
                '<Rectangle Self="r1"><Image Self="i1">'
                '<Link Self="l1" LinkResourceURI="file:Links/logo.ai"/></Image></Rectangle>'
                '<TextFrame Self="tf1" ParentStory="s"/>'
            ),
        )
        container = Path(tmpdir) / "user_upload.zip"
        with zipfile.ZipFile(container, "w") as zf:
            zf.writestr("manual.idml", inner.read_bytes())
            for name, data in extra_links.items():
                zf.writestr(name, data)
        return str(container)

    def test_ai_reference_uses_converted_svg_sibling(self):
        """A .ai link is served by the same-stem .svg the user converted."""
        import tempfile
        import zipfile

        with tempfile.TemporaryDirectory() as tmpdir:
            container = self._idml_referencing_ai(
                tmpdir, {"Links/logo.svg": b"<svg xmlns='http://www.w3.org/2000/svg'/>"}
            )
            resolved = _resolve_idml_source(container)
            names = zipfile.ZipFile(resolved).namelist()
            self.assertTrue(any(n.endswith("logo.svg") for n in names))
            md = _idml_to_markdown(resolved)
            self.assertIn("logo.svg", md)

    def test_percent_encoded_link_resolves_sibling_with_spaces(self):
        """InDesign percent-encodes spaces in link URIs; the sibling must still match."""
        import tempfile
        import zipfile

        with tempfile.TemporaryDirectory() as tmpdir:
            body = (
                '<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/Heading1">'
                '<CharacterStyleRange><Content>正文</Content></CharacterStyleRange>'
                "</ParagraphStyleRange>"
            )
            inner = Path(tmpdir) / "manual.idml"
            _write_idml_package(
                inner,
                story_files={"Story_u1.xml": _idml_story(body)},
                spread_bodies=(
                    '<Rectangle Self="r1"><Image Self="i1">'
                    '<Link Self="l1" LinkResourceURI="file:Links/flow%20workflow.ai"/>'
                    "</Image></Rectangle>"
                    '<TextFrame Self="tf1" ParentStory="s"/>'
                ),
            )
            container = Path(tmpdir) / "user_upload.zip"
            with zipfile.ZipFile(container, "w") as zf:
                zf.writestr("manual.idml", inner.read_bytes())
                zf.writestr(
                    "Links/flow workflow.svg",
                    b"<svg xmlns='http://www.w3.org/2000/svg'/>",
                )
            resolved = _resolve_idml_source(container)
            names = zipfile.ZipFile(resolved).namelist()
            self.assertTrue(any(n.endswith("flow workflow.svg") for n in names))
            md = _idml_to_markdown(resolved)
            # The image package's own name (space included) is preserved.
            self.assertIn("flow workflow.svg", md)

    def test_ai_reference_prefers_svg_over_ai(self):
        import tempfile
        import zipfile

        with tempfile.TemporaryDirectory() as tmpdir:
            container = self._idml_referencing_ai(
                tmpdir,
                {
                    "Links/logo.ai": b"%PDF fake-ai",
                    "Links/logo.svg": b"<svg xmlns='http://www.w3.org/2000/svg'/>",
                },
            )
            resolved = _resolve_idml_source(container)
            names = zipfile.ZipFile(resolved).namelist()
            self.assertFalse(any(n.endswith(".ai") for n in names))
            md = _idml_to_markdown(resolved)
            self.assertIn("logo.svg", md)

    def test_ai_reference_without_svg_is_skipped(self):
        import tempfile

        with tempfile.TemporaryDirectory() as tmpdir:
            container = self._idml_referencing_ai(tmpdir, {"Links/logo.ai": b"%PDF fake-ai"})
            resolved = _resolve_idml_source(container)
            md = _idml_to_markdown(resolved)
            self.assertNotIn(".ai", md)
            self.assertNotIn("logo", md)

    def test_leading_image_forms_cover(self):
        import tempfile
        import zipfile

        with tempfile.TemporaryDirectory() as tmpdir:
            idml_path = Path(tmpdir) / "manual.idml"
            body = (
                '<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/Heading1">'
                '<CharacterStyleRange><Content>正文</Content></CharacterStyleRange>'
                "</ParagraphStyleRange>"
            )
            _write_idml_package(
                idml_path,
                story_files={"Story_u1.xml": _idml_story(body)},
                spread_bodies=(
                    '    <Rectangle Self="rect1" ItemTransform="1 0 0 1 20 10">'
                    '<Image Self="img1"><Link Self="lnk1" LinkResourceURI="file:Links/cover.png"/></Image>'
                    "</Rectangle>\n"
                    '    <TextFrame Self="tf1" ParentStory="s" ItemTransform="1 0 0 1 20 60"/>'
                ),
            )
            with zipfile.ZipFile(idml_path, "a") as zf:
                zf.writestr("Links/cover.png", b"\x89PNG\r\n\x1a\nfake")
            md = _idml_to_markdown(str(idml_path))
            # Cover must be a single leading section, not inserted mid-document.
            self.assertTrue(md.startswith("# Cover"))
            self.assertIn("cover.png", md)
            self.assertLess(md.index("cover.png"), md.index("正文"))
            self.assertEqual(md.count("# Cover"), 1)

    def test_real_world_title_l_styles_map_to_heading_levels(self):
        """MGI/DNBSEQ manuals name headings "<group>:Title-L<n>-<lang>"."""
        import tempfile

        styles = (
            '    <ParagraphStyle Self="ParagraphStyle/For reagent kit only:Title-L1-with number-CN"'
            ' Name="For reagent kit only:Title-L1-with number-CN"/>\n'
            '    <ParagraphStyle Self="ParagraphStyle/For reagent kit only:Title-L2-with number-CN"'
            ' Name="For reagent kit only:Title-L2-with number-CN"/>\n'
            '    <ParagraphStyle Self="ParagraphStyle/Text:L2-CN" Name="Text:L2-CN"/>\n'
        )
        body = (
            '<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/For reagent kit only:Title-L1-with number-CN">'
            "<CharacterStyleRange><Content>介绍</Content></CharacterStyleRange></ParagraphStyleRange>"
            '<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/For reagent kit only:Title-L2-with number-CN">'
            "<CharacterStyleRange><Content>预期用途</Content></CharacterStyleRange></ParagraphStyleRange>"
            '<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/Text:L2-CN">'
            "<CharacterStyleRange><Content>正文内容。</Content></CharacterStyleRange></ParagraphStyleRange>"
        )
        with tempfile.TemporaryDirectory() as tmpdir:
            idml_path = Path(tmpdir) / "manual.idml"
            _write_idml_package(
                idml_path,
                story_files={"Story_u1.xml": _idml_story(body)},
                spread_bodies='    <TextFrame Self="tf1" ParentStory="s" ItemTransform="1 0 0 1 20 60"/>',
                styles_body=styles,
            )
            md = _idml_to_markdown(str(idml_path))
            self.assertIn("# 介绍", md)
            self.assertIn("## 预期用途", md)
            # The heading style contains the word "number"; it must not be
            # mistaken for an ordered-list item.
            self.assertNotIn("1. 介绍", md)
            self.assertNotIn("1. 预期用途", md)

    def test_table_of_contents_styles_are_dropped(self):
        import tempfile

        styles = (
            '    <ParagraphStyle Self="ParagraphStyle/Contents :Level 1-CN" Name="Contents :Level 1-CN"/>\n'
            '    <ParagraphStyle Self="ParagraphStyle/Heading1" Name="Heading 1"/>\n'
        )
        body = (
            '<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/Heading1">'
            "<CharacterStyleRange><Content>介绍</Content></CharacterStyleRange></ParagraphStyleRange>"
            '<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/Contents :Level 1-CN">'
            "<CharacterStyleRange><Content>第1章 介绍1</Content></CharacterStyleRange></ParagraphStyleRange>"
        )
        with tempfile.TemporaryDirectory() as tmpdir:
            idml_path = Path(tmpdir) / "manual.idml"
            _write_idml_package(
                idml_path,
                story_files={"Story_u1.xml": _idml_story(body)},
                spread_bodies='    <TextFrame Self="tf1" ParentStory="s" ItemTransform="1 0 0 1 20 60"/>',
                styles_body=styles,
            )
            md = _idml_to_markdown(str(idml_path))
            self.assertIn("# 介绍", md)
            self.assertNotIn("第1章", md)

    def _callout_table(self, icon_name, body_text):
        """Build an IDML callout box: single column, icon-only header cell."""
        link = (
            '<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/Text:Symbols-CN">'
            f'<CharacterStyleRange><Link LinkResourceURI="file:C:/doc/Links/{icon_name}" />'
            "</CharacterStyleRange></ParagraphStyleRange>"
        )
        return (
            '<Table Self="t1">'
            '<Row Self="r0" Name="0"/><Row Self="r1" Name="1"/>'
            '<Column Self="co0" Name="0"/>'
            f'<Cell Self="c0" Name="0:0" RowSpan="1" ColumnSpan="1">{link}</Cell>'
            '<Cell Self="c1" Name="1:0" RowSpan="1" ColumnSpan="1">'
            '<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/Table:Warning sentence-CN">'
            f"<CharacterStyleRange><Content>{body_text}</Content></CharacterStyleRange>"
            "</ParagraphStyleRange></Cell>"
            "</Table>"
        )

    def test_callout_icon_selects_note_type(self):
        cases = {
            "label000081_Warning.ai": "warning",
            "label000018_Danger.ai": "warning",
            "label000008_Caution.ai": "caution",
            "label000073_Tips.ai": "tip",
            "icon_noteEN.ai": "tip",
        }
        for icon_name, expected in cases.items():
            with self.subTest(icon=icon_name):
                table = self._callout_table(icon_name, "示例内容")
                lines = _idml_table_to_markdown(
                    __import__("xml.etree.ElementTree", fromlist=["ET"]).fromstring(table)
                )
                self.assertEqual(lines[0], f"| [[NOTE:{expected}]] |")

    def test_callout_marker_selects_caution_type(self):
        xml = _content_to_dita_xml(
            "注意",
            "| [[NOTE:caution]] |\n| :--- |\n| 握住载片时避免触碰小孔。 |",
            "concept",
            "DTC041207",
            "zh-CN",
            topic_kind="concept",
        )
        self.assertIn('<note type="caution">', xml)
        self.assertIn("握住载片时避免触碰小孔。", xml)

    def test_two_column_icon_callout_becomes_note(self):
        """A real IDML callout is a 2-column table: an icon column plus a body
        column. It must surface as a <note>, typed from the icon, not a table."""
        import xml.etree.ElementTree as ET

        table = (
            '<Table Self="t1" HeaderRowCount="0" BodyRowCount="1" ColumnCount="2">'
            '<Row Self="r0" Name="0"/>'
            '<Column Self="co0" Name="0"/><Column Self="co1" Name="1"/>'
            '<Cell Self="c0" Name="0:0" RowSpan="1" ColumnSpan="1">'
            '<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/Text:Symbols-CN">'
            '<CharacterStyleRange><Link LinkResourceURI="file:C:/doc/Links/label000081_Warning.ai"/></CharacterStyleRange>'
            "</ParagraphStyleRange></Cell>"
            '<Cell Self="c1" Name="1:0" RowSpan="1" ColumnSpan="1">'
            '<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/Table:Warning sentence-CN">'
            "<CharacterStyleRange><Content>请勿将手置于仓门之间。</Content></CharacterStyleRange>"
            "</ParagraphStyleRange></Cell>"
            "</Table>"
        )
        lines = _idml_table_to_markdown(ET.fromstring(table))
        self.assertEqual(lines[0], "| [[NOTE:warning]] |")
        self.assertIn("请勿将手置于仓门之间。", "\n".join(lines))

    def test_numbered_list_br_runs_become_separate_steps(self):
        """<Br/> inside one InDesign numbered-list paragraph separates steps."""
        import tempfile

        body = (
            '<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/StepsList">'
            "<CharacterStyleRange><Content>第一步。</Content><Br/>"
            "<Content>第二步。</Content><Br/>"
            "<Content>第三步。</Content></CharacterStyleRange>"
            "</ParagraphStyleRange>"
        )
        with tempfile.TemporaryDirectory() as tmpdir:
            idml_path = Path(tmpdir) / "manual.idml"
            _write_idml_package(
                idml_path,
                story_files={"Story_u1.xml": _idml_story(body)},
                spread_bodies='    <TextFrame Self="tf1" ParentStory="s" ItemTransform="1 0 0 1 20 60"/>',
            )
            md = _idml_to_markdown(str(idml_path))
        self.assertIn("1. 第一步。", md)
        self.assertIn("1. 第二步。", md)
        self.assertIn("1. 第三步。", md)
        xml = _content_to_dita_xml("步骤", md, "task", "DTO041212", "zh-CN", topic_kind="task")
        self.assertEqual(xml.count("<li>"), 3)

    def test_trailing_step_icon_stays_inside_last_list_item(self):
        """An inline button icon on the final step must not escape after </ol>."""
        md = (
            "操作步骤如下：\n\n"
            "1. 第一步。\n\n"
            "1. 第二步，点击按钮。\n\n"
            "![icon_nextstep.svg](image/icon_nextstep.svg)\n"
        )
        xml = _content_to_dita_xml("装载", md, "task", "DTO041214", "zh-CN", topic_kind="task")
        self.assertIn("icon_nextstep.svg", xml)
        self.assertLess(xml.index("icon_nextstep.svg"), xml.rindex("</li>"))
        self.assertLess(xml.index("icon_nextstep.svg"), xml.index("</ol>"))

    def test_inline_icon_before_note_keeps_its_position(self):
        """An inline icon queued before a callout stays ahead of the note."""
        md = (
            "如在方案选择界面点击，界面出现提示信息。\n\n"
            "![icon_laststep.svg](image/icon_laststep.svg)\n\n"
            "| [[NOTE:warning]] |\n| :--- |\n| 点击确定后将退出测序。 |\n"
        )
        xml = _content_to_dita_xml("回顾参数", md, "task", "DTO041215", "zh-CN", topic_kind="task")
        self.assertIn('<note type="warning">', xml)
        self.assertLess(xml.index("icon_laststep.svg"), xml.index('<note type="warning">'))

    def test_bracket_control_names_become_uicontrol(self):
        """Chinese 【】 control names map to <uicontrol> without the brackets."""
        md = "1. 在【测序类型】右侧选择测序类型，如【PE150APPC】。\n\n1. 点击【确定】。\n"
        xml = _content_to_dita_xml("选择测序方案", md, "task", "DTO041216", "zh-CN", topic_kind="task")
        self.assertIn("<uicontrol>测序类型</uicontrol>", xml)
        self.assertIn("<uicontrol>PE150APPC</uicontrol>", xml)
        self.assertIn("<uicontrol>确定</uicontrol>", xml)
        self.assertNotIn("【", xml)

    def test_numeric_bracket_marker_is_not_a_uicontrol(self):
        """【1】 is a note enumeration marker, not a UI control name."""
        md = "注意：\n\n【1】提前取出DNA Clean Beads。\n\n【2】提前取出Lysis Buffer。\n"
        xml = _content_to_dita_xml("试剂准备", md, "concept", "DTC041217", "zh-CN", topic_kind="concept")
        self.assertIn("<li><p>提前取出DNA Clean Beads。</p></li>", xml)
        self.assertNotIn("<uicontrol>1</uicontrol>", xml)

    def test_english_step_bold_becomes_uicontrol(self):
        """Bold in an English operation step names a UI control."""
        md = "1. Click [[B]]OK[[/B]] to continue.\n"
        xml = _content_to_dita_xml("Run", md, "task", "DTO041218", "en-US", topic_kind="task")
        self.assertIn("<uicontrol>OK</uicontrol>", xml)
        self.assertNotIn("<b>OK</b>", xml)

    def test_english_prose_bold_stays_bold(self):
        """Bold outside a step (prose) is not a UI control in English."""
        md = "This section is [[B]]important[[/B]].\n"
        xml = _content_to_dita_xml("Notes", md, "concept", "DTC041219", "en-US", topic_kind="concept")
        self.assertIn("<b>important</b>", xml)

    def test_english_square_bracket_label_becomes_uicontrol(self):
        xml = _content_to_dita_xml(
            "Software login",
            "Double-click to open [DNBelabC-YellowR16]. Password is [123456]. See note [1].",
            "task",
            "DTO041231",
            "en-US",
            topic_kind="task",
        )
        self.assertIn("<uicontrol>DNBelabC-YellowR16</uicontrol>", xml)
        self.assertIn("<uicontrol>123456</uicontrol>", xml)
        self.assertNotIn("<uicontrol>1</uicontrol>", xml)

    def test_idml_bullet_with_warning_keyword_stays_a_list_item(self):
        """IDML warnings are callout tables; a bulleted item must not be
        promoted to a note and folded into the previous item."""
        md = (
            "- 本产品仅用于科学研究，使用前请仔细阅读产品说明书。\n\n"
            "- 实验前请熟悉和掌握需使用的各种仪器的操作方法和注意事项。\n\n"
            "- 所有样本及试剂应避免直接接触皮肤和眼睛，切勿吞咽。\n\n"
            "- 所有样本和各种废弃物均应按相关法规规定进行污染物处理。\n"
        )
        xml = _content_to_dita_xml(
            "注意事项", md, "concept", "DTC041220", "zh-CN",
            topic_kind="concept", source_format="idml",
        )
        self.assertEqual(xml.count("<li>"), 4)
        self.assertNotIn("<note", xml)

    def test_idml_silk_screen_run_becomes_uicontrol(self):
        """English on-screen text uses a "Silk screen" character style (which
        renders bold); it maps to <uicontrol> like the Chinese 【】 convention."""
        import tempfile

        styles = (
            '    <ParagraphStyle Self="ParagraphStyle/StepsList" Name="Numbered list:Level 2-EN">\n'
            "      <Properties><BulletsAndNumberingListType>NumberedList</BulletsAndNumberingListType></Properties>\n"
            "    </ParagraphStyle>\n"
            '    <CharacterStyle Self="CharacterStyle/Silk screen-EN" Name="Silk screen-EN"/>\n'
        )
        body = (
            '<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/StepsList">'
            '<CharacterStyleRange><Content>Tap </Content></CharacterStyleRange>'
            '<CharacterStyleRange AppliedCharacterStyle="CharacterStyle/Silk screen-EN">'
            "<Content>Login</Content></CharacterStyleRange>"
            '<CharacterStyleRange><Content> to continue.</Content></CharacterStyleRange>'
            "</ParagraphStyleRange>"
        )
        with tempfile.TemporaryDirectory() as tmpdir:
            idml_path = Path(tmpdir) / "manual_en.idml"
            _write_idml_package(
                idml_path,
                story_files={"Story_u1.xml": _idml_story(body)},
                spread_bodies='    <TextFrame Self="tf1" ParentStory="s" ItemTransform="1 0 0 1 20 60"/>',
                styles_body=styles,
            )
            md = _idml_to_markdown(str(idml_path))
        self.assertIn("[[UI]]Login[[/UI]]", md)
        xml = _content_to_dita_xml("Run", md, "task", "DTO041221", "en-US", topic_kind="task")
        self.assertIn("<uicontrol>Login</uicontrol>", xml)
        self.assertNotIn("[[UI]]", xml)

    def test_idml_italic_character_style_becomes_i(self):
        """Cross reference character styles (N/c/V, book titles) become <i>."""
        import tempfile

        body = (
            '<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/Body">'
            '<CharacterStyleRange><Content>其中，</Content></CharacterStyleRange>'
            '<CharacterStyleRange AppliedCharacterStyle="CharacterStyle/Cross reference">'
            "<Content>N</Content></CharacterStyleRange>"
            '<CharacterStyleRange><Content>为长度，</Content></CharacterStyleRange>'
            '<CharacterStyleRange AppliedCharacterStyle="CharacterStyle/Cross reference-in para.-CN">'
            "<Content>c</Content></CharacterStyleRange>"
            '<CharacterStyleRange><Content>为浓度。</Content></CharacterStyleRange>'
            "</ParagraphStyleRange>"
        )
        with tempfile.TemporaryDirectory() as tmpdir:
            idml_path = Path(tmpdir) / "manual_italic.idml"
            _write_idml_package(
                idml_path,
                story_files={"Story_u1.xml": _idml_story(body)},
                spread_bodies='    <TextFrame Self="tf1" ParentStory="s" ItemTransform="1 0 0 1 20 60"/>',
            )
            md = _idml_to_markdown(str(idml_path))
        self.assertIn("[[I]]N[[/I]]", md)
        self.assertIn("[[I]]c[[/I]]", md)
        xml = _content_to_dita_xml(
            "文库定量", md, "concept", "DTC041240", "zh-CN",
            topic_kind="concept", source_format="idml",
        )
        self.assertIn("<i>N</i>为长度", xml)
        self.assertIn("<i>c</i>为浓度", xml)

    def test_idml_icon_stays_inline_between_text(self):
        """An inline button icon keeps its position inside the step text."""
        import tempfile

        body = (
            '<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/StepsList">'
            '<CharacterStyleRange><Content>点击</Content>'
            '<Rectangle Self="re1"><PDF Self="pdf1">'
            '<Link LinkResourceURI="file:C:/doc/Links/icon_nextstep.ai"/>'
            "</PDF></Rectangle>"
            "<Content>，托盘将自动收回至仪器内。</Content></CharacterStyleRange>"
            "</ParagraphStyleRange>"
        )
        with tempfile.TemporaryDirectory() as tmpdir:
            idml_path = Path(tmpdir) / "manual_icon.idml"
            _write_idml_package(
                idml_path,
                story_files={"Story_u1.xml": _idml_story(body)},
                spread_bodies='    <TextFrame Self="tf1" ParentStory="s" ItemTransform="1 0 0 1 20 60"/>',
            )
            md = _idml_to_markdown(str(idml_path))
        # The icon reference resolves to nothing (not packaged), so the step
        # text simply stays intact with no stray icon escaping it.
        self.assertIn("1. 点击，托盘将自动收回至仪器内。", md)
        xml = _content_to_dita_xml("装载", md, "task", "DTO041222", "zh-CN", topic_kind="task")
        self.assertEqual(xml.count("<li>"), 1)

    def test_idml_cross_reference_becomes_xref(self):
        """An InDesign cross-reference parses to [[XREF:title#page]]; DITA
        emits an empty xref and leaves page text to IME."""
        import tempfile

        body = (
            '<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/Body">'
            "<CharacterStyleRange><Content>具体操作可参考</Content></CharacterStyleRange>"
            '<CrossReferenceSource Self="xref1" AppliedFormat="uda6">'
            "<CharacterStyleRange><Content>第</Content></CharacterStyleRange>"
            '<CharacterStyleRange PageNumberType="TextVariable">'
            '<TextVariableInstance Name="TV XRefPageNumber" ResultText="12"/>'
            "</CharacterStyleRange>"
            "<CharacterStyleRange><Content>页“准备测序试剂槽”</Content></CharacterStyleRange>"
            "</CrossReferenceSource>"
            '<CharacterStyleRange><Content>。</Content></CharacterStyleRange>'
            "</ParagraphStyleRange>"
        )
        with tempfile.TemporaryDirectory() as tmpdir:
            idml_path = Path(tmpdir) / "manual_xref.idml"
            _write_idml_package(
                idml_path,
                story_files={"Story_u1.xml": _idml_story(body)},
                spread_bodies='    <TextFrame Self="tf1" ParentStory="s" ItemTransform="1 0 0 1 20 60"/>',
            )
            md = _idml_to_markdown(str(idml_path))
        self.assertIn("具体操作可参考[[XREF:准备测序试剂槽#12]]。", md)
        self.assertNotIn("第页", md)

        xml = _content_to_dita_xml(
            "试剂盒暂存", md, "concept", "DTC041230", "zh-CN",
            topic_kind="concept", source_format="idml",
            xref_map={"准备测序试剂槽": "DTO041031"},
        )
        self.assertIn(
            '<xref href="DTO041031.dita" cms:title="准备测序试剂槽"/>',
            xml,
        )
        self.assertNotIn("第12 页准备测序试剂槽", xml)
        self.assertNotIn("[[XREF:", xml)

    def test_unresolved_cross_reference_falls_back_to_title_text(self):
        """A cross-reference to an unknown target keeps its visible title."""
        md = "参考[[XREF:不存在的章节]]。\n"
        xml = _content_to_dita_xml(
            "正文", md, "concept", "DTC041231", "zh-CN",
            topic_kind="concept", source_format="idml", xref_map={},
        )
        self.assertIn("参考不存在的章节。", xml)
        self.assertNotIn("<xref", xml)
        self.assertNotIn("[[XREF:", xml)

    def test_xref_inside_table_cell_keeps_page_number(self):
        """Page-bearing xref markers stay in one table cell as an empty xref."""
        md = (
            "| 说明 |\n| :--- |\n"
            "| 具体操作可参考[[XREF:试剂槽底部胶塞歪斜或脱落#30]]。 |\n"
        )
        xml = _content_to_dita_xml(
            "装载载片", md, "task", "DTO041027", "zh-CN",
            topic_kind="task", source_format="idml",
            xref_map={"试剂槽底部胶塞歪斜或脱落": "DTC041040"},
        )
        self.assertIn(
            '<xref href="DTC041040.dita" cms:title="试剂槽底部胶塞歪斜或脱落"/>',
            xml,
        )
        self.assertNotIn("第30 页", xml)
        self.assertNotIn("[[XREF:", xml)

    def test_idml_table_column_widths_become_colspec(self):
        """IDML Column/@SingleColumnWidth maps to relative CALS colwidth."""
        import xml.etree.ElementTree as ET

        table = (
            '<Table Self="t1" HeaderRowCount="1" BodyRowCount="1" ColumnCount="2">'
            '<Row Self="r0" Name="0"/><Row Self="r1" Name="1"/>'
            '<Column Self="co0" Name="0" SingleColumnWidth="87.37"/>'
            '<Column Self="co1" Name="1" SingleColumnWidth="203.86"/>'
            '<Cell Self="c0" Name="0:0" RowSpan="1" ColumnSpan="1">'
            "<ParagraphStyleRange><CharacterStyleRange><Content>组分</Content>"
            "</CharacterStyleRange></ParagraphStyleRange></Cell>"
            '<Cell Self="c1" Name="1:0" RowSpan="1" ColumnSpan="1">'
            "<ParagraphStyleRange><CharacterStyleRange><Content>体积</Content>"
            "</CharacterStyleRange></ParagraphStyleRange></Cell>"
            '<Cell Self="c2" Name="0:1" RowSpan="1" ColumnSpan="1">'
            "<ParagraphStyleRange><CharacterStyleRange><Content>A</Content>"
            "</CharacterStyleRange></ParagraphStyleRange></Cell>"
            '<Cell Self="c3" Name="1:1" RowSpan="1" ColumnSpan="1">'
            "<ParagraphStyleRange><CharacterStyleRange><Content>1 μL</Content>"
            "</CharacterStyleRange></ParagraphStyleRange></Cell>"
            "</Table>"
        )
        lines = _idml_table_to_markdown(ET.fromstring(table))
        self.assertEqual(lines[0], "[[COLW:0.30*,0.70*]]")
        xml = _content_to_dita_xml(
            "组分", "\n".join(lines), "concept", "DTC041241", "zh-CN",
            topic_kind="concept", source_format="idml",
        )
        self.assertIn('<colspec colname="COLSPEC0" colwidth="0.30*"/>', xml)
        self.assertIn('<colspec colname="COLSPEC1" colwidth="0.70*"/>', xml)

    def test_figure_title_style_marks_caption_and_wraps_image(self):
        """A Figure-title paragraph becomes a <fig><title> bound to its image."""
        import tempfile

        styles = (
            '    <ParagraphStyle Self="ParagraphStyle/For reagent kit only:Figure-title-CN"'
            ' Name="For reagent kit only:Figure-title-CN"/>\n'
            '    <ParagraphStyle Self="ParagraphStyle/For reagent kit only:Title-L1-with number-CN"'
            ' Name="For reagent kit only:Title-L1-with number-CN"/>\n'
        )
        body = (
            '<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/For reagent kit only:Title-L1-with number-CN">'
            "<CharacterStyleRange><Content>准备测序试剂槽</Content></CharacterStyleRange></ParagraphStyleRange>"
            '<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/For reagent kit only:Figure-title-CN">'
            "<CharacterStyleRange><Content>竖直放置试剂槽</Content></CharacterStyleRange></ParagraphStyleRange>"
        )
        with tempfile.TemporaryDirectory() as tmpdir:
            idml_path = Path(tmpdir) / "manual.idml"
            _write_idml_package(
                idml_path,
                story_files={"Story_u1.xml": _idml_story(body)},
                spread_bodies='    <TextFrame Self="tf1" ParentStory="s" ItemTransform="1 0 0 1 20 60"/>',
                styles_body=styles,
            )
            md = _idml_to_markdown(str(idml_path))
            self.assertIn("[[FIGCAP]]竖直放置试剂槽", md)

        xml = _content_to_dita_xml(
            "准备测序试剂槽",
            "![试剂槽.svg](image/试剂槽.svg)\n[[FIGCAP]]竖直放置试剂槽",
            "task",
            "DTO041023",
            "zh-CN",
            topic_kind="task",
        )
        self.assertIn(
            '<fig><title>竖直放置试剂槽</title><image href="image/试剂槽.svg" placement="break"></image></fig>',
            xml,
        )

    def test_table_cells_are_placed_by_name_attribute(self):
        """IDML stores cells as direct children positioned by Name="col:row"."""
        import tempfile

        def _cell(name, text):
            return (
                f'<Cell Self="c{name}" Name="{name}" RowSpan="1" ColumnSpan="1">'
                '<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/Text:L2-CN">'
                f"<CharacterStyleRange><Content>{text}</Content></CharacterStyleRange>"
                "</ParagraphStyleRange></Cell>"
            )

        table = (
            '<Table Self="t1">'
            '<Row Self="r0" Name="0"/><Row Self="r1" Name="1"/>'
            '<Column Self="co0" Name="0"/><Column Self="co1" Name="1"/>'
            + _cell("0:0", "版本") + _cell("1:0", "4.0")
            + _cell("0:1", "发布日期") + _cell("1:1", "2026年10月9日")
            + "</Table>"
        )
        body = (
            '<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/Heading1">'
            "<CharacterStyleRange><Content>版本记录</Content></CharacterStyleRange></ParagraphStyleRange>"
            '<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/Text:L2-CN">'
            f"<CharacterStyleRange>{table}</CharacterStyleRange></ParagraphStyleRange>"
        )
        with tempfile.TemporaryDirectory() as tmpdir:
            idml_path = Path(tmpdir) / "manual.idml"
            _write_idml_package(
                idml_path,
                story_files={"Story_u1.xml": _idml_story(body)},
                spread_bodies='    <TextFrame Self="tf1" ParentStory="s" ItemTransform="1 0 0 1 20 60"/>',
                styles_body=(
                    '    <ParagraphStyle Self="ParagraphStyle/Heading1" Name="Heading 1"/>\n'
                    '    <ParagraphStyle Self="ParagraphStyle/Text:L2-CN" Name="Text:L2-CN"/>\n'
                ),
            )
            md = _idml_to_markdown(str(idml_path))
            self.assertIn("| 版本 | 4.0 |", md)
            self.assertIn("| 发布日期 | 2026年10月9日 |", md)

    def test_paragraph_with_anchored_table_keeps_its_own_text(self):
        """A callout table anchored inside a paragraph must not swallow the
        paragraph's own text."""
        import tempfile

        def _cell(name, text):
            return (
                f'<Cell Self="c{name}" Name="{name}" RowSpan="1" ColumnSpan="1">'
                '<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/Text:L2-CN">'
                f"<CharacterStyleRange><Content>{text}</Content></CharacterStyleRange>"
                "</ParagraphStyleRange></Cell>"
            )

        table = (
            '<Table Self="t1">'
            '<Row Self="r0" Name="0"/><Row Self="r1" Name="1"/>'
            '<Column Self="co0" Name="0"/><Column Self="co1" Name="1"/>'
            + _cell("0:0", "危害") + _cell("1:0", "说明")
            + _cell("0:1", "夹手") + _cell("1:1", "请勿将手置于仓门之间。")
            + "</Table>"
        )
        body = (
            '<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/Heading1">'
            "<CharacterStyleRange><Content>操作</Content></CharacterStyleRange></ParagraphStyleRange>"
            '<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/Text:L2-CN">'
            "<CharacterStyleRange><Content>此时，试剂仓中的托盘进行初始化。</Content><Br/></CharacterStyleRange>"
            f"<CharacterStyleRange>{table}</CharacterStyleRange>"
            "</ParagraphStyleRange>"
        )
        with tempfile.TemporaryDirectory() as tmpdir:
            idml_path = Path(tmpdir) / "manual.idml"
            _write_idml_package(
                idml_path,
                story_files={"Story_u1.xml": _idml_story(body)},
                spread_bodies='    <TextFrame Self="tf1" ParentStory="s" ItemTransform="1 0 0 1 20 60"/>',
                styles_body=(
                    '    <ParagraphStyle Self="ParagraphStyle/Heading1" Name="Heading 1"/>\n'
                    '    <ParagraphStyle Self="ParagraphStyle/Text:L2-CN" Name="Text:L2-CN"/>\n'
                ),
            )
            md = _idml_to_markdown(str(idml_path))
            self.assertIn("此时，试剂仓中的托盘进行初始化。", md)
            self.assertIn("| 危害 | 说明 |", md)

    def test_first_page_body_br_becomes_separate_paragraphs(self):
        """关于说明书 uses one paragraph style with <Br/> between logical paras."""
        import tempfile

        body = (
            '<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/Title:Title in first page-CN">'
            "<CharacterStyleRange><Content>关于说明书</Content></CharacterStyleRange>"
            "</ParagraphStyleRange>"
            '<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/Text:Text in first page-CN">'
            "<CharacterStyleRange><Content>本说明书适用于DNBSEQ-E25RS高通量测序试剂套装。</Content></CharacterStyleRange>"
            "<Br/>"
            "<CharacterStyleRange><Content>本说明书的读者为终端用户。</Content></CharacterStyleRange>"
            "</ParagraphStyleRange>"
        )
        with tempfile.TemporaryDirectory() as tmpdir:
            idml_path = Path(tmpdir) / "manual.idml"
            _write_idml_package(
                idml_path,
                story_files={"Story_u1.xml": _idml_story(body)},
                spread_bodies='    <TextFrame Self="tf1" ParentStory="s" ItemTransform="1 0 0 1 20 60"/>',
                styles_body=(
                    '    <ParagraphStyle Self="ParagraphStyle/Title:Title in first page-CN" Name="Title:Title in first page-CN"/>\n'
                    '    <ParagraphStyle Self="ParagraphStyle/Text:Text in first page-CN" Name="Text:Text in first page-CN"/>\n'
                ),
            )
            md = _idml_to_markdown(str(idml_path))
        self.assertIn("本说明书适用于DNBSEQ-E25RS高通量测序试剂套装。", md)
        self.assertIn("本说明书的读者为终端用户。", md)
        self.assertNotIn(
            "本说明书适用于DNBSEQ-E25RS高通量测序试剂套装。 本说明书的读者为终端用户。",
            md,
        )
        xml = _content_to_dita_xml(
            "关于说明书", md.split("# 关于说明书", 1)[-1], "concept", "DTC041001", "zh-CN",
            topic_kind="concept", source_format="idml",
        )
        self.assertIn("<p>本说明书适用于DNBSEQ-E25RS高通量测序试剂套装。</p>", xml)
        self.assertIn("<p>本说明书的读者为终端用户。</p>", xml)

    def test_table_cell_bulleted_list_becomes_ul(self):
        """版本记录修订摘要 uses Table:Bulleted list with <Br/> between items."""
        import tempfile

        def _cell(name, style, *runs):
            inner = "".join(
                f"<CharacterStyleRange><Content>{text}</Content></CharacterStyleRange>"
                if text != "<Br/>"
                else "<Br/>"
                for text in runs
            )
            return (
                f'<Cell Self="c{name}" Name="{name}" RowSpan="1" ColumnSpan="1">'
                f'<ParagraphStyleRange AppliedParagraphStyle="{style}">'
                f"{inner}</ParagraphStyleRange></Cell>"
            )

        table = (
            '<Table Self="t1" HeaderRowCount="1">'
            '<Row Self="r0" Name="0"/><Row Self="r1" Name="1"/>'
            '<Column Self="co0" Name="0"/><Column Self="co1" Name="1"/><Column Self="co2" Name="2"/>'
            + _cell("0:0", "ParagraphStyle/Table:Table heading-CN", "版本")
            + _cell("1:0", "ParagraphStyle/Table:Table heading-CN", "发布日期")
            + _cell("2:0", "ParagraphStyle/Table:Table heading-CN", "修订内容摘要")
            + _cell("0:1", "ParagraphStyle/Table:Single sentence-CN", "8.0")
            + _cell("1:1", "ParagraphStyle/Table:Single sentence-CN", "2026年10月9日")
            + _cell(
                "2:1",
                "ParagraphStyle/Table:Bulleted list-L1-CN",
                "载片运输温度由10 ℃~25 ℃ 改为2 ℃~25 ℃",
                "<Br/>",
                "仪器操作描述优化",
            )
            + "</Table>"
        )
        body = (
            '<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/Heading1">'
            "<CharacterStyleRange><Content>版本记录</Content></CharacterStyleRange></ParagraphStyleRange>"
            + table
        )
        with tempfile.TemporaryDirectory() as tmpdir:
            idml_path = Path(tmpdir) / "manual.idml"
            _write_idml_package(
                idml_path,
                story_files={"Story_u1.xml": _idml_story(body)},
                spread_bodies='    <TextFrame Self="tf1" ParentStory="s" ItemTransform="1 0 0 1 20 60"/>',
                styles_body=(
                    '    <ParagraphStyle Self="ParagraphStyle/Heading1" Name="Heading 1"/>\n'
                    '    <ParagraphStyle Self="ParagraphStyle/Table:Table heading-CN" Name="Table:Table heading-CN"/>\n'
                    '    <ParagraphStyle Self="ParagraphStyle/Table:Single sentence-CN" Name="Table:Single sentence-CN"/>\n'
                    '    <ParagraphStyle Self="ParagraphStyle/Table:Bulleted list-L1-CN" Name="Table:Bulleted list-L1-CN">\n'
                    "      <Properties><BulletsAndNumberingListType>BulletList</BulletsAndNumberingListType></Properties>\n"
                    "    </ParagraphStyle>\n"
                ),
            )
            md = _idml_to_markdown(str(idml_path))
        self.assertIn(
            "[[UL]]载片运输温度由10 ℃~25 ℃ 改为2 ℃~25 ℃[[LI]]仪器操作描述优化[[/UL]]",
            md,
        )
        xml = _content_to_dita_xml(
            "版本记录", md.split("# 版本记录", 1)[-1], "concept", "DTC041002", "zh-CN",
            topic_kind="concept", source_format="idml",
        )
        self.assertIn(
            "<entry><ul><li><p>载片运输温度由10 ℃~25 ℃ 改为2 ℃~25 ℃</p></li><li><p>仪器操作描述优化</p></li></ul></entry>",
            xml,
        )

    def test_running_header_only_story_is_dropped(self):
        import tempfile

        with tempfile.TemporaryDirectory() as tmpdir:
            idml_path = Path(tmpdir) / "manual.idml"
            _write_idml_package(
                idml_path,
                story_files={
                    "Story_body.xml": _idml_story(
                        '<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/Heading1">'
                        "<CharacterStyleRange><Content>介绍</Content></CharacterStyleRange>"
                        "</ParagraphStyleRange>",
                        self_id="body",
                    ),
                    "Story_head.xml": _idml_story(
                        '<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/$ID/NormalParagraphStyle">'
                        "<CharacterStyleRange><Content>内部资料</Content></CharacterStyleRange>"
                        "</ParagraphStyleRange>",
                        self_id="head",
                    ),
                },
                spread_bodies=(
                    '    <TextFrame Self="tf1" ParentStory="body" ItemTransform="1 0 0 1 20 60"/>\n'
                    '    <TextFrame Self="tf2" ParentStory="head" ItemTransform="1 0 0 1 20 10"/>'
                ),
            )
            md = _idml_to_markdown(str(idml_path))
            self.assertIn("# 介绍", md)
            self.assertNotIn("内部资料", md)

    def test_two_row_table_header_stays_in_thead(self):
        """HeaderRowCount=2 must keep both rows in <thead>, not leak row 1."""
        import xml.etree.ElementTree as ET

        def cell(name, text, rowspan=1, colspan=1):
            return (
                f'<Cell Name="{name}" RowSpan="{rowspan}" ColumnSpan="{colspan}">'
                '<ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/Table body">'
                f"<CharacterStyleRange><Content>{text}</Content></CharacterStyleRange>"
                "</ParagraphStyleRange></Cell>"
            )

        table = (
            '<Table Self="t1" HeaderRowCount="2" BodyRowCount="1" ColumnCount="3">'
            + cell("0:0", "型号", rowspan=2)
            + cell("1:0", "解冻方式", colspan=2)
            + cell("1:1", "2 ℃~8 ℃冰箱（h）")
            + cell("2:1", "15 ℃~25 ℃室温（h）")
            + cell("0:2", "FCL SE100")
            + cell("1:2", "6")
            + cell("2:2", "3.5~4.5")
            + "</Table>"
        )
        lines = _idml_table_to_markdown(ET.fromstring(table))
        self.assertEqual(lines[0], "[[THEAD:2]]")
        xml = _content_to_dita_xml(
            "试剂槽解冻时间",
            "\n".join(lines),
            "concept",
            "DTC041230",
            "zh-CN",
            topic_kind="concept",
        )
        thead = xml.split("<thead>", 1)[1].split("</thead>", 1)[0]
        tbody = xml.split("<tbody>", 1)[1].split("</tbody>", 1)[0]
        self.assertEqual(thead.count("<row>"), 2)
        self.assertIn('morerows="1"', thead)
        self.assertIn("2 ℃~8 ℃冰箱", thead)
        self.assertNotIn("2 ℃~8 ℃冰箱", tbody)

    def test_table_cell_icon_keeps_inline_position(self):
        """An icon inside a table cell must not disappear during conversion."""
        import xml.etree.ElementTree as ET

        table = (
            '<Table Self="t1" HeaderRowCount="1" BodyRowCount="1" ColumnCount="2">'
            '<Cell Name="0:0"><ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/Table heading">'
            "<CharacterStyleRange><Content>项目</Content></CharacterStyleRange></ParagraphStyleRange></Cell>"
            '<Cell Name="1:0"><ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/Table heading">'
            "<CharacterStyleRange><Content>说明</Content></CharacterStyleRange></ParagraphStyleRange></Cell>"
            '<Cell Name="0:1"><ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/Table body">'
            "<CharacterStyleRange><Content>点击</Content>"
            '<Link LinkResourceURI="file:C:/doc/Links/icon_emergencystop.ai"/>'
            "<Content>可提前终止测序。</Content></CharacterStyleRange></ParagraphStyleRange></Cell>"
            '<Cell Name="1:1"><ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/Table body">'
            "<CharacterStyleRange><Content>说明文字</Content></CharacterStyleRange></ParagraphStyleRange></Cell>"
            "</Table>"
        )
        resolver = lambda uri: "image/icon_emergencystop.svg"
        lines = _idml_table_to_markdown(ET.fromstring(table), None, resolver)
        self.assertIn(
            "![icon_emergencystop.svg](image/icon_emergencystop.svg)",
            "\n".join(lines),
        )
        xml = _content_to_dita_xml(
            "测序",
            "\n".join(lines),
            "concept",
            "DTC041031",
            "zh-CN",
            topic_kind="concept",
        )
        entry = xml.split("点击", 1)[1]
        self.assertTrue(
            entry.startswith(
                '<image href="image/icon_emergencystop.svg"></image>可提前终止测序。'
            )
        )

    def test_sentence_list_callout_becomes_note_ul(self):
        """A "sentence list" callout body is a ul inside the note."""
        import xml.etree.ElementTree as ET

        table = (
            '<Table Self="t1" HeaderRowCount="0" BodyRowCount="1" ColumnCount="2">'
            '<Cell Name="0:0"><ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/Text:Symbols-CN">'
            '<CharacterStyleRange><Link LinkResourceURI="file:C:/doc/Links/icon_noteCH.ai"/>'
            "</CharacterStyleRange></ParagraphStyleRange></Cell>"
            '<Cell Name="1:0"><ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/Table:Warning sentence list-CN">'
            "<CharacterStyleRange><Content>第一条提示。</Content><Br/>"
            "<Content>第二条提示。</Content></CharacterStyleRange></ParagraphStyleRange></Cell>"
            "</Table>"
        )
        lines = _idml_table_to_markdown(ET.fromstring(table))
        self.assertIn("| - 第一条提示。 |", lines)
        self.assertIn("| - 第二条提示。 |", lines)
        xml = _content_to_dita_xml(
            "提示",
            "\n".join(lines),
            "concept",
            "DTC041231",
            "zh-CN",
            topic_kind="concept",
        )
        note = xml.split('<note type="tip">', 1)[1].split("</note>", 1)[0]
        self.assertEqual(note.count("<li>"), 2)
        self.assertIn("<ul>", note)

    def test_nested_list_levels_are_preserved(self):
        """Level 3 bullets and level 4 steps nest under the parent step."""
        md = (
            "操作步骤如下：\n\n"
            "[[LVL:2]]1. 确认方案无误后，点击进入下一步。\n\n"
            "[[LVL:3]]- 如需退出测序，执行如下操作：\n\n"
            "[[LVL:4]]1. 点击确定。\n\n"
            "[[LVL:4]]1. 点击下一步。\n\n"
            "[[LVL:3]]- 如需继续测序，点击取消。\n"
        )
        xml = _content_to_dita_xml(
            "选择测序方案", md, "task", "DTO041232", "zh-CN", topic_kind="task"
        )
        step_start = xml.index("确认方案无误后")
        sub_bullet = xml.index("<ul>", step_start)
        sub_step = xml.index("<ol>", sub_bullet)
        next_bullet = xml.index("如需继续测序", sub_bullet)
        self.assertLess(step_start, sub_bullet)
        self.assertLess(sub_bullet, sub_step)
        self.assertLess(sub_step, next_bullet)
        self.assertEqual(xml.count("</ol>"), 2)
        self.assertEqual(xml.count("</ul>"), 1)

    def test_non_list_status_icon_stays_inline(self):
        """A status icon inside a prose paragraph stays in its sentence."""
        import xml.etree.ElementTree as ET

        body = (
            '<Story><ParagraphStyleRange AppliedParagraphStyle="ParagraphStyle/Text:L3-CN">'
            '<CharacterStyleRange><Content>如自检失败，点击</Content></CharacterStyleRange>'
            '<CharacterStyleRange AppliedCharacterStyle="CharacterStyle/icon width">'
            '<Link LinkResourceURI="file:C:/doc/Links/icon_status sequencing_1_1.ai"/>'
            "</CharacterStyleRange>"
            '<CharacterStyleRange><Content>重试。多次自检失败，请联系技术支持。</Content>'
            "</CharacterStyleRange></ParagraphStyleRange></Story>"
        )
        blocks = []
        _idml_walk(
            ET.fromstring(body),
            {},
            {},
            blocks,
            image_resolver=lambda uri: "image/icon_status_sequencing_1_1.svg",
        )
        self.assertEqual(len(blocks), 1)
        self.assertEqual(blocks[0]["kind"], "para")
        self.assertIn(
            "点击![icon_status_sequencing_1_1.svg](image/icon_status_sequencing_1_1.svg)重试。",
            blocks[0]["text"],
        )

    def test_level_marker_does_not_hide_numbered_steps(self):
        """A [[LVL:n]] prefix must still let the section be typed as a task."""
        content = "操作步骤如下：\n\n[[LVL:2]]1. 点击确定。\n\n[[LVL:2]]1. 点击下一步。\n"
        self.assertEqual(_section_dita_type("选择测序方案", content), "task")

    def test_list_body_paragraph_stays_inside_parent_item(self):
        """A formula paragraph after nested bullets hangs on the parent li."""
        md = (
            "[[LVL:2]]- 文库浓度要求\n\n"
            "[[LVL:3]]- 对于MGI接头文库，要求初始dsDNA文库浓度不小于2 fmol/µL。\n\n"
            "ng与fmol换算公式如下：\n\n"
            "![formula.svg](image/formula.svg)\n\n"
            "330为碱基平均分子量。\n\n"
            "[[LVL:2]]- 文库量要求\n"
        )
        xml = _content_to_dita_xml(
            "文库浓度及所需量的要求", md, "concept", "DTC041016", "zh-CN",
            topic_kind="concept", source_format="idml",
        )
        conc_li = xml.split("文库浓度要求", 1)[1].split("文库量要求", 1)[0]
        self.assertIn("<ul>", conc_li)
        self.assertIn("对于MGI接头文库", conc_li)
        self.assertIn("<p>ng与fmol换算公式如下：</p>", conc_li)
        self.assertIn('<image href="image/formula.svg"', conc_li)
        self.assertIn("<p>330为碱基平均分子量。</p>", conc_li)
        formula_pos = conc_li.index("ng与fmol换算公式如下")
        nested_close = conc_li.index("</ul>")
        self.assertLess(nested_close, formula_pos)
        self.assertNotIn("ng与fmol换算公式如下", xml.rsplit("</ul>", 1)[-1])

    def test_xref_link_text_omits_section_number(self):
        """Resolved xref is empty; cms:title drops 9.1 numbering."""
        xml = _content_to_dita_xml(
            "测定DNB浓度",
            "否则，参考[[XREF:9.1DNB浓度不合格#30]]。\n",
            "concept",
            "DTC041021",
            "zh-CN",
            topic_kind="concept",
            source_format="idml",
            xref_map={"DNB浓度不合格": "DTO041034"},
        )
        self.assertIn(
            '<xref href="DTO041034.dita" cms:title="DNB浓度不合格"/>',
            xml,
        )
        self.assertNotIn("9.1", xml)
        self.assertNotIn("第30 页", xml)

    def test_xref_link_text_spaces_alnum_before_cjk(self):
        """cms:title keeps the compact title; unresolved fallback spaces alnum/CJK."""
        xml = _content_to_dita_xml(
            "DNB浓度不合格",
            "当DNB浓度不符合[[XREF:测定DNB浓度#12]]的要求时。\n",
            "task",
            "DTO041034",
            "zh-CN",
            topic_kind="task",
            source_format="idml",
            xref_map={"测定DNB浓度": "DTC041021"},
        )
        self.assertIn(
            '<xref href="DTC041021.dita" cms:title="测定DNB浓度"/>',
            xml,
        )
        self.assertNotIn("第12 页测定DNB 浓度", xml)

        unresolved = _content_to_dita_xml(
            "DNB浓度不合格",
            "当DNB浓度不符合[[XREF:测定DNB浓度#12]]的要求时。\n",
            "task",
            "DTO041034",
            "zh-CN",
            topic_kind="task",
            source_format="idml",
            xref_map={},
        )
        self.assertIn("第12 页“测定DNB 浓度”", unresolved)
        self.assertNotIn("<xref", unresolved)

    def test_list_body_paragraph_stays_before_note_and_figure(self):
        """A step continuation paragraph stays ahead of the following note/fig."""
        md = (
            "[[LVL:2]]1. 用扫码枪扫描测序载片塑料包装上的二维码。\n\n"
            "如扫码失败，可点击文本框，用弹出的软键盘手动输入信息。\n\n"
            "| [[NOTE:warning]] |\n| :--- |\n| 载片装载完成后，请及时进行测序。 |\n\n"
            "![screen.svg](image/screen.svg)\n\n"
            "[[LVL:2]]1. 下一步。\n"
        )
        xml = _content_to_dita_xml(
            "装载载片", md, "task", "DTO041027", "zh-CN",
            topic_kind="task", source_format="idml",
        )
        first_li = xml.split("<li>", 1)[1].split("</li>", 1)[0]
        para_pos = first_li.index("如扫码失败")
        note_pos = first_li.index('<note type="warning">')
        fig_pos = first_li.index("screen.svg")
        self.assertLess(para_pos, note_pos)
        self.assertLess(note_pos, fig_pos)

    def test_list_body_paragraph_stays_before_nested_list(self):
        """A continuation paragraph before nested bullets stays ahead of them."""
        md = (
            "[[LVL:2]]1. 确认方案无误后，点击进入下一步。\n\n"
            "如点击图标，界面出现提示信息。\n\n"
            "[[LVL:3]]- 如需退出测序，执行如下操作：\n\n"
            "[[LVL:3]]- 如需继续测序，点击取消。\n"
        )
        xml = _content_to_dita_xml(
            "回顾参数", md, "task", "DTO041026", "zh-CN",
            topic_kind="task", source_format="idml",
        )
        first_li = xml.split("<li>", 1)[1].rsplit("</li>", 1)[0]
        self.assertLess(first_li.index("如点击图标"), first_li.index("<ul>"))

    def test_note_after_parent_continuation_stays_on_parent(self):
        """Notes after a parent-level paragraph hang on the parent, not the nested li."""
        md = (
            "[[LVL:2]]- 若未生成报告，请根据提示信息处理。\n\n"
            "[[LVL:3]]- 若为设备问题，请联系技术支持。\n\n"
            "[[LVL:3]]- 若因用户操作不当，请重新测序。\n\n"
            "如下界面显示后，测序完成。\n\n"
            "| [[NOTE:warning]] |\n| :--- |\n| 请勿将手伸入仓门。 |\n"
        )
        xml = _content_to_dita_xml(
            "测序完成", md, "concept", "DTC041030", "zh-CN",
            topic_kind="concept", source_format="idml",
        )
        parent = xml.split("若未生成报告", 1)[1]
        nested_li = parent.split("若因用户操作不当", 1)[1].split("</li>", 1)[0]
        self.assertNotIn('<note type="warning">', nested_li)
        self.assertLess(parent.index("</ul>"), parent.index("如下界面显示后"))
        self.assertLess(parent.index("如下界面显示后"), parent.index('<note type="warning">'))


class NativeMarkdownToDitaTest(unittest.TestCase):
    def test_headingless_markdown_becomes_one_section(self):
        sections = _parse_content("md", None, "试剂使用前请平衡至室温。\n\n请勿反复冻融。")
        self.assertEqual(len(sections), 1)
        self.assertEqual(sections[0]["title"], "文档")
        self.assertIn("试剂使用前请平衡至室温。", sections[0]["content"])

    def test_preamble_before_first_heading_is_kept(self):
        md = "本手册说明仪器操作。\n\n# 装载试剂\n\n将试剂取出。\n"
        sections = _parse_content("md", None, md)
        self.assertEqual(sections[0]["title"], "正文")
        self.assertIn("本手册说明仪器操作。", sections[0]["content"])
        self.assertEqual(sections[1]["title"], "装载试剂")

    def test_setext_and_yaml_title(self):
        md = "---\ntitle: Yellow 16 用户手册\n---\n\n操作说明\n======\n\n第一步。\n"
        sections = _parse_content("md", None, md)
        self.assertEqual(sections[0]["title"], "操作说明")
        self.assertIn("第一步。", sections[0]["content"])

    def test_yaml_title_used_when_document_has_no_heading(self):
        md = "---\ntitle: 产品简介\n---\n\n这是正文。\n"
        sections = _parse_content("md", None, md)
        self.assertEqual(sections[0]["title"], "产品简介")
        self.assertIn("这是正文。", sections[0]["content"])

    def test_native_bold_italic_code_and_link(self):
        md = "点击 **确定** 后查看 *结果*，命令为 `dnb seq`，详见 [手册](https://example.com/manual)。"
        xml = _content_to_dita_xml(
            "操作", md, "concept", "DTC041301", "zh-CN",
            topic_kind="concept", source_format="md",
        )
        self.assertIn("<b>确定</b>", xml)
        self.assertIn("<i>结果</i>", xml)
        self.assertIn("dnb seq", xml)
        self.assertNotIn("<codeph>", xml)
        self.assertIn(
            '<xref href="https://example.com/manual" scope="external" format="html">手册</xref>',
            xml,
        )

    def test_file_format_inline_code_becomes_italic(self):
        md = "输入支持 `.fq` 或 `.fq.gz`，输出为 `.arc`，命令为 `zarc -c`。"
        xml = _content_to_dita_xml(
            "压缩", md, "concept", "DTC041320", "zh-CN",
            topic_kind="concept", source_format="md",
        )
        self.assertIn("<i>.fq</i>", xml)
        self.assertIn("<i>.fq.gz</i>", xml)
        self.assertIn("<i>.arc</i>", xml)
        self.assertIn("zarc -c", xml)
        self.assertNotIn("<codeph>", xml)

    def test_plain_file_format_token_becomes_italic(self):
        xml = _content_to_dita_xml(
            "参数", "输出文件名称（.arc后缀）\n", "concept", "DTC041321", "zh-CN",
            topic_kind="concept", source_format="md",
        )
        self.assertIn("<i>.arc</i>后缀", xml)
        self.assertNotIn("<codeph>", xml)

    def test_blockquote_note_becomes_caution(self):
        xml = _content_to_dita_xml(
            "准备",
            "> **注意：** 请先平衡试剂。\n",
            "concept",
            "DTC041302",
            "zh-CN",
            topic_kind="concept",
            source_format="md",
        )
        self.assertIn('<note type="caution">', xml)
        self.assertIn("请先平衡试剂。", xml)

    def test_fenced_code_becomes_codeblock(self):
        md = "示例：\n\n```python\nprint(1)\nprint(2)\n```\n"
        xml = _content_to_dita_xml(
            "示例", md, "concept", "DTC041303", "zh-CN",
            topic_kind="concept", source_format="md",
        )
        self.assertIn('<codeblock outputclass="python" xml:space="preserve">', xml)
        self.assertIn("print(1)", xml)
        self.assertIn("print(2)", xml)
        self.assertNotIn("<b>", xml.split("<codeblock", 1)[1].split("</codeblock>", 1)[0])

    def test_indented_nested_list_stays_nested(self):
        md = (
            "1. 确认方案。\n"
            "   - 如需退出，点击取消。\n"
            "   - 如需继续，点击确定。\n"
            "2. 开始测序。\n"
        )
        xml = _content_to_dita_xml(
            "回顾参数", md, "task", "DTO041304", "zh-CN",
            topic_kind="task", source_format="md",
        )
        self.assertIn("<ul>", xml)
        self.assertLess(xml.index("<ol>"), xml.index("<ul>"))
        self.assertLess(xml.index("如需退出"), xml.index("</ol>"))
        self.assertGreater(xml.index("</ul>"), xml.index("如需退出"))
        self.assertLess(xml.index("</ul>"), xml.rindex("</ol>"))
        self.assertLess(xml.index("如需退出"), xml.index("开始测序。"))

    def test_horizontal_rule_is_dropped(self):
        xml = _content_to_dita_xml(
            "说明",
            "上一段。\n\n---\n\n下一段。\n",
            "concept",
            "DTC041305",
            "zh-CN",
            topic_kind="concept",
            source_format="md",
        )
        self.assertIn("上一段。", xml)
        self.assertIn("下一段。", xml)
        self.assertNotIn("<p>---</p>", xml)

    def test_normalize_skips_fenced_bold(self):
        md = "```\n**not bold**\n```\n"
        normalized = _normalize_native_markdown(md)
        self.assertIn("**not bold**", normalized)
        self.assertNotIn("[[B]]", normalized)

    def test_indented_fenced_code_stays_codeblock(self):
        md = (
            "1. 执行安装命令\n"
            "   ```bash\n"
            "   cd /path/to/script_dir\n"
            "   bash zarc_install.sh -P /opt/\n"
            "   ```\n"
        )
        xml = _content_to_dita_xml(
            "安装步骤", md, "task", "DTO041306", "zh-CN",
            topic_kind="task", source_format="md",
        )
        self.assertIn('<codeblock outputclass="bash" xml:space="preserve">', xml)
        self.assertIn("bash zarc_install.sh -P /opt/", xml)
        self.assertNotIn("<p>cd /path/to/script_dir</p>", xml)
        self.assertNotIn("[[B]]", _normalize_native_markdown("   ```\n   **not bold**\n   ```\n"))

    def test_hash_comments_inside_fence_are_not_headings(self):
        md = (
            "## 压缩示例\n\n"
            "**使用自定义参考基因组**\n"
            "```bash\n"
            "# 先为自定义参考基因组建索引\n"
            "zarc -i -r /path/to/hg19.fa\n"
            "\n"
            "# 再进行压缩\n"
            "zarc -c -t 16 -1 sample.fq -r /path/to/hg19.fa -o sample.arc\n"
            "```\n"
        )
        sections = _parse_content("md", None, md)
        titles = [sec["title"] for sec in sections]
        self.assertEqual(titles, ["压缩示例"])
        xml = _content_to_dita_xml(
            "压缩示例", sections[0]["content"], "concept", "DTC041307", "zh-CN",
            topic_kind="concept", source_format="md",
        )
        self.assertIn("# 先为自定义参考基因组建索引", xml)
        self.assertIn("zarc -i -r /path/to/hg19.fa", xml)
        self.assertIn("zarc -c -t 16 -1 sample.fq", xml)

    def test_english_markdown_bold_stays_bold_in_steps(self):
        xml = _content_to_dita_xml(
            "Installation Steps",
            "1. **Place the installation script**\n   Copy `zarc_install.sh`.\n",
            "task",
            "DTO041308",
            "en-US",
            topic_kind="task",
            source_format="md",
        )
        self.assertIn("<b>Place the installation script</b>", xml)
        self.assertNotIn("<uicontrol>Place the installation script</uicontrol>", xml)

    def test_unindented_paragraph_closes_native_markdown_list(self):
        md = (
            "试用版限制如下：\n"
            "- 上限 15 TB\n"
            "- 有效期 90 天\n"
            "\n"
            "若已购买正式授权，请按以下步骤激活：\n"
            "\n"
            "1. 申请许可证文件\n"
        )
        xml = _content_to_dita_xml(
            "软件激活", md, "task", "DTO041309", "zh-CN",
            topic_kind="task", source_format="md",
        )
        last_li = xml.split("<ul>", 1)[1].split("</ul>", 1)[0]
        self.assertNotIn("若已购买正式授权", last_li)
        self.assertLess(xml.index("</ul>"), xml.index("若已购买正式授权"))
        self.assertLess(xml.index("若已购买正式授权"), xml.index("<ol>"))

    def test_indented_paragraph_stays_inside_native_markdown_list_item(self):
        md = (
            "1. **申请许可证文件**\n"
            "   联系技术支持。\n"
        )
        xml = _content_to_dita_xml(
            "软件激活", md, "task", "DTO041310", "zh-CN",
            topic_kind="task", source_format="md",
        )
        first_li = xml.split("<li>", 1)[1].split("</li>", 1)[0]
        self.assertIn("申请许可证文件", first_li)
        self.assertIn("联系技术支持。", first_li)

    def test_markdown_section_reference_becomes_xref(self):
        xml = _content_to_dita_xml(
            "常见问题",
            "手动安装对应依赖包（参见 3.2 节第 5 点）。\n",
            "concept",
            "DTC041311",
            "zh-CN",
            topic_kind="concept",
            source_format="md",
            xref_map={"3.2": "DTO041007", "DTO041007": "安装步骤"},
        )
        self.assertIn(
            '<xref href="DTO041007.dita" cms:title="安装步骤">安装步骤</xref>',
            xml,
        )
        self.assertIn("第 5 点", xml)

    def test_english_section_reference_becomes_xref(self):
        xml = _content_to_dita_xml(
            "FAQ",
            "Manually install the packages (see Section 3.2, step 5).\n",
            "concept",
            "DTC041312",
            "en-US",
            topic_kind="concept",
            source_format="md",
            xref_map={"3.2": "DTO041007", "DTO041007": "Installation steps"},
        )
        self.assertIn(
            '<xref href="DTO041007.dita" cms:title="Installation steps">Installation steps</xref>',
            xml,
        )
        self.assertIn("step 5", xml)


if __name__ == "__main__":
    unittest.main()
