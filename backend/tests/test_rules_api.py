"""规则库导入/导出接口回归。

覆盖两类已修复的问题：
1. static 路径（/export、/import-template、/import）必须注册在 /{rule_id} 之前，
   否则会被当作 rule_id 解析失败；
2. 规则库导入导出统一走 Excel：导出不因默认分页截断且带 severity/language，
   导入按行校验必填列、重复 rule_no 跳过而不是报错。
"""
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.api.auth import require_admin
from app.api.rules import RULE_EXPORT_COLUMNS, XLSX_MEDIA_TYPE
from app.database import Base, get_db
from app.main import app
from app.models.rule import Rule as RuleModel

COLUMN_TITLES = [title for title, _field, _width in RULE_EXPORT_COLUMNS]
COLUMN_INDEX = {field: index for index, (_title, field, _width) in enumerate(RULE_EXPORT_COLUMNS)}


@pytest.fixture()
def client():
    # StaticPool：TestClient 的请求跑在独立线程，内存库需共享同一连接，
    # 否则每个线程会各自新建一个空库。
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    session_local = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    Base.metadata.create_all(bind=engine)

    def override_get_db():
        db = session_local()
        try:
            yield db
        finally:
            db.close()

    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[require_admin] = lambda: {"username": "admin", "role": "admin"}
    try:
        yield TestClient(app), session_local
    finally:
        app.dependency_overrides.clear()


def _build_xlsx(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(COLUMN_TITLES)
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _sheet_rows(payload):
    workbook = load_workbook(BytesIO(payload))
    sheet = workbook.worksheets[0]
    return [[cell for cell in row] for row in sheet.iter_rows(values_only=True)]


def test_static_rule_routes_are_registered_before_dynamic_rule_id():
    paths = [route.path for route in app.routes if getattr(route, "path", "").startswith("/api/rules")]
    dynamic = paths.index("/api/rules/{rule_id}")
    for static in ("/api/rules/export", "/api/rules/import-template", "/api/rules/import"):
        assert paths.index(static) < dynamic


def test_export_rules_returns_xlsx_with_severity_and_language(client):
    test_client, session_local = client
    db = session_local()
    try:
        db.add(RuleModel(
            rule_no="EXT-T001", category="拼写", description="示例规则", regex=r"\bteh\b",
            example="teh", suggestion="the", audit_basis="依据",
            severity="serious", language="en",
        ))
        db.commit()
    finally:
        db.close()

    response = test_client.get("/api/rules/export")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(XLSX_MEDIA_TYPE)
    header, first = _sheet_rows(response.content)[:2]
    assert header[0] == "规则编号"
    assert first[0] == "EXT-T001"
    # 导出把枚举值转成页面上显示的中文标签，便于直接回填维护。
    assert first[COLUMN_INDEX["severity"]] == "严重"
    assert first[COLUMN_INDEX["language"]] == "英文"


def test_export_rules_is_not_truncated_by_default_page_limit(client):
    test_client, session_local = client
    db = session_local()
    try:
        db.add_all([
            RuleModel(rule_no=f"EXT-B{i:03d}", category="拼写", description="d", regex="r")
            for i in range(120)
        ])
        db.commit()
    finally:
        db.close()

    response = test_client.get("/api/rules/export")

    assert response.status_code == 200
    # 表头 + 120 条数据
    assert len(_sheet_rows(response.content)) == 121


def test_import_template_has_data_sheet_and_filling_guide(client):
    test_client, _ = client

    response = test_client.get("/api/rules/import-template")

    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == ["规则库", "填写说明"]
    assert [cell.value for cell in workbook["规则库"][1]] == COLUMN_TITLES
    assert workbook["填写说明"]["A1"].value == "列名"


def test_import_rules_creates_new_and_skips_existing(client):
    test_client, session_local = client
    db = session_local()
    try:
        db.add(RuleModel(rule_no="R-DUP", category="拼写", description="已存在", regex="r"))
        db.commit()
    finally:
        db.close()

    payload = _build_xlsx([
        ["R-NEW", "拼写", "新增", "r", "", "", "依据", "致命", "中文"],
        ["R-DUP", "拼写", "库中已存在", "r", "", "", "", "", ""],
    ])

    response = test_client.post(
        "/api/rules/import",
        files={"file": ("rules.xlsx", payload, XLSX_MEDIA_TYPE)},
    )

    assert response.status_code == 200
    body = response.json()
    assert (body["created"], body["duplicates"], body["errors"]) == (1, 1, [])

    exported = _sheet_rows(test_client.get("/api/rules/export").content)
    assert {row[0] for row in exported[1:]} == {"R-DUP", "R-NEW"}
    new_row = next(row for row in exported[1:] if row[0] == "R-NEW")
    # 导出显示中文标签（页面阅读用）
    assert new_row[COLUMN_INDEX["severity"]] == "致命"
    assert new_row[COLUMN_INDEX["language"]] == "中文"
    # DB 必须存枚举值，否则严重级判断与语言过滤全部失效（P1 回归防护）
    db = session_local()
    try:
        stored = db.query(RuleModel).filter(RuleModel.rule_no == "R-NEW").one()
        assert stored.severity == "fatal"
        assert stored.language == "cn"
    finally:
        db.close()


def test_import_rules_reports_row_errors_without_aborting_valid_rows(client):
    test_client, _ = client

    payload = _build_xlsx([
        ["R-OK", "拼写", "合法", "r", "", "", "", "", ""],
        ["", "拼写", "缺少规则编号", "r", "", "", "", "", ""],
        ["R-BAD", "拼写", "严重程度非法", "r", "", "", "", "很重要", ""],
        ["R-MISSING", "拼写", "", "", "", "", "", "", ""],
    ])

    response = test_client.post(
        "/api/rules/import",
        files={"file": ("rules.xlsx", payload, XLSX_MEDIA_TYPE)},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["created"] == 1
    assert {error["row"] for error in body["errors"]} == {3, 4, 5}
    assert "严重程度" in next(error["message"] for error in body["errors"] if error["row"] == 4)
    assert "规则描述" in next(error["message"] for error in body["errors"] if error["row"] == 5)


def test_import_rules_rejects_missing_required_column(client):
    test_client, _ = client
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["规则编号", "分类", "规则描述"])  # 缺「正则」列
    sheet.append(["R-1", "拼写", "d"])
    buffer = BytesIO()
    workbook.save(buffer)

    response = test_client.post(
        "/api/rules/import",
        files={"file": ("rules.xlsx", buffer.getvalue(), XLSX_MEDIA_TYPE)},
    )

    assert response.status_code == 400
    assert "正则" in response.json()["detail"]


def test_import_rules_rejects_non_xlsx_file(client):
    test_client, _ = client

    response = test_client.post(
        "/api/rules/import",
        files={"file": ("rules.json", b'[{"rule_no":"R1"}]', "application/json")},
    )

    assert response.status_code == 400
    assert "xlsx" in response.json()["detail"]


def test_dynamic_rule_id_route_still_resolves_after_reordering(client):
    test_client, session_local = client
    db = session_local()
    try:
        rule = RuleModel(
            rule_no="EXT-T002", category="拼写", description="d", regex="r",
            example="", suggestion="", audit_basis="",
        )
        db.add(rule)
        db.commit()
        rule_id = rule.id
    finally:
        db.close()

    response = test_client.get(f"/api/rules/{rule_id}")

    assert response.status_code == 200
    assert response.json()["rule_no"] == "EXT-T002"


def test_bulk_json_endpoint_still_creates_rules_without_duplicates(client):
    test_client, session_local = client
    db = session_local()
    try:
        db.add(RuleModel(rule_no="R-DUP", category="拼写", description="已存在", regex="r"))
        db.commit()
    finally:
        db.close()

    payload = [
        {"rule_no": "R-NEW", "category": "拼写", "description": "新增", "regex": "r"},
        {"rule_no": "R-DUP", "category": "拼写", "description": "库中已存在", "regex": "r"},
        {"rule_no": "R-NEW", "category": "拼写", "description": "同批重复", "regex": "r"},
    ]

    response = test_client.post("/api/rules/bulk", json=payload)

    assert response.status_code == 200
    assert response.json()["created"] == 1
    assert response.json()["total"] == 3


def test_normalize_choice_accepts_chinese_labels_english_keys_and_uppercase():
    from app.api.rules import _normalize_choice, SEVERITY_LABELS, LANGUAGE_LABELS

    cases = [
        # (输入, labels, 期望)
        ("致命", SEVERITY_LABELS, "fatal"),
        ("严重", SEVERITY_LABELS, "serious"),
        ("一般", SEVERITY_LABELS, "general"),
        ("建议", SEVERITY_LABELS, "suggestion"),
        ("fatal", SEVERITY_LABELS, "fatal"),
        ("FATAL", SEVERITY_LABELS, "fatal"),
        ("Serious", SEVERITY_LABELS, "serious"),
        ("中文", LANGUAGE_LABELS, "cn"),
        ("英文", LANGUAGE_LABELS, "en"),
        ("中英通用", LANGUAGE_LABELS, "both"),
        ("cn", LANGUAGE_LABELS, "cn"),
        ("CN", LANGUAGE_LABELS, "cn"),
        ("", SEVERITY_LABELS, "general"),
        (None, SEVERITY_LABELS, "general"),
        ("很重要", SEVERITY_LABELS, None),
        ("EN-CN", LANGUAGE_LABELS, None),
    ]
    for raw, labels, expected in cases:
        default = "general" if labels is SEVERITY_LABELS else "both"
        assert _normalize_choice(raw, labels, default) == expected, f"raw={raw!r}"

    # 默认值回落
    assert _normalize_choice("", SEVERITY_LABELS, "suggestion") == "suggestion"


def test_import_export_roundtrip_keeps_enum_values(client):
    """导出（中文标签文件）直接回导，DB 枚举值必须保持不变。"""
    test_client, session_local = client
    db = session_local()
    try:
        db.add(RuleModel(
            rule_no="RT-001", category="拼写", description="回灌测试", regex=r"\bteh\b",
            example="teh", suggestion="the", audit_basis="依据",
            severity="serious", language="en",
        ))
        db.commit()
    finally:
        db.close()

    # 1. 导出 -> 中文标签 xlsx
    exported = test_client.get("/api/rules/export")
    assert exported.status_code == 200

    # 2. 原样回导
    response = test_client.post(
        "/api/rules/import",
        files={"file": ("roundtrip.xlsx", exported.content, XLSX_MEDIA_TYPE)},
    )
    assert response.status_code == 200
    body = response.json()
    assert body["duplicates"] == 1  # 同 rule_no 应全部跳过，不产生污染副本

    # 3. DB 仍为枚举值
    db = session_local()
    try:
        stored = db.query(RuleModel).filter(RuleModel.rule_no == "RT-001").one()
        assert stored.severity == "serious"
        assert stored.language == "en"
    finally:
        db.close()
