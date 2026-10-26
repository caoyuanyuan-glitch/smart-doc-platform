import os
import tempfile

from app.utils.param_compare import (
    compare_params,
    extract_params,
    _values_consistent,
    resolve_param_targets,
    run_param_compare,
    select_target_params,
)


def test_compare_params_copies_chapter_fields():
    params_a = [
        {"name": "重量", "value": "15 kg", "source": "text", "chapter": "技术规格"},
        {"name": "功率", "value": "200 W", "source": "text", "chapter": "电气参数"},
    ]
    params_b = [
        {"name": "重量", "value": "16 kg", "source": "text", "chapter": "主要技术参数"},
        {"name": "噪音", "value": "50 dB", "source": "text", "chapter": "操作环境"},
    ]
    rows = compare_params(params_a, params_b)
    by_name = {r["param"]: r for r in rows}
    assert by_name["重量"]["chapter_a"] == "技术规格"
    assert by_name["重量"]["chapter_b"] == "主要技术参数"
    assert by_name["功率"]["match"] == "仅A有"
    assert by_name["功率"]["chapter_a"] == "电气参数"
    assert by_name["功率"]["chapter_b"] == ""
    assert by_name["噪音"]["match"] == "仅B有"
    assert by_name["噪音"]["chapter_a"] == ""
    assert by_name["噪音"]["chapter_b"] == "操作环境"


def test_markdown_params_attach_heading():
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "spec.md")
        with open(path, "w", encoding="utf-8") as f:
            f.write("# 技术规格\n\n重量：15 kg\n尺寸：300 mm\n\n# 操作环境\n\n工作温度：10 ℃\n")
        params = extract_params(path)
        by_name = {p["name"]: p for p in params}
        assert by_name["重量"]["chapter"] == "技术规格"
        assert by_name["尺寸"]["chapter"] == "技术规格"
        assert by_name["工作温度"]["chapter"] == "操作环境"


def test_xlsx_params_use_sheet_as_chapter():
    from openpyxl import Workbook

    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "spec.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "电气参数"
        ws["A1"] = "参数名称"
        ws["B1"] = "数值"
        ws["A2"] = "功率"
        ws["B2"] = "200 W"
        wb.save(path)
        params = extract_params(path)
        assert params
        assert params[0]["chapter"] == "电气参数"


def test_pdf_params_use_bookmark_title():
    import fitz

    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "spec.pdf")
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), "Weight: 15 kg")
        doc.set_toc([[1, "技术规格", 1]])
        doc.save(path)
        doc.close()
        params = extract_params(path)
        by_name = {p["name"]: p for p in params}
        assert "Weight" in by_name
        assert by_name["Weight"]["chapter"] == "技术规格"


def test_default_targets_drop_noise_keep_spec_names():
    with tempfile.TemporaryDirectory() as tmp:
        a = os.path.join(tmp, "a.md")
        b = os.path.join(tmp, "b.md")
        with open(a, "w", encoding="utf-8") as f:
            f.write(
                "# 技术规格\n"
                "外形尺寸：400 mm×300 mm×200 mm\n"
                "净重：15 kg\n"
                "工作温度：10 ℃~30 ℃\n"
                "湿度：20%~80%\n"
                "大气压力范围：70 kPa~106 kPa\n"
                "电源电压：100-240 V\n"
                "额定功率：200 W\n"
                "LCD尺寸：12.1 英寸\n"
                "分辨率：1280×800 像素\n"
                "基因测序仪：DNBSEQ-E25\n"
                "y 网口：用于主机通信\n"
            )
        with open(b, "w", encoding="utf-8") as f:
            f.write(
                "# 技术规格\n"
                "外形尺寸：400 mm×300 mm×200 mm\n"
                "净重：18 kg\n"
                "工作温度：10 ℃~30 ℃\n"
                "湿度：20%~80%\n"
                "大气压力：70 kPa~106 kPa\n"
                "电源电压：100-240 V\n"
                "额定功率：200 W\n"
                "LCD尺寸：12.1 英寸\n"
                "分辨率：1280×800 像素\n"
            )
        result = run_param_compare(a, b)
        names = {row["param"] for row in result["results"]}
        assert "重量" in names
        assert "仪器尺寸" in names
        assert "分辨率" in names
        assert "LCD尺寸" in names
        assert "基因测序仪" not in names
        assert "y 网口" not in names
        by_name = {row["param"]: row for row in result["results"]}
        assert by_name["重量"]["match"] == "不一致"
        assert by_name["重量"]["value_a"] == "15 kg"
        assert by_name["重量"]["value_b"] == "18 kg"
        assert [row["param"] for row in result["results"]] == [
            "预期用途", "适用范围", "结构组成", "符合标准清单", "产品性能-技术指标",
            "仪器尺寸", "重量", "操作/运行环境温度", "湿度", "大气压",
            "电源电压", "额定功率", "LCD尺寸", "分辨率",
        ]


def test_custom_param_names_only():
    params = [
        {"name": "净重", "value": "15 kg", "source": "text", "chapter": "技术规格"},
        {"name": "基因测序仪", "value": "E25", "source": "text", "chapter": ""},
    ]
    targets = resolve_param_targets("重量")
    selected = select_target_params(params, "净重：15 kg", targets)
    assert [p["name"] for p in selected] == ["重量"]
    assert selected[0]["value"] == "15 kg"


def test_run_param_compare_respects_param_names():
    with tempfile.TemporaryDirectory() as tmp:
        a = os.path.join(tmp, "a.md")
        b = os.path.join(tmp, "b.md")
        with open(a, "w", encoding="utf-8") as f:
            f.write(
                "# 技术规格\n"
                "外形尺寸：348 mm×312 mm×257 mm\n"
                "净重：15 kg\n"
                "工作温度：10 ℃~30 ℃\n"
            )
        with open(b, "w", encoding="utf-8") as f:
            f.write(
                "# 技术规格\n"
                "外形尺寸：400 mm×300 mm×200 mm\n"
                "净重：18 kg\n"
                "工作温度：10 ℃~30 ℃\n"
            )
        result = run_param_compare(a, b, param_names="重量")
        names = [row["param"] for row in result["results"]]
        assert names == ["重量"]
        assert result["param_names"] == ["重量"]
        assert result["results"][0]["match"] == "不一致"


def test_consistency_rules_ignore_format_noise():
    dim_a = "348 mm ( 长 )×312 mm ( 宽 )×257 mm ( 高)"
    dim_b = "348 mm（长）×312 mm（宽）×257 mm（高）"
    assert _values_consistent("仪器尺寸", dim_a, dim_b)

    pressure_a = "70 kPa~106 kPa"
    pressure_b = "70 kPa~106 kPa（海拔不超过3000 米）"
    assert not _values_consistent("大气压", pressure_a, pressure_b)

    assert _values_consistent("操作/运行环境温度", "20°C - 25°C", "18°C - 26°C")
    assert not _values_consistent("额定功率", "1,500 W", "1,800 W")
    assert not _values_consistent("重量", "15 kg", "18 kg")


def test_matrix_keeps_all_requested_targets():
    with tempfile.TemporaryDirectory() as tmp:
        a = os.path.join(tmp, "a.md")
        b = os.path.join(tmp, "b.md")
        with open(a, "w", encoding="utf-8") as f:
            f.write("# specs\nweight: 15 kg\n")
        with open(b, "w", encoding="utf-8") as f:
            f.write("# specs\nweight: 15 kg\n")
        result = run_param_compare(a, b)
        names = [row["param"] for row in result["results"]]
        assert names == [
            "预期用途", "适用范围", "结构组成", "符合标准清单", "产品性能-技术指标",
            "仪器尺寸", "重量", "操作/运行环境温度", "湿度", "大气压",
            "电源电压", "额定功率", "LCD尺寸", "分辨率",
        ]
        assert all(row["match"] == "缺失" for row in result["results"])


def test_short_spec_labels_map_to_canonical_targets():
    with tempfile.TemporaryDirectory() as tmp:
        a = os.path.join(tmp, "a.md")
        b = os.path.join(tmp, "b.md")
        body = (
            "# 技术规格\n"
            "外形尺寸：348 mm×312 mm×257 mm\n"
            "重量：约15 kg\n"
            "温度：15 ℃~30 ℃\n"
            "湿度：20%~80%，无冷凝\n"
            "大气压：70 kPa~106 kPa\n"
            "电源：AC 100-240 V，50/60 Hz\n"
            "功率：200 W\n"
            "LCD：12.1 英寸\n"
            "分辨率：1280×800 像素\n"
        )
        with open(a, "w", encoding="utf-8") as f:
            f.write(body)
        with open(b, "w", encoding="utf-8") as f:
            f.write(body)
        result = run_param_compare(a, b)
        by_name = {row["param"]: row for row in result["results"]}
        assert len(result["results"]) == 14
        assert by_name["操作/运行环境温度"]["value_a"] == "15 ℃~30 ℃"
        assert by_name["电源电压"]["value_a"].startswith("AC 100-240 V")
        assert by_name["额定功率"]["value_a"] == "200 W"
        assert by_name["LCD尺寸"]["value_a"] == "12.1 英寸"
        assert by_name["操作/运行环境温度"]["match"] == "一致"
        assert by_name["操作/运行环境温度"]["chapter_a"] == "技术规格"
        assert by_name["适用范围"]["match"] == "缺失"


def test_rejects_lcd_type_and_voltage_sentence():
    with tempfile.TemporaryDirectory() as tmp:
        a = os.path.join(tmp, "a.md")
        b = os.path.join(tmp, "b.md")
        body = (
            "# 服务计划\n"
            "LCD：类型\n"
            "# 电源要求\n"
            "电压：不应超过标称电压的±10%\n"
            "# 技术规格\n"
            "LCD尺寸：12.1 英寸\n"
            "电源电压：AC 100-240 V\n"
            "额定功率：350 W\n"
        )
        with open(a, "w", encoding="utf-8") as f:
            f.write(body)
        with open(b, "w", encoding="utf-8") as f:
            f.write(body)
        result = run_param_compare(a, b)
        by_name = {row["param"]: row for row in result["results"]}
        assert by_name["LCD尺寸"]["value_a"] == "12.1 英寸"
        assert by_name["LCD尺寸"]["chapter_a"] == "技术规格"
        assert by_name["电源电压"]["value_a"] == "AC 100-240 V"
        assert by_name["额定功率"]["value_a"] == "350 W"


def test_section_targets_and_generic_chapter_skipped():
    with tempfile.TemporaryDirectory() as tmp:
        a = os.path.join(tmp, "a.md")
        b = os.path.join(tmp, "b.md")
        body = (
            "# 说明书\n"
            "封面文字\n"
            "# 预期用途\n"
            "用于高通量基因测序。\n"
            "# 结构组成\n"
            "主机、电源适配器、数据线。\n"
            "# 符合标准\n"
            "YY 0648、GB 4793.1\n"
            "# 技术规格\n"
            "温度：15 ℃~30 ℃\n"
        )
        with open(a, "w", encoding="utf-8") as f:
            f.write(body)
        with open(b, "w", encoding="utf-8") as f:
            f.write(body)
        result = run_param_compare(a, b)
        by_name = {row["param"]: row for row in result["results"]}
        assert "用于高通量基因测序" in by_name["预期用途"]["value_a"]
        assert by_name["预期用途"]["chapter_a"] == "预期用途"
        assert "主机" in by_name["结构组成"]["value_a"]
        assert "YY 0648" in by_name["符合标准清单"]["value_a"]
        assert by_name["操作/运行环境温度"]["chapter_a"] == "技术规格"
        assert by_name["操作/运行环境温度"]["chapter_a"] != "说明书"


def test_section_values_use_body_not_toc_location():
    with tempfile.TemporaryDirectory() as tmp:
        a = os.path.join(tmp, "a.md")
        b = os.path.join(tmp, "b.md")
        body = (
            "# 目录\n"
            "1.1 预期用途 ........ 2\n"
            "1.2 适用范围 ........ 3\n"
            "1.3 结构组成 ........ 4\n"
            "1.4 符合标准 ........ 5\n"
            "# 1.1 预期用途\n"
            "用于高通量基因测序检测。\n"
            "# 1.2 适用范围\n"
            "适用于医疗机构的临床基因检测实验室。\n"
            "# 1.3 结构组成\n"
            "主机、电源适配器、数据线。\n"
            "# 1.4 符合标准\n"
            "YY 0648、GB 4793.1\n"
        )
        with open(a, "w", encoding="utf-8") as f:
            f.write(body)
        with open(b, "w", encoding="utf-8") as f:
            f.write(body)
        result = run_param_compare(a, b)
        by_name = {row["param"]: row for row in result["results"]}
        assert "用于高通量基因测序检测" in by_name["预期用途"]["value_a"]
        assert "1.1" not in by_name["预期用途"]["value_a"]
        assert "临床基因检测" in by_name["适用范围"]["value_a"]
        assert "1.2" not in by_name["适用范围"]["value_a"]
        assert "主机" in by_name["结构组成"]["value_a"]
        assert "1.3" not in by_name["结构组成"]["value_a"]
        assert "YY 0648" in by_name["符合标准清单"]["value_a"]
        assert by_name["适用范围"]["match"] == "一致"


def test_section_inline_heading_keeps_body():
    text_a = "1.2 适用范围 适用于人全血样本的基因检测。\n1.3 结构组成 主机和电源适配器。\n"
    text_b = "1.2 适用范围 适用于人全血样本的基因检测。\n1.3 结构组成 主机和电源适配器。\n"
    with tempfile.TemporaryDirectory() as tmp:
        a = os.path.join(tmp, "a.md")
        b = os.path.join(tmp, "b.md")
        with open(a, "w", encoding="utf-8") as f:
            f.write(text_a)
        with open(b, "w", encoding="utf-8") as f:
            f.write(text_b)
        result = run_param_compare(a, b)
        by_name = {row["param"]: row for row in result["results"]}
        assert "人全血样本" in by_name["适用范围"]["value_a"]
        assert by_name["适用范围"]["value_a"] != "1.2 适用范围"
        assert "主机" in by_name["结构组成"]["value_a"]


def test_standards_list_skips_toc_and_keeps_gb_body():
    with tempfile.TemporaryDirectory() as tmp:
        a = os.path.join(tmp, "a.md")
        b = os.path.join(tmp, "b.md")
        body = (
            "# 目录\n"
            "符合标准清单\n"
            "产品性能-技术指标\n"
            "制造商信息\n"
            "产品中有害物质的名称及含量\n"
            "索引\n"
            "# 符合标准清单\n"
            "本仪器符合以下标准：\n"
            "GB/T 18268.1\n"
            "YY 0648\n"
            "# 产品性能-技术指标\n"
            "测序通量：500 M reads\n"
            "Q30：>85%\n"
            "# 制造商信息\n"
            "华大智造\n"
        )
        with open(a, "w", encoding="utf-8") as f:
            f.write(body)
        with open(b, "w", encoding="utf-8") as f:
            f.write(body.replace("YY 0648", "YY 0793").replace("500 M", "800 M"))
        result = run_param_compare(a, b)
        by_name = {row["param"]: row for row in result["results"]}
        assert "YY 0648" in by_name["符合标准清单"]["value_a"]
        assert "GB/T 18268.1" in by_name["符合标准清单"]["value_a"]
        assert "制造商信息" not in by_name["符合标准清单"]["value_a"]
        assert "索引" not in by_name["符合标准清单"]["value_a"]
        assert "测序通量" in by_name["产品性能-技术指标"]["value_a"]
        assert "500 M" in by_name["产品性能-技术指标"]["value_a"]
        assert "800 M" in by_name["产品性能-技术指标"]["value_b"]
        assert by_name["产品性能-技术指标"]["match"] == "不一致"
        assert by_name["符合标准清单"]["match"] == "不一致"


def test_old_standards_label_resolves():
    targets = resolve_param_targets("符合法规标准清单")
    assert [t["label"] for t in targets] == ["符合标准清单"]


def test_standards_list_rejects_concatenated_toc():
    with tempfile.TemporaryDirectory() as tmp:
        a = os.path.join(tmp, "a.md")
        b = os.path.join(tmp, "b.md")
        body = (
            "符合标准清单 产品性能-技术指标 制造商信息 产品中有害物质的名称及含量 索引\n"
            "本仪器符合以下标准：\n"
            "GB/T 18268.1\n"
            "YY 0648\n"
        )
        with open(a, "w", encoding="utf-8") as f:
            f.write(body)
        with open(b, "w", encoding="utf-8") as f:
            f.write(body)
        result = run_param_compare(a, b)
        by_name = {row["param"]: row for row in result["results"]}
        assert "YY 0648" in by_name["符合标准清单"]["value_a"]
        assert "制造商信息" not in by_name["符合标准清单"]["value_a"]
        assert "索引" not in by_name["符合标准清单"]["value_a"]


def test_standards_table_diff_not_hidden_by_shared_prefix():
    intro = "# 符合标准清单\n本仪器符合以下标准：\n\n"
    rest_a = "1 GB 4793.1 安全要求\n2 YY 0648 体外诊断医疗设备\n# 制造商信息\n华大智造\n"
    rest_b = "1 GB 4793.1 安全要求\n2 YY 0793 医用电气设备\n# 制造商信息\n华大智造\n"
    with tempfile.TemporaryDirectory() as tmp:
        a = os.path.join(tmp, "a.md")
        b = os.path.join(tmp, "b.md")
        with open(a, "w", encoding="utf-8") as f:
            f.write(intro + rest_a)
        with open(b, "w", encoding="utf-8") as f:
            f.write(intro + rest_b)
        result = run_param_compare(a, b)
        by_name = {row["param"]: row for row in result["results"]}
        assert "YY 0648" in by_name["符合标准清单"]["value_a"]
        assert "YY 0793" in by_name["符合标准清单"]["value_b"]
        assert by_name["符合标准清单"]["match"] == "不一致"


def test_standards_same_ids_different_text_are_inconsistent():
    assert not _values_consistent(
        "符合标准清单",
        "本仪器符合以下标准：GB 4793.1 安全要求",
        "本仪器符合以下标准：GB 4793.1 实验室设备安全",
    )
    assert not _values_consistent(
        "符合标准清单",
        "本仪器符合以下标准：GB 4793.1",
        "本仪器符合以下标准：GB 4793.1 YY 0648 体外诊断",
    )
    assert _values_consistent(
        "符合标准清单",
        "本仪器符合以下标准：GB 4793.1、YY 0648",
        "本仪器符合以下标准： GB 4793.1、 YY 0648",
    )


def test_structure_section_stops_at_blank_and_drops_duplicate():
    body = (
        "# 结构组成\n"
        "本产品由测序仪主机、电源适配器、数据线组成。\n"
        "本产品由测序仪主机、电源适配器、数据线组成。\n"
        "\n"
        "后续安装步骤很长很长，不应进入结构组成。\n"
        "# 技术规格\n"
        "重量：15 kg\n"
    )
    with tempfile.TemporaryDirectory() as tmp:
        a = os.path.join(tmp, "a.md")
        b = os.path.join(tmp, "b.md")
        with open(a, "w", encoding="utf-8") as f:
            f.write(body)
        with open(b, "w", encoding="utf-8") as f:
            f.write(body)
        result = run_param_compare(a, b)
        by_name = {row["param"]: row for row in result["results"]}
        val = by_name["结构组成"]["value_a"]
        assert "测序仪主机" in val
        assert val.count("测序仪主机") == 1
        assert "安装步骤" not in val
        assert "15 kg" not in val


def test_scope_section_clips_working_principle():
    body = (
        "# 适用范围\n"
        "该产品在临床上用于对来源于人体样本的DNA进行测序。禁忌症：无。"
        "工作原理 基因测序仪采用先进的DNA纳米球以及联合探针锚定聚合测序核心技术，"
        "制备DNB示意图如下。\n"
        "# 结构组成\n"
        "主机和电源适配器。\n"
    )
    with tempfile.TemporaryDirectory() as tmp:
        a = os.path.join(tmp, "a.md")
        b = os.path.join(tmp, "b.md")
        with open(a, "w", encoding="utf-8") as f:
            f.write(body)
        with open(b, "w", encoding="utf-8") as f:
            f.write(body)
        result = run_param_compare(a, b)
        by_name = {row["param"]: row for row in result["results"]}
        val = by_name["适用范围"]["value_a"]
        assert "人体样本" in val
        assert "工作原理" not in val
        assert "纳米球" not in val


def test_structure_section_formats_module_table():
    body = (
        "# 结构组成\n"
        "基因测序仪由主机、计算模块和控制软件及耗材组成。各部件功能说明如下："
        "部件 功能 主机 主体模块 负责试剂储存、样本加载、数据采集等。"
        "外壳模块 负责对主机起到固定支撑、遮蔽光线的作用。"
        "显示模块 负责显示器的控制界面和相关信息以及用户交互。"
        "电源模块 负责主机电力供应，为系统及模块供电。"
        "控制模块 负责集成系统控制，完成与主体模块的通讯及数据储存。"
        "计算模块 负责处理原始数据，计算获得碱基序列。\n"
        "# 技术规格\n"
        "重量：15 kg\n"
    )
    with tempfile.TemporaryDirectory() as tmp:
        a = os.path.join(tmp, "a.md")
        b = os.path.join(tmp, "b.md")
        with open(a, "w", encoding="utf-8") as f:
            f.write(body)
        with open(b, "w", encoding="utf-8") as f:
            f.write(body.replace("试剂储存", "试剂冷藏"))
        result = run_param_compare(a, b)
        by_name = {row["param"]: row for row in result["results"]}
        val = by_name["结构组成"]["value_a"]
        assert "主体模块 | " in val
        assert "计算模块 | " in val
        assert "工作原理" not in val
        assert by_name["结构组成"]["match"] == "不一致"
        assert "完成与主体模块的通讯" in val


def test_structure_keeps_modules_after_page_break():
    body_a = (
        "# 结构组成\n"
        "基因测序仪由主机、计算模块和控制软件（软件版本V1）及耗材组成。"
        "耗材为集成测序载片（E25 FCL）。\n"
        "各部件功能说明如下：\n"
        "主体模块 负责试剂储存、样本加载、数据采集等。\n"
        "外壳模块 负责对主机起到固定支撑、遮蔽光线的作用。\n"
        "显示模块 负责显示仪器的控制界面和相关信息以及用户交互。\n"
        "电源模块 负责主机电力供应，为系统及模块供电。\n"
        "控制模块 负责集成系统控制，完成与主体模块的通讯及数据储存。\n"
        "计算模块 负责处理原始数据，计算获得碱基序列。\n"
        "控制软件\n"
        "y 控制功能：控制软硬件协同工作，完成测序工作。\n"
        "y 碱基识别。\n"
        "耗材 集成测序载片（E25 FCL）货号为940-000599-00。\n"
        "# 基本部件\n"
        "前视图\n"
    )
    body_b = (
        "# 结构组成\n"
        "基因测序仪由主机、计算模块和控制软件（软件版本：V2）及耗材组成。"
        "耗材为集成测序载片（E25 FCL、E25 FCU 和E25 FCS）。\n"
        "各部件功能说明如下：\n"
        "主体模块 负责试剂储存、样本加载、数据采集等。\n"
        "外壳模块 负责对主机起到固定支撑、遮蔽光线的作用。\n"
        "显示模块 负责显示仪器的控制界面和相关信息以及用户交互。\n"
        "电源模块 负责主机电力供应，为系统及模块供电。\n"
        "控制模块 负责集成系统控制，完成与主体模块的通讯及数据储存。\n"
        "\n"
        "---PAGE---\n"
        "10\n"
        "仪器概述\n"
        "说明书\n"
        "部件\n"
        "功能\n"
        "计算模块 负责处理原始数据，计算获得碱基序列。\n"
        "控制软件\n"
        "y 控制功能：控制软硬件协同工作，完成测序工作。\n"
        "y 碱基识别。\n"
        "耗材 集成测序载片（E25 FCL、E25 FCU 和E25 FCS）货号如下。\n"
        "# 基本部件\n"
        "前视图\n"
    )
    with tempfile.TemporaryDirectory() as tmp:
        a = os.path.join(tmp, "a.md")
        b = os.path.join(tmp, "b.md")
        with open(a, "w", encoding="utf-8") as f:
            f.write(body_a)
        with open(b, "w", encoding="utf-8") as f:
            f.write(body_b)
        result = run_param_compare(a, b)
        by_name = {row["param"]: row for row in result["results"]}
        va = by_name["结构组成"]["value_a"]
        vb = by_name["结构组成"]["value_b"]
        assert "软件版本V1" in va
        assert "软件版本：V2" in vb
        assert "计算模块 | " in va
        assert "计算模块 | " in vb
        assert "控制软件 | " in va
        assert "控制软件 | " in vb
        assert "E25 FCL" in va
        assert "E25 FCU" in vb
        assert "940-000599-00" in va
        assert "前视图" not in va
        assert by_name["结构组成"]["match"] == "不一致"


def test_standards_list_includes_safety_block_diff():
    head = (
        "# 符合标准清单\n"
        "本仪器符合以下标准：\n"
        "项目\n符合标准\n"
        "电磁兼容性\n"
        "y GB/T 18268.1-2010 测量、控制和实验室用的电设备 电磁兼容性要求 第1 部分：通用要求\n"
        "y GB/T 18268.26-2010 测量、控制和实验室用的电设备 电磁兼容性要求 第26 部分：特殊要求\n"
    )
    rest_a = (
        "安全要求\n"
        "y GB 4793.1-2007 测量、控制和实验室用电气设备的安全要求 第1 部分：通用要求\n"
        "y YY 0648-2008 体外诊断医用设备的专用要求\n"
        "---PAGE---\n"
        "电磁兼容性\n"
        "y GB/T 18268.1-2010 测量、控制和实验室用的电设备 电磁兼容性要求 第1 部分：通用要求\n"
        "y YY 0648-2008 体外诊断医用设备的专用要求\n"
        "# 产品性能-技术指标\n"
        "测序通量：500 M reads\n"
    )
    rest_b = (
        "安全要求\n"
        "y GB 4793-2024 测量、控制和实验室用电气设备安全技术规范\n"
        "y YY/T 0648-2025 体外诊断医用设备的专用要求\n"
        "# 产品性能-技术指标\n"
        "测序通量：800 M reads\n"
    )
    with tempfile.TemporaryDirectory() as tmp:
        a = os.path.join(tmp, "a.md")
        b = os.path.join(tmp, "b.md")
        with open(a, "w", encoding="utf-8") as f:
            f.write(head + rest_a)
        with open(b, "w", encoding="utf-8") as f:
            f.write(head + rest_b)
        result = run_param_compare(a, b)
        by_name = {row["param"]: row for row in result["results"]}
        va = by_name["符合标准清单"]["value_a"]
        vb = by_name["符合标准清单"]["value_b"]
        compact_a = va.replace(" ", "")
        compact_b = vb.replace(" ", "")
        assert "GB/T18268.1-2010" in compact_a
        assert "GB4793.1-2007" in compact_a
        assert "YY0648-2008" in compact_a
        assert "GB4793-2024" in compact_b
        assert "YY/T0648-2025" in compact_b
        assert by_name["符合标准清单"]["match"] == "不一致"
        assert "---PAGE---" not in va
        assert compact_a.count("GB4793.1-2007") == 1
