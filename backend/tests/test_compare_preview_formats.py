import os
import tempfile
import io

import fitz

from app.api.compare import _diff_page_boxes
from app.utils.compare_preview import build_highlighted_preview, extra_page_boxes


def _write(path, content):
    with open(path, "w", encoding="utf-8") as f:
        f.write(content)


def test_md_preview_marks_insert_and_change():
    with tempfile.TemporaryDirectory() as tmp:
        a = os.path.join(tmp, "a.md")
        b = os.path.join(tmp, "b.md")
        _write(a, "# Title\n\nHello world\n\nKeep me")
        _write(b, "# Title\n\nHello WORLD\n\nKeep me\n\nNew para")
        html_a = build_highlighted_preview(a, "a.md", b, "b.md", "a")
        html_b = build_highlighted_preview(b, "b.md", a, "a.md", "b")
        assert "hl-chg" in html_a
        assert "hl-ins" in html_b
        assert "New para" in html_b


def test_xlsx_preview_marks_changed_cell():
    from openpyxl import Workbook

    with tempfile.TemporaryDirectory() as tmp:
        a = os.path.join(tmp, "a.xlsx")
        b = os.path.join(tmp, "b.xlsx")
        wa = Workbook()
        sa = wa.active
        sa.title = "Sheet1"
        sa["A1"] = "Name"
        sa["B1"] = "Old"
        wa.save(a)
        wb = Workbook()
        sb = wb.active
        sb.title = "Sheet1"
        sb["A1"] = "Name"
        sb["B1"] = "New"
        sb["C1"] = "Extra"
        wb.save(b)
        html_a = build_highlighted_preview(a, "a.xlsx", b, "b.xlsx", "a")
        html_b = build_highlighted_preview(b, "b.xlsx", a, "a.xlsx", "b")
        assert "hl-chg" in html_a
        assert "hl-ins" in html_b
        assert "Extra" in html_b


def test_docx_preview_marks_changed_paragraph():
    from docx import Document

    with tempfile.TemporaryDirectory() as tmp:
        a = os.path.join(tmp, "a.docx")
        b = os.path.join(tmp, "b.docx")
        da = Document()
        da.add_heading("Intro", level=1)
        da.add_paragraph("Alpha text")
        da.save(a)
        db = Document()
        db.add_heading("Intro", level=1)
        db.add_paragraph("Beta text")
        db.add_paragraph("Added line")
        db.save(b)
        html_a = build_highlighted_preview(a, "a.docx", b, "b.docx", "a")
        html_b = build_highlighted_preview(b, "b.docx", a, "a.docx", "b")
        assert "hl-chg" in html_a or "hl-del" in html_a
        assert "hl-chg" in html_b or "hl-ins" in html_b
        assert "Added line" in html_b


def test_pdf_image_and_format_kinds():
    doc_a = fitz.open()
    doc_b = fitz.open()
    try:
        pa = doc_a.new_page(width=300, height=400)
        pb = doc_b.new_page(width=300, height=400)
        pa.insert_text((40, 60), "Same line", fontsize=12)
        pb.insert_text((40, 60), "Same line", fontsize=20)
        from PIL import Image

        def png_bytes(color):
            buf = io.BytesIO()
            Image.new("RGB", (16, 16), color).save(buf, format="PNG")
            return buf.getvalue()

        pa.insert_image(fitz.Rect(40, 120, 120, 200), stream=png_bytes((255, 0, 0)))
        pb.insert_image(fitz.Rect(40, 120, 120, 200), stream=png_bytes((0, 255, 0)))
        extra_a = extra_page_boxes(pa, pb, "a")
        extra_b = extra_page_boxes(pa, pb, "b")
        kinds_a = {k for _, k in extra_a}
        kinds_b = {k for _, k in extra_b}
        assert "fmt" in kinds_a
        assert "fmt" in kinds_b
        assert "img" in kinds_a
        assert "img" in kinds_b
        merged = [k for _, k in _diff_page_boxes(pa, pb, "a")]
        assert "fmt" in merged
        assert "img" in merged
    finally:
        doc_a.close()
        doc_b.close()
