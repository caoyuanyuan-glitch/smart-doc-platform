"""Preview helpers for compare: extra PDF kinds and Word/MD/Excel HTML diffs."""
import hashlib
import html
import os
import re
from difflib import SequenceMatcher


def _file_ext(name):
    return os.path.splitext(name or "")[1].lower()


def _norm_text(text):
    return re.sub(r"\s+", "", text or "").strip()


def _font_key(name):
    name = name or ""
    if "+" in name:
        name = name.split("+", 1)[-1]
    return name.lower()


def extract_page_images(page):
    items = []
    seen = set()
    doc = page.parent
    for img in page.get_images(full=True) or []:
        xref = img[0]
        if xref in seen:
            continue
        seen.add(xref)
        try:
            blob = (doc.extract_image(xref) or {}).get("image") or b""
        except Exception:
            blob = b""
        digest = hashlib.md5(blob).hexdigest() if blob else f"xref-{xref}"
        try:
            rects = list(page.get_image_rects(xref) or [])
        except Exception:
            rects = []
        for r in rects:
            items.append({"hash": digest, "bbox": (float(r.x0), float(r.y0), float(r.x1), float(r.y1))})
    if items:
        return items
    d = page.get_text("dict") or {}
    for block in d.get("blocks", []):
        if block.get("type") != 1:
            continue
        blob = block.get("image") or b""
        digest = hashlib.md5(blob).hexdigest() if blob else f"block-{len(items)}"
        bbox = tuple(block.get("bbox") or ())
        if len(bbox) == 4:
            items.append({"hash": digest, "bbox": bbox})
    return items


def diff_page_images(page_a, page_b, for_side):
    ia = extract_page_images(page_a)
    ib = extract_page_images(page_b)
    hashes_a = {x["hash"] for x in ia}
    hashes_b = {x["hash"] for x in ib}
    boxes = []
    src = ia if for_side == "a" else ib
    other = hashes_b if for_side == "a" else hashes_a
    for img in src:
        if img["hash"] not in other and img.get("bbox"):
            boxes.append((img["bbox"], "img"))
    return boxes


def extract_page_spans(page):
    spans = []
    d = page.get_text("dict") or {}
    for block in d.get("blocks", []):
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            for s in line.get("spans", []) or []:
                text = _norm_text(s.get("text") or "")
                if not text:
                    continue
                bbox = tuple(s.get("bbox") or ())
                if len(bbox) != 4:
                    continue
                spans.append({
                    "text": text,
                    "bbox": bbox,
                    "font": _font_key(s.get("font") or ""),
                    "size": round(float(s.get("size") or 0), 1),
                    "color": s.get("color"),
                })
    return spans


def diff_page_format(page_a, page_b, for_side):
    sa = extract_page_spans(page_a)
    sb = extract_page_spans(page_b)
    na = [x["text"] for x in sa]
    nb = [x["text"] for x in sb]
    sm = SequenceMatcher(None, na, nb, autojunk=False)
    boxes = []
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag != "equal":
            continue
        for k in range(i2 - i1):
            a = sa[i1 + k]
            b = sb[j1 + k]
            if a["font"] == b["font"] and a["size"] == b["size"] and a["color"] == b["color"]:
                continue
            src = a if for_side == "a" else b
            boxes.append((src["bbox"], "fmt"))
    return boxes


def extra_page_boxes(page_a, page_b, for_side):
    boxes = []
    boxes.extend(diff_page_images(page_a, page_b, for_side))
    boxes.extend(diff_page_format(page_a, page_b, for_side))
    return boxes


def _char_highlight(text_a, text_b, side):
    a = text_a or ""
    b = text_b or ""
    sm = SequenceMatcher(None, a, b, autojunk=False)
    parts = []
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if side == "a":
            chunk = html.escape(a[i1:i2])
            if not chunk:
                continue
            if tag == "equal":
                parts.append(chunk)
            elif tag == "delete":
                parts.append(f'<span class="hl-del">{chunk}</span>')
            elif tag == "replace":
                parts.append(f'<span class="hl-chg">{chunk}</span>')
        else:
            chunk = html.escape(b[j1:j2])
            if not chunk:
                continue
            if tag == "equal":
                parts.append(chunk)
            elif tag == "insert":
                parts.append(f'<span class="hl-ins">{chunk}</span>')
            elif tag == "replace":
                parts.append(f'<span class="hl-chg">{chunk}</span>')
    return "".join(parts)


_ATOMIC_TAGS = {"table", "pre", "blockquote"}


def highlight_block_html(blocks_a, blocks_b, side):
    na = [x.get("key", "") for x in blocks_a]
    nb = [x.get("key", "") for x in blocks_b]
    sm = SequenceMatcher(None, na, nb, autojunk=False)
    out = []
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal":
            src = blocks_a if side == "a" else blocks_b
            start = i1 if side == "a" else j1
            for k in range(i2 - i1):
                out.append(src[start + k].get("html") or "")
            continue
        if tag == "delete" and side == "a":
            for block in blocks_a[i1:i2]:
                out.append(f'<div class="hl-block hl-del">{block.get("html") or ""}</div>')
            continue
        if tag == "insert" and side == "b":
            for block in blocks_b[j1:j2]:
                out.append(f'<div class="hl-block hl-ins">{block.get("html") or ""}</div>')
            continue
        if tag != "replace":
            continue
        if (i2 - i1) == 1 and (j2 - j1) == 1:
            a = blocks_a[i1]
            b = blocks_b[j1]
            if a.get("tag") in _ATOMIC_TAGS or b.get("tag") in _ATOMIC_TAGS:
                src = a if side == "a" else b
                out.append(f'<div class="hl-block hl-chg">{src.get("html") or ""}</div>')
                continue
            inner = _char_highlight(a.get("text", ""), b.get("text", ""), side)
            tag_name = a.get("tag") if side == "a" else b.get("tag")
            tag_name = tag_name or "p"
            out.append(f"<{tag_name}>{inner}</{tag_name}>")
            continue
        if side == "a":
            for block in blocks_a[i1:i2]:
                out.append(f'<div class="hl-block hl-del">{block.get("html") or ""}</div>')
        else:
            for block in blocks_b[j1:j2]:
                out.append(f'<div class="hl-block hl-ins">{block.get("html") or ""}</div>')
    return "".join(out)


def _wrap_blocks(html_body, kind):
    return f'<div class="doc-preview-html doc-preview-{kind}">{html_body}</div>'


def _md_blocks(path):
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        text = f.read()
    from app.utils.md_compare import parse_markdown_blocks
    return parse_markdown_blocks(text)


def _docx_blocks(path):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(path)
    blocks = []

    def para_block(p):
        text = p.text or ""
        style = (p.style.name if p.style else "") or ""
        tag = "p"
        if style.startswith("Heading"):
            try:
                level = int(style.replace("Heading", "").strip() or "1")
            except Exception:
                level = 1
            tag = f"h{min(max(level, 1), 3)}"
        return {
            "key": _norm_text(text),
            "text": text,
            "tag": tag,
            "html": f"<{tag}>{html.escape(text) or '&nbsp;'}</{tag}>",
        }

    def table_block(table):
        rows_html = []
        keys = []
        texts = []
        for row in table.rows:
            cells = []
            for cell in row.cells:
                value = (cell.text or "").strip()
                keys.append(_norm_text(value))
                texts.append(value)
                cells.append(f"<td>{html.escape(value)}</td>")
            rows_html.append("<tr>" + "".join(cells) + "</tr>")
        text = "\n".join(texts)
        return {
            "key": "|".join(keys),
            "text": text,
            "tag": "table",
            "html": '<table class="doc-preview-table">' + "".join(rows_html) + "</table>",
        }

    body = doc.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            blocks.append(para_block(Paragraph(child, doc)))
        elif child.tag == qn("w:tbl"):
            blocks.append(table_block(Table(child, doc)))
    return blocks


def _cell_str(value):
    if value is None:
        return ""
    return str(value)


def _xlsx_sheet_maps(path, max_rows=200, max_cols=40):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheets = []
    try:
        for ws in wb.worksheets:
            rows = []
            for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
                if i > max_rows:
                    break
                cells = [_cell_str(v) for v in list(row)[:max_cols]]
                rows.append(cells)
            sheets.append({"title": ws.title, "rows": rows})
    finally:
        wb.close()
    return sheets


def _xlsx_highlight_html(path, peer_path, side):
    sheets_a = _xlsx_sheet_maps(path if side == "a" else peer_path)
    sheets_b = _xlsx_sheet_maps(peer_path if side == "a" else path)
    this = sheets_a if side == "a" else sheets_b
    other = sheets_b if side == "a" else sheets_a
    other_by_name = {s["title"]: s for s in other}
    used = set()
    sections = []
    for idx, sheet in enumerate(this):
        peer = other_by_name.get(sheet["title"])
        if peer:
            used.add(sheet["title"])
        elif idx < len(other) and other[idx]["title"] not in used:
            peer = other[idx]
            used.add(peer["title"])
        else:
            peer = {"title": "", "rows": []}
        rows_html = []
        max_r = max(len(sheet["rows"]), len(peer["rows"]))
        for r in range(max_r):
            this_row = sheet["rows"][r] if r < len(sheet["rows"]) else []
            peer_row = peer["rows"][r] if r < len(peer["rows"]) else []
            max_c = max(len(this_row), len(peer_row))
            cells = []
            for c in range(max_c):
                this_val = this_row[c] if c < len(this_row) else ""
                peer_val = peer_row[c] if c < len(peer_row) else ""
                va = this_val if side == "a" else peer_val
                vb = peer_val if side == "a" else this_val
                value = this_val
                cls = ""
                if side == "a":
                    if va and not vb:
                        cls = "hl-del"
                    elif va and vb and va != vb:
                        cls = "hl-chg"
                else:
                    if vb and not va:
                        cls = "hl-ins"
                    elif va and vb and va != vb:
                        cls = "hl-chg"
                attr = f' class="{cls}"' if cls else ""
                inner = html.escape(value)
                if cls == "hl-chg":
                    inner = _char_highlight(va, vb, side)
                cells.append(f"<td{attr}>{inner}</td>")
            if cells:
                rows_html.append("<tr>" + "".join(cells) + "</tr>")
        sections.append(
            f'<section class="doc-preview-sheet"><h4>{html.escape(sheet["title"])}</h4>'
            f'<table class="doc-preview-table">{"".join(rows_html)}</table></section>'
        )
    return _wrap_blocks("".join(sections), "xlsx")


def build_highlighted_preview(path, file_name, peer_path, peer_name, side):
    ext = _file_ext(file_name) or _file_ext(path)
    peer_ext = _file_ext(peer_name) or _file_ext(peer_path)
    if ext in {".xlsx", ".xlsm"} and peer_ext in {".xlsx", ".xlsm"}:
        return _xlsx_highlight_html(path, peer_path, side)
    if ext in {".md", ".markdown", ".txt"} and peer_ext in {".md", ".markdown", ".txt"}:
        blocks_a = _md_blocks(path if side == "a" else peer_path)
        blocks_b = _md_blocks(peer_path if side == "a" else path)
        if side == "b":
            body = highlight_block_html(blocks_a, blocks_b, "b")
        else:
            body = highlight_block_html(blocks_a, blocks_b, "a")
        return _wrap_blocks(body, "md")
    if ext == ".docx" and peer_ext == ".docx":
        blocks_a = _docx_blocks(path if side == "a" else peer_path)
        blocks_b = _docx_blocks(peer_path if side == "a" else path)
        body = highlight_block_html(blocks_a, blocks_b, "a" if side == "a" else "b")
        return _wrap_blocks(body, "docx")
    return None
