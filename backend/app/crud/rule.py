import re
from pathlib import Path

from sqlalchemy.orm import Session
from app.models.rule import Rule
from app.schemas.rule import RuleCreate, RuleUpdate


REVIEW_RULE_LIBRARY_SEED_PATH = Path(__file__).resolve().parents[2] / "seed" / "review_rule_library_seed.xlsx"
REVIEW_RULE_LIBRARY_SEED_SHEET = "规则库"
REVIEW_RULE_LIBRARY_SEED_META_SHEET = "元信息"
REVIEW_RULE_LIBRARY_SEED_COLUMNS = {
    "rule_id": "规则ID",
    "category": "分类",
    "severity": "严重程度",
    "rule_content": "规则内容",
    "applicable_scenarios": "适用场景",
}
SEED_SCENARIO_DELIMITER = "、"

# 严重程度映射：种子数据中的中文 → 系统英文标识
SEVERITY_MAP = {
    "致命": "fatal",
    "严重": "serious",
    "一般": "general",
    "建议": "suggestion",
}


DISABLED_RULE_REGEX = r"(?!)"

# 规则描述特征 -> 实际用于扫描文档的正则。第一项为按正则匹配规则描述的键。
# 只有能确定性表达成模式匹配的规则才给出正则；语义、版式类规则统一返回
# DISABLED_RULE_REGEX，改由 AI 审核或人工判断，不能靠猜测生成模式。
RULE_PATTERN_MAP = (
    # 版式/视觉类规则无法用单一正则判定
    (r"图标文字底部|图片和图注|图片中标注文字|图片中文字体", DISABLED_RULE_REGEX),
    (r"产品中有害物质的名称及含有物质表", DISABLED_RULE_REGEX),
    (r"仅可交互UI元素", DISABLED_RULE_REGEX),
    (r"统一使用双引号.*单引号", DISABLED_RULE_REGEX),
    (r"标点符号", DISABLED_RULE_REGEX),
    (r"公司官网地址", r"https?://[^\s]+mgi[^\s]*"),
    (r"多余的?空格|空行", r"[ \t]{2,}|\n{3,}"),
    (r"乘号", r"\*[×xX]?\s*\d+|\d+\s*\*"),
    (r"现成.*现场", r"现场(?:情况)?"),
    (r"不避免", r"不避免"),
    (r"手工冰箱", r"手工冰箱"),
    (r"使用限期|限期.*期限", r"限期"),
    (r"成语", r"周而复始|恰如其分|千丝万缕|不言而喻|一目了然|举足轻重"),
    (r"文言化", r"未尽事宜|鉴于|据此|兹"),
    (r"Cat\.?\s*No", r"Cat\.?\s*No\.?"),
)


def _convert_rule_content_to_regex(rule_content: str) -> str:
    """把外部规则库的规则描述转成可执行正则。

    只转换能确定性表达成模式匹配的规则。历史上这里会把规则描述切成 2-4 字的中文
    片段拼接成正则，这类正则只能匹配规则描述自身的措辞：对文档只会产生误报
    （例如把正确的“期限”判为问题、匹配到本应存在的表格标题片段），对英文文档
    则完全无效。
    """
    content = str(rule_content or "").strip()
    if not content:
        return DISABLED_RULE_REGEX

    # 明确声明“不列为错误/问题”的条目属于误报抑制口径，不参与匹配
    if re.search(r"不列为(?:错误|问题)|不属于错误|属于转换 ?artifact", content):
        return DISABLED_RULE_REGEX

    for keyword_pattern, scan_pattern in RULE_PATTERN_MAP:
        if re.search(keyword_pattern, content):
            return scan_pattern
    return DISABLED_RULE_REGEX


def _rule_language_for_pattern(pattern: str) -> str:
    """外部规则库是中文评审规则库，只有模式本身含拉丁字符时才适用于英文文档。"""
    if pattern == DISABLED_RULE_REGEX:
        return "cn"
    return "both" if re.search(r"[A-Za-z0-9]", pattern) else "cn"


def _load_review_rule_library_seed():
    """读取 Excel 种子，返回 (元信息, 规则列表)。"""
    from openpyxl import load_workbook

    workbook = load_workbook(REVIEW_RULE_LIBRARY_SEED_PATH, data_only=True)
    sheet = (
        workbook[REVIEW_RULE_LIBRARY_SEED_SHEET]
        if REVIEW_RULE_LIBRARY_SEED_SHEET in workbook.sheetnames
        else workbook.worksheets[0]
    )
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return {}, []

    header = [str(value or "").strip() for value in rows[0]]
    position = {
        field: header.index(title)
        for field, title in REVIEW_RULE_LIBRARY_SEED_COLUMNS.items()
        if title in header
    }

    def cell(row, field):
        index = position.get(field)
        if index is None or index >= len(row):
            return ""
        return str(row[index] or "").strip()

    rules = []
    for row in rows[1:]:
        if not any(str(value or "").strip() for value in row):
            continue
        rules.append({
            "rule_id": cell(row, "rule_id"),
            "category": cell(row, "category"),
            "severity": cell(row, "severity"),
            "rule_content": cell(row, "rule_content"),
            "applicable_scenarios": [
                item for item in cell(row, "applicable_scenarios").split(SEED_SCENARIO_DELIMITER) if item
            ],
        })

    meta = {}
    if REVIEW_RULE_LIBRARY_SEED_META_SHEET in workbook.sheetnames:
        for row in workbook[REVIEW_RULE_LIBRARY_SEED_META_SHEET].iter_rows(values_only=True):
            if row and str(row[0] or "").strip():
                meta[str(row[0]).strip()] = str(row[1] or "").strip() if len(row) > 1 else ""
    return meta, rules


def seed_external_review_rules(db: Session):
    if not REVIEW_RULE_LIBRARY_SEED_PATH.exists():
        return 0

    meta, seed_rules = _load_review_rule_library_seed()
    source = meta.get("来源") or "外部评审规则库"
    export_date = meta.get("导出日期", "")
    created = 0
    updated = 0

    for item in seed_rules:
        original_rule_id = str(item.get("rule_id", "")).strip()
        if not original_rule_id:
            continue

        rule_no = f"EXT-{original_rule_id}"
        rule_content = item.get("rule_content") or ""
        category = item.get("category") or "其他"
        chinese_severity = item.get("severity", "一般")
        severity = SEVERITY_MAP.get(chinese_severity, "general")
        scenarios = "、".join(item.get("applicable_scenarios") or []) or "通用"

        # 将规则内容转为可执行的正则表达式
        regex = _convert_rule_content_to_regex(rule_content)
        language = _rule_language_for_pattern(regex)

        example = f"来源: {source} | 适用场景: {scenarios}"
        audit_basis = f"{source}{' | 导出日期: ' + export_date if export_date else ''}"
        existing = get_rule_by_no(db, rule_no)
        if existing:
            changed = False
            updates = {
                "category": category,
                "description": rule_content,
                "regex": regex,
                "example": example,
                "suggestion": rule_content,
                "audit_basis": audit_basis,
                "severity": severity,
                "language": language,
            }
            for field, value in updates.items():
                if getattr(existing, field) != value:
                    setattr(existing, field, value)
                    changed = True
            if changed:
                updated += 1
            continue

        db.add(Rule(
            rule_no=rule_no,
            category=category,
            description=rule_content,
            regex=regex,
            example=example,
            suggestion=rule_content,
            audit_basis=audit_basis,
            severity=severity,
            language=language,
        ))
        created += 1

    if created or updated:
        db.commit()
    return created + updated

def create_rule(db: Session, rule: RuleCreate):
    db_rule = Rule(
        rule_no=rule.rule_no,
        category=rule.category,
        description=rule.description,
        regex=rule.regex,
        example=rule.example,
        suggestion=rule.suggestion,
        audit_basis=rule.audit_basis,
        severity=rule.severity,
        language=rule.language
    )
    db.add(db_rule)
    db.commit()
    db.refresh(db_rule)
    return db_rule

def get_rule(db: Session, rule_id: int):
    return db.query(Rule).filter(Rule.id == rule_id).first()

def get_rule_by_no(db: Session, rule_no: str):
    return db.query(Rule).filter(Rule.rule_no == rule_no).first()

def get_rules(db: Session, skip: int = 0, limit: int = 100):
    return db.query(Rule).offset(skip).limit(limit).all()

def update_rule(db: Session, rule_id: int, rule_update: RuleUpdate):
    rule = db.query(Rule).filter(Rule.id == rule_id).first()
    if rule:
        if rule_update.category is not None:
            rule.category = rule_update.category
        if rule_update.description is not None:
            rule.description = rule_update.description
        if rule_update.regex is not None:
            rule.regex = rule_update.regex
        if rule_update.example is not None:
            rule.example = rule_update.example
        if rule_update.suggestion is not None:
            rule.suggestion = rule_update.suggestion
        if rule_update.audit_basis is not None:
            rule.audit_basis = rule_update.audit_basis
        if rule_update.severity is not None:
            rule.severity = rule_update.severity
        if rule_update.language is not None:
            rule.language = rule_update.language
        db.commit()
        db.refresh(rule)
    return rule

def delete_rule(db: Session, rule_id: int):
    rule = db.query(Rule).filter(Rule.id == rule_id).first()
    if rule:
        db.delete(rule)
        db.commit()
    return rule

def bulk_create_rules(db: Session, rules: list[RuleCreate]):
    db_rules = []
    seen_rule_nos = set()
    for rule in rules:
        # rule_no 唯一：跳过库中已存在的，也跳过同一批里重复出现的，
        # 否则 add_all + commit 会触发唯一约束错误。
        if rule.rule_no in seen_rule_nos or get_rule_by_no(db, rule.rule_no):
            continue
        seen_rule_nos.add(rule.rule_no)
        db_rules.append(Rule(
            rule_no=rule.rule_no,
            category=rule.category,
            description=rule.description,
            regex=rule.regex,
            example=rule.example,
            suggestion=rule.suggestion,
            audit_basis=rule.audit_basis,
            severity=rule.severity,
            language=rule.language
        ))
    if db_rules:
        db.add_all(db_rules)
        db.commit()
    return len(db_rules)

def bulk_delete_rules(db: Session, rule_ids: list[int]):
    count = 0
    for rule_id in rule_ids:
        rule = db.query(Rule).filter(Rule.id == rule_id).first()
        if rule:
            db.delete(rule)
            count += 1
    db.commit()
    return count
